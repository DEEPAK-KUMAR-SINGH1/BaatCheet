import io
import asyncio
import sys
from pathlib import Path

import pytest


pytest.importorskip("langchain_chroma")

BACKEND_DIR = Path(__file__).resolve().parents[1] / "backend"
sys.path.insert(0, str(BACKEND_DIR))


@pytest.fixture()
def isolated_db(tmp_path, monkeypatch):
    import db

    test_db = tmp_path / "chatbot-rag-test.db"
    monkeypatch.setattr(db, "DATABASE_PATH", str(test_db))
    db.init_all_tables()
    return db


def test_extracts_text_csv_markdown_doc_docx_and_pdf(tmp_path, monkeypatch):
    import fitz
    from docx import Document

    import rag_engine

    markdown = tmp_path / "sample.md"
    markdown.write_text("# Notes\n\nMarkdown retrieval needle.", encoding="utf-8")

    csv_file = tmp_path / "sample.csv"
    csv_file.write_text("name,value\nneedle,42\n", encoding="utf-8")

    doc_file = tmp_path / "sample.doc"
    doc_file.write_bytes(b"legacy")
    monkeypatch.setattr(rag_engine, "_extract_doc", lambda path: "legacy DOC retrieval needle")

    docx_file = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("DOCX paragraph retrieval needle.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Table key"
    table.rows[0].cells[1].text = "Table value"
    doc.save(docx_file)

    pdf_file = tmp_path / "sample.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "PDF retrieval needle.")
    pdf.save(pdf_file)

    assert "Markdown retrieval needle" in rag_engine.extract_document_pages(str(markdown), markdown.name)[0]["text"]
    assert "Row 1: name=needle; value=42" in rag_engine.extract_document_pages(str(csv_file), csv_file.name)[0]["text"]
    assert "legacy DOC retrieval needle" in rag_engine.extract_document_pages(str(doc_file), doc_file.name)[0]["text"]

    docx_text = rag_engine.extract_document_pages(str(docx_file), docx_file.name)[0]["text"]
    assert "DOCX paragraph retrieval needle" in docx_text
    assert "Table key | Table value" in docx_text

    assert "PDF retrieval needle" in rag_engine.extract_document_pages(str(pdf_file), pdf_file.name)[0]["text"]


def test_retrieval_falls_back_to_indexed_chunks_when_vector_search_fails(isolated_db, monkeypatch):
    import database
    import rag_engine

    database.create_thread("thread-a", "owner@example.com")
    database.create_workspace("owner@example.com", "Knowledge", "workspace-a")
    database.save_document_record(
        doc_id="doc-a",
        thread_id="thread-a",
        workspace_id="workspace-a",
        filename="policy.md",
        file_type=".md",
        status="indexed",
        chunk_count=1,
    )
    database.replace_document_chunks("doc-a", [
        {
            "doc_id": "doc-a",
            "chunk_index": 0,
            "page": 1,
            "content": "The warranty period is 24 months for the retrieval needle.",
        }
    ])

    def broken_vectorstore(_workspace_id):
        raise RuntimeError("embedding service unavailable")

    monkeypatch.setattr(rag_engine, "get_vectorstore", broken_vectorstore)

    context, sources = rag_engine.retrieve_context_with_sources("workspace-a", "warranty period", k=3)

    assert "24 months" in context
    assert sources[0]["filename"] == "policy.md"
    assert sources[0]["chunk_index"] == 0


def test_upload_returns_failed_document_record_for_retry(isolated_db, tmp_path, monkeypatch):
    from fastapi import HTTPException, UploadFile

    import database
    import document_service

    database.create_thread("thread-a", "owner@example.com")
    database.create_workspace("owner@example.com", "Knowledge", "workspace-a")
    monkeypatch.setattr(document_service, "UPLOADS_DIR", str(tmp_path))

    def fail_index(doc_id, *, workspace_id, user_id=None):
        database.update_document_status(doc_id, "failed", error="synthetic indexing error")
        raise HTTPException(400, "synthetic indexing error")

    monkeypatch.setattr(document_service, "index_existing_document", fail_index)

    upload = UploadFile(filename="notes.md", file=io.BytesIO(b"# Notes"))
    doc = asyncio.run(
        document_service.upload_and_index_document(
            thread_id="thread-a",
            workspace_id="workspace-a",
            file=upload,
            user_id="owner@example.com",
            raise_on_failure=False,
        )
    )

    assert doc["status"] == "failed"
    assert doc["error"] == "synthetic indexing error"
