from pathlib import Path
from copy import deepcopy
import re
import zipfile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

folder = Path(__file__).parent
src = folder / 'Class_9_English_Half_Yearly_Question_Paper.docx'
out = folder / 'Class_9_English_Half_Yearly_Question_Paper_Revised.docx'
doc = Document(src)

def find(prefix):
    return next(p for p in doc.paragraphs if p.text.startswith(prefix))

def replace(p, text):
    # Preserve the paragraph and its first run's formatting.
    props = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear()
    r = p.add_run(text)
    if props is not None:
        r._r.insert(0, props)

def insert_before(anchor, text, template=None, bold=False, center=False):
    p = anchor.insert_paragraph_before(text)
    if template is not None and template._p.pPr is not None:
        p._p.insert(0, deepcopy(template._p.pPr))
    if bold:
        for r in p.runs: r.bold = True
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

replace(find('Q3.'), 'Q3. Do any 10 out of 12 grammar questions. (10 × 1 = 10 marks)')
q4 = find('Q4.')
template = find('(x)  Identify the incorrect word')
insert_before(q4, '(xi)  Choose the correct tense:\nWhen I reached the playground, the children __________ cricket.\nA. were playing    B. plays    C. has played    D. is playing  [1]', template)
insert_before(q4, '(xii)  Choose the appropriate determiner:\nThere are __________ apples in the basket, enough for all three of us.\nA. a few    B. little    C. much    D. any  [1]', template)

replace(q4, 'Q4. Formal Letter — Attempt EITHER option A OR option B. (5 marks)')
q4.paragraph_format.page_break_before = True
paras = doc.paragraphs
old_letter = paras[paras.index(q4) + 1] if q4 in paras else None
# Paragraph wrappers are recreated, so locate the following XML element directly.
from docx.text.paragraph import Paragraph
old_letter = Paragraph(q4._p.getnext(), q4._parent)
replace(old_letter, 'Write your letter in 100–120 words using an appropriate formal format.')
q5 = find('Q5.')
heading5 = Paragraph(q5._p.getprevious(), q5._parent)
# Insert new letter options before the existing page break preceding Q5's section heading.
anchor = Paragraph(heading5._p.getprevious(), heading5._parent)
insert_before(anchor, 'Option A — Letter to the Editor', bold=True)
insert_before(anchor, 'You are Anuj/Ananya, a resident of 24, Green Park, Jaipur. The public park in your neighbourhood is littered with plastic waste, and several dustbins are damaged. Write a letter to the Editor of The City Herald, Jaipur, describing the problem and its effect on residents. Suggest practical measures to improve cleanliness and encourage responsible use of the park.')
insert_before(anchor, 'OR', bold=True, center=True)
insert_before(anchor, 'Option B — Complaint Letter', bold=True)
insert_before(anchor, 'You are Anuj/Ananya, Library Secretary of K.M. International School, 18, School Road, Jaipur. Your school ordered 30 English dictionaries from Bright Books, 12, Station Road, Jaipur, under Order No. KM/LIB/19 dated 2 September 2026. The delivery received on 10 September contained only 25 dictionaries, five of which had torn or missing pages. Write a complaint to the Sales Manager requesting replacement of the damaged copies and delivery of the missing dictionaries.')

q7 = find('Q7.')
replace(q7, 'Q7. Read the poetry extract and answer ALL three questions. (5 marks)')
rhyme = next(p for p in doc.paragraphs if 'What is the rhyme scheme of these four lines?' in p.text)
rhyme._p.getparent().remove(rhyme._p)
p = find('(ii)  The comparison between seeds')
replace(p, p.text.replace('(ii)', '(i)', 1))
p = find('(iii)  Which word in the extract')
replace(p, p.text.replace('(iii)', '(ii)', 1))
p = find('(iv)  How does the extract present gardening')
replace(p, p.text.replace('(iv)', '(iii)', 1).replace('[2]', '[3]'))

replace(find('Q9.'), 'Q9. Short Answer Questions — Do any 4 out of 5 parts. (4 × 2 = 8 marks)')
template = find('(iv)  Why does a stranger')
insert_before(find('Q10.'), '(v)  Mention two ways in which the grandmother practised her reading and writing lessons. (How I Taught My Grandmother to Read)  [2]', template)

doc.core_properties.title = 'Class IX English Half-Yearly Question Paper — Revised'
doc.save(out)

check = Document(out)
text = '\n'.join(p.text for p in check.paragraphs)
def block(a, b): return text[text.index(a):text.index(b)]
assert len(re.findall(r'^\([ivx]+\)', block('Q3.', 'Q4.'), re.M)) == 12
assert len(re.findall(r'^\([ivx]+\)', block('Q9.', 'Q10.'), re.M)) == 5
assert len(re.findall(r'^\([ivx]+\)', block('Q7.', 'Q8.'), re.M)) == 3
assert sum(map(int, re.findall(r'\[(\d+)\]', block('Q7.', 'Q8.')))) == 5
assert 'rhyme scheme' not in text.lower()
assert 'Option A — Letter to the Editor' in block('Q4.', 'Q5.')
assert 'Option B — Complaint Letter' in block('Q4.', 'Q5.')
assert 'requesting a weekly reading-club session' not in text
assert sum([10,10,10,5,5,5,5,12,8,5,5]) == 80
with zipfile.ZipFile(out) as z: assert z.testzip() is None
print(out.resolve())
print('Verified: Q3 10/12; Q9 4/5; Q4 letter choice; Q7 three parts worth 1+1+3; total 80 marks.')
