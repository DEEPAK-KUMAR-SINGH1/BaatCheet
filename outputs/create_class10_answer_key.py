from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import re, zipfile

ROOT = Path(__file__).parent
SOURCE = ROOT / 'Class_10_English_Half_Yearly_Question_Paper.docx'
OUT = ROOT / 'Class_10_English_Half_Yearly_Answer_Key_and_Marking_Scheme.docx'
original = Document(SOURCE)
original_text = '\n'.join(p.text for p in original.paragraphs)
assert 'bridge-builder' in original_text and 'By the time the teacher arrived' in original_text
assert 'SET B' not in original_text

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.7)
sec.header_distance = sec.footer_distance = Inches(.28)
style=doc.styles['Normal']; style.font.name='Times New Roman'; style.font.size=Pt(11)
style.paragraph_format.space_after=Pt(6); style.paragraph_format.line_spacing=1.03
for name in ['Heading 1','Heading 2']:
    s=doc.styles[name]; s.font.name='Times New Roman'; s.font.size=Pt(12)
    s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_before=Pt(9); s.paragraph_format.space_after=Pt(7)

def p(text='',bold=False,size=None,center=False):
    para=doc.add_paragraph(); r=para.add_run(text); r.bold=bold
    if size: r.font.size=Pt(size)
    if center: para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return para
def h(text): doc.add_paragraph(text,'Heading 2')
def page(): doc.add_page_break()
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'
    for c,text in zip(t.rows[0].cells,headers):
        c.text=str(text)
        for r in c.paragraphs[0].runs: r.bold=True
    header=t.rows[0]._tr.get_or_add_trPr(); header.append(OxmlElement('w:tblHeader'))
    for values in rows:
        for c,text in zip(t.add_row().cells,values): c.text=str(text)
    if widths:
        t.autofit=False
        for col,w in zip(t.columns,widths): col.width=Inches(w)
        for row in t.rows:
            for c,w in zip(row.cells,widths): c.width=Inches(w)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            for para in c.paragraphs:
                para.paragraph_format.space_before=Pt(3)
                para.paragraph_format.space_after=Pt(3)
                for r in para.runs: r.font.size=Pt(10.5)
    p('')
def link(label,url):
    para=doc.add_paragraph(); para.paragraph_format.space_after=Pt(4)
    a=OxmlElement('w:hyperlink'); a.set(qn('r:id'),para.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); props=OxmlElement('w:rPr')
    size=OxmlElement('w:sz'); size.set(qn('w:val'),'18'); props.append(size); r.append(props)
    t=OxmlElement('w:t'); t.text=label; r.append(t); a.append(r); para._p.append(a)

hp=sec.header.paragraphs[0]; hp.text='K.M. INTERNATIONAL SCHOOL | CLASS X ENGLISH | TEACHER COPY'
hp.alignment=WD_ALIGN_PARAGRAPH.CENTER; hp.runs[0].font.size=Pt(9)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
fp.add_run('Answer Key & Marking Scheme | Page ')
f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); fp._p.append(f)
fp.add_run(' of '); f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'NUMPAGES'); fp._p.append(f)

p('K.M. INTERNATIONAL SCHOOL',True,15,True)
p('HALF-YEARLY EXAMINATION — CLASS X',True,13,True)
p('ENGLISH LANGUAGE & LITERATURE (184)',True,12,True)
p('ANSWER KEY AND MARKING SCHEME',True,13,True)
p('Maximum Marks: 80 | Examination Time: 3 Hours',True,11,True)
p('For the original question paper supplied in this request (not Set B).',center=True)
h('Mark allocation')
table(['Question','Component / attempt requirement','Marks'],[
    ['Q1','Discursive passage: all 10 parts',10],
    ['Q2','Case-based passage: all 10 parts',10],
    ['Q3','Grammar: any 10 of 12',10],
    ['Q4','Analytical paragraph',5],
    ['Q5','Formal letter: either A or B',5],
    ['Q6','Prose extract: 1 + 1 + 2',4],
    ['Q7','Poetry extract: 1 + 1 + 1 + 1 + 2',6],
    ['Q8','Two prose + two poetry answers; 3 marks each',12],
    ['Q9','Any 4 of 5 supplementary answers; 2 marks each',8],
    ['Q10','Either A or B; 120–150 words',6],
    ['Q11','Either A or B; 100–120 words',4],
    ['TOTAL','Reading 20 + Grammar/Writing 20 + Literature 40',80]
],[.7,5.4,.75])
h('Guidance for examiners')
for t in [
    'This is a marking scheme prepared for this school paper, not an official CBSE marking scheme. The answers are indicative; reward accurate paraphrases and defensible interpretations supported by the text.',
    'For objective items, award 1 or 0. A clearly correct option letter or answer is sufficient; if both are given and conflict, award 0. For written responses, use the stated point allocations; half marks may be used where indicated or for partially developed points.',
    'For this scheme, when extra answers are attempted, count the best permitted number, respecting Q8’s two-prose/two-poetry requirement. A school rule announced before the examination takes precedence.',
    'Do not deduct twice for the same error. Minor language slips should not lose comprehension/content marks when meaning is clear. Assess language separately where an expression component is provided.',
    'Treat word limits as guidance for focus and development. Do not impose a separate automatic penalty; assess omissions, repetition or irrelevance under the relevant content/organisation criterion.'
]: p(t,size=10)

page(); h('SECTION A — READING SKILLS (20 MARKS)')
h('Q1. Discursive passage — Learning from mistakes (10 marks)')
p('All ten parts: 1 mark each. Accept equivalent wording for short answers.')
q1=[
    ['(i)','B — Examining and correcting mistakes supports learning.'],
    ['(ii)','The difference between what we intended to do and what actually happened. Both sides of the comparison must be conveyed.'],
    ['(iii)','She has identified where the bridge bends and changes the supports to correct the weakness / test an improved design.'],
    ['(iv)','Reflection.'],
    ['(v)','C — Explain how your example supports the main point.'],
    ['(vi)','False. A supportive classroom still maintains standards of correctness.'],
    ['(vii)','Fear of ridicule / being laughed at / humiliation. Any one equivalent expression.'],
    ['(viii)','Acknowledge.'],
    ['(ix)','Mistakes help only when learners examine and correct them; carelessly repeating an error does not produce learning.'],
    ['(x)','Apply the feedback to revise the work; identify the weakness and make a relevant change before resubmitting. Accept one concrete action consistent with this advice.']
]
table(['Part','Expected answer'],q1,[.6,6.25])
p('Partial-credit guidance: for Q1(ii), award ½ if only intended action or actual outcome is stated. For Q1(ix), award ½ for a relevant but incomplete distinction. Do not require explanations for the MCQ or true/false items.',size=10)

page(); h('Q2. Case-based passage — Travel to school (10 marks)')
p('All ten parts: 1 mark each. Correct numerical answers receive full credit without working unless the answer itself is ambiguous.')
q2=[
    ['(i)','A — To count each student once in the totals.'],
    ['(ii)','School bus (70 students in July; 80 in August). Naming the mode is sufficient.'],
    ['(iii)','20 students: 50 − 30 = 20.'],
    ['(iv)','15%: (30 ÷ 200) × 100 = 15%.'],
    ['(v)','50% decrease: [(60 − 30) ÷ 60] × 100 = 50%.'],
    ['(vi)','False. Public-transport users decreased from 20 to 10.'],
    ['(vii)','Any one: long distance from school; unsafe road crossings; family schedules that limit other options.'],
    ['(viii)','The survey measures travel choices, not actual emissions, fuel consumption or journey lengths; therefore it cannot quantify pollution reduction.'],
    ['(ix)','B — The combined number walking and cycling increased (50 to 80).'],
    ['(x)','Provide safer road crossings, such as a supervised crossing, crossing guard or suitable pedestrian crossing. Accept one relevant safety measure.']
]
table(['Part','Expected answer'],q2,[.6,6.25])
h('Calculation and interpretation notes')
p('Q2(iv): ½ for the correct fraction 30/200 and ½ for the correct percentage if the calculation is shown but incomplete. Q2(v): ½ for a correct percentage-decrease setup and ½ for 50%. Do not confuse the 15-percentage-point fall in car use (30% to 15%) with the requested 50% relative decrease.',size=10)
p('Q2(viii): award ½ for the incomplete observation that pollution was not directly measured; award 1 when the mismatch between the survey data and the claimed pollution reduction is clear. Q2(x): “better bus access” alone does not directly resolve the unsafe-crossing problem.',size=10)

page(); h('SECTION B — GRAMMAR & WRITING (20 MARKS)')
h('Q3. Grammar — Any 10 of 12 (10 marks)')
q3=[
    ['(i)','B — little','1'],['(ii)','B — had completed','1'],
    ['(iii)','C — must not','1'],['(iv)','B — receives','1'],
    ['(v)','B — since','1'],['(vi)','A — either','1'],
    ['(vii)','are planting','1'],['(viii)','are → is: The list of selected candidates is on the notice board.','1'],
    ['(ix)','win → winning: Our team is confident of winning the final match.','1'],
    ['(x)','was preparing (expected backshift)','1'],
    ['(xi)','B — could','1'],
    ['(xii)','Delete about: We discussed the proposal during the meeting.','1']
]
table(['Part','Answer / correction','Marks'],q3,[.6,5.6,.65])
h('Acceptable alternatives and scoring')
p('Award 1 for the correct answer/correction and 0 otherwise; do not split the mark for identifying an error without correcting it. Accept a fully corrected sentence even if the student does not use an error → correction table.')
p('Q3(x): “is preparing” is also grammatical if the reported activity is still current. Because the question provides no reporting-time context, accept this unshifted form as an alternative; “was preparing” is the standard school backshift answer.')
p('Q3(xii): also accept discussed → talked or spoke, producing “We talked/spoke about the proposal during the meeting.” These repair the sentence by changing one word. Do not insist on deletion as the only solution.')
p('Maximum for Q3 remains 10, even if all twelve items are answered correctly.',True)

page(); h('Q4. Analytical paragraph (5 marks)')
p('Required length: 100–120 words. Assess a connected analytical paragraph, not a list of unexplained figures.')
table(['Criterion','Allocation'],[
    ['Content — 2','1: accurate overview and key figures, including the most/least preferred activities. 1: at least two relevant comparisons or grouped observations, without invented causes.'],
    ['Organisation — 2','1: clear opening and logically ordered comparisons. 1: cohesive links and an overall observation; an integrated concluding sentence is sufficient.'],
    ['Accuracy — 1','Grammar, vocabulary, spelling and punctuation support clear meaning.']
],[1.45,5.4])
h('Model answer')
models={}
models['Q4']='The table presents the main leisure preferences of 200 Class X students. Sports and outdoor games lead with 60 students, or 30%, while art and music attract the fewest, at 20 students, or 10%. Social media ranks second with 50 students, followed by films and online videos with 40. Reading accounts for 30 students, representing 15% of the sample. Sports attract twice as many students as reading and three times as many as art and music. Together, social media and viewing films or videos account for 45%, exceeding sports alone by 15 percentage points. Overall, the preferences vary considerably across the five activities.'
p(models['Q4'])
h('Useful content checks')
p('Sports 60/30%; social media 50/25%; films/videos 40/20%; reading 30/15%; art/music 20/10%. Total 200/100%. Other accurate comparisons are equally valid. A student need not reproduce every statistic for full marks.')
p('Do not reward unsupported explanations such as “students dislike reading because it is boring”, or treat the survey as a record of exact hours spent on each activity. It records each respondent’s main leisure activity.',size=10)
h('Q5. Formal letter — Common rubric (5 marks)')
table(['Criterion','Marks / basis'],[
    ['Format','1: appropriate addresses, date, subject, salutation and formal closing; ½ for a substantially recognisable but incomplete format.'],
    ['Content','2: use the option-specific allocation on the following pages.'],
    ['Organisation and tone','1: logical sequence, clear purpose/request and suitable formal tone.'],
    ['Language accuracy','1: effective sentence construction, spelling and punctuation; ½ for noticeable errors that do not obscure meaning.']
],[1.6,5.25])
p('Assess either A or B. Accept Aarav/Aarohi or Rohan/Riya as appropriate. Sensible dates and conventional formal closings are acceptable. Model body paragraphs follow the 100–120-word limit; address blocks and the closing are shown separately.',size=10)

page(); h('Q5(A). Letter to the Editor — Model and content allocation')
p('Content: 1 mark for the obstruction and its effect on student safety; 1 mark for practical remedies and a clear request for attention/action.')
p('24, Green Park\nJaipur\n16 September 2026\n\nThe Editor\nThe City Herald\nJaipur')
p('Subject: Unsafe parking on the footpath near our school',True)
p('Sir/Madam,')
models['Q5A']='Through the columns of your newspaper, I wish to draw attention to vehicles parked on the footpath near our school. The obstruction forces children to walk on the road alongside moving traffic, especially during arrival and dispersal. It also leaves little safe space for younger pupils and other pedestrians.\n\nThe authorities should enforce parking rules and keep the footpath clear. A designated drop-off area, visible signs and regular monitoring would help. The school could also request parents and drivers to avoid blocking pedestrian access.\n\nPlease highlight this problem so that the concerned authorities take prompt action and students can reach school safely.'
for t in models['Q5A'].split('\n\n'): p(t)
p('Yours faithfully,\nAarohi')
h('What to accept')
p('Accept other reasonable measures, such as designated parking away from the gate or monitoring at peak school hours. The letter should connect the obstruction with road exposure and student safety. Exact reproduction of this model is unnecessary.')

page(); h('Q5(B). Letter of complaint — Model and content allocation')
p('Content: 1 mark for the order/delivery details and both faults (½ for each substantially correct group); 1 mark for requesting both missing items and replacements by the required deadline.')
p('K.M. International School\n18, School Road\nJaipur\n16 September 2026\n\nThe Sales Manager\nSunrise Sports\n12, Station Road\nJaipur')
p('Subject: Short delivery and defective footballs — Order KM/SPORTS/27',True)
p('Sir/Madam,')
models['Q5B']='Our school ordered 20 footballs under Order No. KM/SPORTS/27 dated 2 September 2026. However, the consignment received on 10 September contained only 16 footballs. Four of the delivered balls have defective valves and cannot be used for practice.\n\nPlease supply the four missing footballs and replace the four defective ones before our sports trials on 20 September 2026. Kindly arrange collection of the defective balls and ensure that the replacements are checked before dispatch.\n\nThis shortage is affecting preparations for the trials. Please acknowledge this complaint and confirm the delivery schedule at the earliest. We expect prompt action so that all 20 footballs are available in usable condition.'
for t in models['Q5B'].split('\n\n'): p(t)
p('Yours faithfully,\nRiya\nSports Secretary')
h('Numerical check')
p('20 ordered − 16 received = 4 missing. Of the 16 received, 4 require replacement. The supplier must therefore deliver 4 missing footballs plus 4 replacements. These are two different problems; do not describe all eight as missing.')

page(); h('SECTION C — LITERATURE (40 MARKS)')
h('Q6. A Letter to God — Prose extract (4 marks)')
table(['Part','Answer and allocation','Marks'],[
    ['(i)','Lencho wrote the words (½), addressing God (½).',1],
    ['(ii)','Crooks.',1],
    ['(iii)','Lencho blames the postal staff for stealing the shortfall (1), although the postmaster and others had actually raised money to help him (1). The benefactors become the accused.',2]
],[.6,5.6,.65])
h('Q7. A Tiger in the Zoo — Poetry extract (6 marks)')
table(['Part','Answer and allocation','Marks'],[
    ['(i)','The tiger is confined in a small cage and deprived of freedom / adequate space for natural movement.',1],
    ['(ii)','B — Bright and distinct.',1],
    ['(iii)','A — Metaphor.',1],
    ['(iv)','Anger / rage / frustrated anger at captivity.',1],
    ['(v)','His movement is quiet and restrained, with soft, silent steps (1), while inwardly he is furious and frustrated by confinement (1). Explain the contrast rather than merely repeating the final phrase.',2]
],[.6,5.6,.65])
p('For Q7(iii), the question specifically asks about comparing the pads with velvet. “Oxymoron” may describe the separate phrase about restrained anger, but it does not answer this item.',size=10)
h('Q8 and Q9 — General marking approach')
p('Q8: assess two answers from Part A and two from Part B. Each answer earns 2 marks for the specified content points plus 1 for clear, coherent expression. Award ½ for a partly developed content point or expression that remains understandable despite noticeable errors. Maximum: 12.')
p('Q9: assess any four answers. Each earns 2 marks for the two stated content points. Language errors alone should not reduce these content marks unless they prevent understanding. Maximum: 8.')

page(); h('Q8. First Flight — Part A: Prose (any 2 × 3 = 6)')
p('Expected length: 40–50 words per answer. The following are marking points; students should develop them as connected answers. For every item, add up to 1 mark for expression to the 2 content marks shown.')
q8a=[
    ['(i) Mandela','1: His idea of freedom grows from childhood pleasures to adult personal liberty and then to freedom for all his people.\n1: This broader understanding makes him accept responsibility for collective liberation, despite personal sacrifice.'],
    ['(ii) His First Flight','1: His mother brings fish near him but stays beyond reach; hunger makes him dive from the ledge.\n1: His wings support him, revealing his ability to fly and replacing fear with confidence.'],
    ['(iii) Anne Frank','1: Family and ordinary friendships do not provide the close confidante with whom she can share her inner thoughts.\n1: She wants her diary, addressed as Kitty, to be a patient, trusted friend for honest self-expression.'],
    ['(iv) A Baker from Goa','1: Bread is integral to celebrations, keeping the baker socially important.\n½ + ½: Any two appropriate examples: sweet bread at marriages; sandwiches at engagements; cakes and bolinhas for Christmas or other festivals.'],
    ['(v) Coorg; Tea from Assam','1: Coorg’s forests, coffee estates, hills or wildlife offer attractive scenery/adventure; give a relevant detail.\n1: Rajvir’s curiosity/knowledge about tea makes him observe plantations, plucking or tea traditions with interest; give a relevant detail.']
]
table(['Part / work','Content allocation: 2 marks per answer'],q8a,[1.4,5.45])
p('In Q8(v), an answer dealing only with Coorg or only with Assam can earn at most 1 content mark. Do not demand a particular legend, date or example where another accurate detail answers the question.',size=10)

page(); h('Q8. First Flight — Part B: Poetry (any 2 × 3 = 6)')
p('Expected length: 40–50 words per answer. Award the 2 content marks below plus up to 1 mark for clear expression.')
q8b=[
    ['(vi) Dust of Snow','1: Snow shaken by a crow from a hemlock tree lifts the speaker’s regretful mood and improves the remaining day.\n1: The ordinary, unexpected incident shows that a small natural experience can bring emotional renewal.'],
    ['(vii) Fire and Ice','1: Fire represents uncontrolled desire/passion, while ice represents hatred or cold indifference.\n1: Both can destroy human relationships or the world; the contrast stresses the danger of extreme emotions.'],
    ['(viii) Wild Animals','1: Explain the comic contrast between cheerful identification advice and dangerous or fatal encounters.\n½ + ½: Two relevant examples, such as recognising a lion while it attacks, a leopard’s repeated leaps, or a bear’s crushing embrace.'],
    ['(ix) The Ball Poem','1: Money can buy another ball but cannot restore its memories or undo the loss; immediate replacement would bypass the lesson.\n1: The boy must learn to accept loss and develop responsibility/resilience.'],
    ['(x) Amanda!','1: Her imagined lives express a desire for freedom, solitude and escape from constant correction.\n½ + ½: Any two: a freely drifting mermaid; an orphan free to wander; Rapunzel enjoying undisturbed isolation.'],
    ['(xi) The Trees','1: Roots loosen, leaves press towards the glass, and cramped branches struggle; the effort makes release seem determined.\n1: Movement into the forest suggests freedom, restoration to a natural home and resistance to confinement.']
]
table(['Part / work','Content allocation: 2 marks per answer'],q8b,[1.4,5.45])
p('Accept a well-supported symbolic reading of The Trees, including liberation from social restrictions. Do not require one exclusive interpretation. For Amanda!, the fantasies are imagined escapes, not events that literally happen to her.',size=10)

page(); h('Q9. Footprints Without Feet (any 4 × 2 = 8 marks)')
p('Expected length: 30–40 words per answer. Award 1 mark for each content point; allow ½ for a substantially correct but incomplete point.')
q9=[
    ['(i) A Triumph of Surgery','1: Tricki needs restricted food and exercise because overfeeding and inactivity have made him unwell.\n1: Recovery through routine care shows that overindulgence, rather than a condition requiring surgery, caused his trouble.'],
    ['(ii) The Thief’s Story','1: Hari regrets betraying Anil’s trust and chooses to return.\n1: Education and the prospect of a respected future matter more than stolen money.'],
    ['(iii) The Midnight Visitor','1: Ausable presents the knock as police arriving to guard the important report, lending urgency to his invented story.\n1: Max fears arrest and tries to escape onto the supposed balcony.'],
    ['(iv) A Question of Trust','1: Eager to please the woman, Horace removes his gloves to offer his cigarette lighter.\n1: Opening the safe without replacing them leaves fingerprints that lead to his arrest.'],
    ['(v) Footprints Without Feet','1: Griffin sets fire to his landlord’s house in revenge and uses invisibility to escape detection.\n1: He exploits science for destructive, selfish ends without responsibility towards others.']
]
table(['Part / work','Content allocation'],q9,[1.65,5.2])
p('Accuracy note for Q9(iii): the actual visitor is Henry, the waiter, whom Ausable expects. The knock is not an arrival of real police. The question’s “unexpected” should be understood from Max’s perspective; do not penalise a student who clarifies this.',size=10)
link('Text checked: The Midnight Visitor — NCERT','https://ncert.nic.in/textbook/pdf/jefp103.pdf')
link('Text checked: A Question of Trust — NCERT','https://ncert.nic.in/textbook/pdf/jefp104.pdf')

page(); h('Q10. First Flight — Long answer (either A or B; 6 marks)')
p('Expected length: 120–150 words. Content: 4 marks, using the option-specific points below. Expression: 2 marks — 1 for organisation/coherence and 1 for language accuracy. Award ½ for partially developed content points. Reasoned alternative conclusions supported by the texts are acceptable.')
h('Option A — His First Flight and Black Aeroplane')
p('Content: seagull’s fear/help (1); pilot’s danger/help (1); comparison of confidence and risk (1); supported judgement (1).')
models['Q10A']='The young seagull and the pilot both face fear, but their situations differ. The seagull doubts his wings, although flying is a necessary stage of growth. His mother uses hunger to draw him beyond the ledge, and the flight reveals an ability he already possesses. The pilot, however, knowingly enters threatening storm clouds because he wants to reach home. When his instruments fail and fuel runs low, he depends on a mysterious aircraft to guide him towards safety. Both experiences show how assistance can restore confidence during danger. Yet the seagull’s risk helps him overcome an unfounded fear, whereas the pilot’s risk follows poor judgement. Courage therefore needs an honest assessment of ability and circumstances; surviving a dangerous choice does not automatically make that choice responsible.'
p(models['Q10A'])
p('Accept interpretations of the mysterious guide; the text leaves his identity unresolved.',size=10)
link('Text checked: Two Stories about Flying — NCERT','https://ncert.nic.in/textbook/pdf/jeff103.pdf')
h('Option B — A Letter to God')
p('Content: faith (1); practical compassion (1); mistaken accusation and irony (1); reasoned thematic conclusion (1).')
models['Q10B']='Lencho’s complete faith in God helps him face the destruction of his crop, but it also prevents him from recognising human kindness. He asks God for money, certain that divine help will arrive. Moved by this confidence, the postmaster decides to protect his belief. He contributes from his salary and collects money from employees and friends, although they cannot raise the full amount. Lencho receives seventy pesos instead of the hundred he requested. Rather than questioning God, he accuses the postal employees of stealing the remainder. The irony is that he condemns the very people who assisted him. The story presents faith as a source of hope and compassion as practical generosity. It also shows how unquestioning assumptions can produce injustice and misunderstanding, even when others act kindly.'
p(models['Q10B'])

page(); h('Q11. Footprints Without Feet — Long answer (either A or B; 4 marks)')
p('Expected length: 100–120 words. Content: 3 marks, using the points below. Expression: 1 mark — ½ for coherence and ½ for language accuracy. Award ½ for partially developed content points. The sample responses are not the only acceptable answers.')
h('Option A — The Thief’s Story')
p('Content: trust/kindness (1); education and return (1); significance of the ending (1).')
models['Q11A']='Anil offers Hari Singh food, shelter and the chance to learn, treating him with trust rather than suspicion. After stealing the money, Hari realises that leaving means losing both a caring relationship and the education that could make him respected. Temporary wealth seems less valuable than this future, so he returns the notes. The next morning, the damp money suggests that Anil knows what happened, yet he neither humiliates Hari nor withdraws his promise to teach him. This quiet forgiveness allows Hari to recover his self-respect. The ending suggests a real opportunity for reform: patient trust and education can encourage a change that punishment alone might not achieve.'
p(models['Q11A'])
link('Text checked: The Thief’s Story — NCERT','https://ncert.nic.in/textbook/pdf/jefp102.pdf')
h('Option B — A Question of Trust')
p('Content: deception (1); compromised caution/evidence (1); judgement about confidence and dishonesty (1).')
models['Q11B']='Horace’s careful planning cannot protect him from his mistaken belief that the young woman owns the house. Her confident manner and familiarity with the dog make her claim convincing. She uses his fear of imprisonment to persuade him to open the safe, pretending that she needs her jewels for a party. Eager to please her, he removes his gloves and later leaves fingerprints. She escapes with the jewels, while he is arrested. His confidence in his judgement proves misplaced. Although he considers himself respectable and steals to buy books, this motive does not make his conduct honest. His experience exposes both his gullibility and his attempt to justify wrongdoing.'
p(models['Q11B'])
link('Text checked: A Question of Trust — NCERT','https://ncert.nic.in/textbook/pdf/jefp104.pdf')
h('Final score check')
p('Q1–Q2: 20 | Q3–Q5: 20 | Q6–Q11: 40 | TOTAL: 80',True)
p('Literature distribution: First Flight prose 16; First Flight poetry 12; Footprints Without Feet 12. Count each internal choice only once.',size=10)

doc.core_properties.title='Class X English Half-Yearly — Answer Key and Marking Scheme'
doc.core_properties.subject='Teacher copy for the original 80-mark question paper; all internal choices covered'
doc.core_properties.author='K.M. International School'
doc.core_properties.comments='Prepared specifically for Class_10_English_Half_Yearly_Question_Paper.docx. Includes indicative keys, model writing, point allocations and valid alternative answers.'
doc.save(OUT)

limits={'Q4':(100,120),'Q5A':(100,120),'Q5B':(100,120),'Q10A':(120,150),'Q10B':(120,150),'Q11A':(100,120),'Q11B':(100,120)}
for k,(lo,hi) in limits.items():
    n=len(models[k].split()); print(f'{k}: {n} words')
    assert lo<=n<=hi,(k,n)
assert len(q1)==len(q2)==10 and len(q3)==12
assert len(q8a)==5 and len(q8b)==6 and len(q9)==5
assert sum([10,10,10,5,5,4,6,12,8,6,4])==80
assert (50-30)==20 and 30/200*100==15 and (60-30)/60*100==50
assert (4+6+6,6+6,8+4)==(16,12,12)
check=Document(OUT)
assert len(check.element.xpath('//w:br[@w:type="page"]'))==12
with zipfile.ZipFile(OUT) as z: assert z.testzip() is None
print(OUT.resolve())
print('Validated: all question parts and choices; 80-mark total; data calculations; model-answer word limits; Word package integrity.')
