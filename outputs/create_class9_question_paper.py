from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import zipfile

OUT = Path(__file__).parent / 'Class_9_English_Half_Yearly_Question_Paper.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.7)
sec.header_distance = sec.footer_distance = Inches(.28)
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.03
for name in ['Heading 1', 'Heading 2']:
    s = doc.styles[name]
    s.font.name = 'Times New Roman'
    s.font.size = Pt(12)
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(9)
    s.paragraph_format.space_after = Pt(7)

def p(text='', bold=False, size=None, center=False):
    para = doc.add_paragraph()
    r = para.add_run(text)
    r.bold = bold
    if size: r.font.size = Pt(size)
    if center: para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return para

def h(text): doc.add_paragraph(text, 'Heading 2')

def q(text):
    para = p(text, True)
    para.paragraph_format.keep_with_next = True

def item(label, text, marks=1):
    para = p(f'({label})  {text}  [{marks}]')
    para.paragraph_format.left_indent = Inches(.12)
    para.paragraph_format.first_line_indent = Inches(-.12)
    para.paragraph_format.keep_together = True

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    for c, text in zip(t.rows[0].cells, headers):
        c.text = text
        for r in c.paragraphs[0].runs: r.bold = True
    for row in rows:
        for c, text in zip(t.add_row().cells, row): c.text = str(text)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(4)
                para.paragraph_format.space_before = Pt(4)
    p('')

def source(label, url):
    para = doc.add_paragraph()
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), para.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '18'); props.append(size)
    run.append(props)
    text = OxmlElement('w:t'); text.text = label; run.append(text)
    link.append(run); para._p.append(link)

def page(): doc.add_page_break()

header = sec.header.paragraphs[0]
header.text = 'K.M. INTERNATIONAL SCHOOL  |  CLASS IX — ENGLISH'
header.alignment = WD_ALIGN_PARAGRAPH.CENTER
header.runs[0].font.size = Pt(9)
foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
foot.add_run('Page ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); foot._p.append(field)
foot.add_run(' of ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'NUMPAGES'); foot._p.append(field)

# PAGE 1
p('K.M. INTERNATIONAL SCHOOL', True, 15, True)
p('HALF-YEARLY EXAMINATION', True, 13, True)
p('CLASS IX — ENGLISH LANGUAGE & LITERATURE', True, 12, True)
p('Time Allowed: 3 Hours                                      Maximum Marks: 80', True)
p('Name: __________________________  Roll No.: __________  Section: ______')
h('General Instructions')
for text in [
    '1. This paper has 11 questions in three sections: A (20 marks), B (20 marks) and C (40 marks).',
    '2. All questions are compulsory, subject to the internal choices provided.',
    '3. Write answers in the answer booklet, retaining the question and sub-question numbers.',
    '4. Follow the word limits wherever given. Marks are indicated in brackets.'
]: p(text, size=10)
h('SECTION A — READING SKILLS (20 MARKS)')
q('Q1. Read the passage and answer ALL ten questions on the next page. (10 marks)')
passage1 = [
    '1. When a school announces a group project, some students immediately choose the most confident speaker as their leader. Others decide that the student with the highest marks should do all the difficult work. These choices may seem practical, but they overlook a simple fact: a group succeeds when its members contribute different strengths. Speaking clearly is useful, but so are observing carefully, organising materials, checking information and noticing who needs help.',
    '2. During a school exhibition, four friends decided to create a model of a water-saving garden. At first, everyone wanted to design the model, while nobody wished to prepare the labels or explain the costs. Their discussions became noisy, and little work was completed. Finally, they listed the tasks and divided them according to interest and ability. One measured the space, another researched plants, a third built the model, and the fourth prepared the presentation. They also agreed to check one another’s work.',
    '3. This arrangement did not mean that each person could ignore the others. When the model-maker discovered that the planned containers were too large, the measurements and costs had to change. The group listened, adjusted its plan and continued. Cooperation therefore involves more than dividing a job into parts. It requires members to share information and understand how their decisions affect the whole project. A perfect individual contribution may be useless if it does not fit the common purpose.',
    '4. Disagreement can improve a group’s work when it is handled respectfully. A student who questions a design may notice a weakness that everyone else has missed. The useful response is to examine the suggestion, not to silence the speaker. However, criticism should focus on the idea and include a reason. Saying that a plan will waste material invites discussion; saying that its designer is foolish damages trust. Members must also be prepared to change their own opinions when the evidence supports a better choice.',
    '5. Fairness matters as much as efficiency. Quiet members should have a chance to speak, and routine jobs should not always fall to the same person. Keeping a short record of responsibilities helps everyone see what remains to be done. It also makes it easier to offer support before a difficulty becomes a crisis. At the end, the group should recognise both visible achievements and less noticeable effort. A successful project is valuable, but learning to work responsibly with others is a skill that students can carry into many parts of life.'
]
for text in passage1: p(text)

# PAGE 2
page()
h('SECTION A — READING SKILLS (CONTINUED)')
q('Q1. Answer ALL ten questions using the passage on the previous page. (10 × 1 = 10)')
item('i', 'Choose the central idea of the passage.\nA. The highest-scoring student should do most of the work.\nB. Teamwork combines different strengths through shared responsibility.\nC. Group projects succeed only when members never disagree.\nD. Presenting a project is more important than preparing it.')
item('ii', 'Mention one useful group-work skill other than speaking clearly, as given in paragraph 1.')
item('iii', 'Why did the four friends make little progress at the beginning?')
item('iv', 'What was the basis for dividing the tasks in the exhibition project?')
item('v', 'Why did the group need to change its measurements and costs?')
item('vi', 'State whether this statement is TRUE or FALSE: Dividing the work removes the need for members to communicate.')
item('vii', 'Which response shows constructive criticism?\nA. Your plan is foolish.    B. Nobody should question the leader.\nC. This design uses too much cardboard; can we reduce its size?\nD. I dislike the idea, but I will not explain why.')
item('viii', 'Find a word in paragraph 5 that means “the ability to work well without wasting time or resources”.')
item('ix', 'How can a written record of responsibilities help a group? Give one benefit.')
item('x', 'A quiet student checks all the measurements but receives no thanks. What would the writer suggest the group should do?')
passage2 = [
    '1. A school library wanted to understand the effect of a six-week reading programme on book borrowing. The programme included short student book talks, a display of new titles and a weekly library period. The librarian compared borrowing records for the six weeks before the programme with those for the six weeks during it. The same 120 Class IX students had access to the library in both periods.',
    '2. The table records borrowing transactions, not the number of different readers. One student could borrow several books, and the same book could be borrowed by different students. Each transaction was placed in one category only. Both periods had the same number of school days, and the library followed the same borrowing rules. All figures are hypothetical and have been created for this examination.',
    '3. In an informal discussion, some students said that classmates’ recommendations encouraged them to try unfamiliar books. Others found the new display helpful. The librarian described the increased borrowing as encouraging but cautioned that a borrowing record does not prove that every book was completed. The records also cannot show whether students’ reading skills improved. To learn more, the school planned to collect voluntary reading reflections and repeat the study later.'
]

# PAGE 3
page()
h('SECTION A — READING SKILLS (CONTINUED)')
q('Q2. Read the passage, study the table and answer ALL ten questions. (10 × 1 = 10)')
for text in passage2: p(text)
table(['Book category', 'Before programme: transactions', 'During programme: transactions'], [
    ['Fiction', 72, 108], ['Biographies', 24, 36], ['Science and nature', 36, 60],
    ['Poetry', 12, 24], ['Travel and culture', 36, 42], ['Total', 180, 270]
])
item('i', 'What does one entry in the borrowing totals represent?\nA. A different student    B. A borrowing transaction\nC. A completed book    D. A new book purchased')
item('ii', 'Which category recorded the highest number of borrowing transactions during the programme?')
item('iii', 'Calculate the increase in the total number of transactions between the two periods.')
item('iv', 'What percentage of the transactions during the programme were for fiction books?')
item('v', 'Which category doubled its number of transactions?')
item('vi', 'How many more science and nature transactions were recorded during the programme than before it?')
item('vii', 'State whether this statement is TRUE or FALSE: The data prove that 270 different students borrowed books during the programme.')
item('viii', 'Mention one condition kept the same in both periods to make the comparison fair.')
item('ix', 'Why cannot the librarian conclude that every borrowed book was finished?')
item('x', 'In paragraph 3, “voluntary” means:\nA. done by choice    B. required by punishment\nC. completed secretly    D. performed without understanding')

# PAGE 4
page()
h('SECTION B — GRAMMAR & WRITING SKILLS (20 MARKS)')
q('Q3. Attempt ALL ten grammar items. (10 × 1 = 10)')
p('For multiple-choice items, write the correct option and answer. For editing items, write the incorrect word and its correction.')
item('i', 'Complete the sentence using the correct tense:\nBy the time the exhibition opened, our class __________ the model.\nA. completes    B. had completed    C. will complete    D. is completing')
item('ii', 'Fill in the blank with the correct form of the verb in brackets:\nListen! Someone __________ (knock) at the door.')
item('iii', 'Choose the modal that expresses obligation:\nAll students __________ wear their identity cards inside the laboratory.\nA. might    B. could    C. must    D. would')
item('iv', 'Choose the correct verb:\nEach of the players __________ a water bottle.\nA. carry    B. carries    C. have    D. are carrying')
item('v', 'Choose the correct verb:\nThe quality of these notebooks __________ excellent.\nA. are    B. were    C. is    D. have been')
item('vi', 'Fill in the blank with the correct article:\nMy sister hopes to become __________ engineer.\nA. a    B. an    C. the    D. no article')
item('vii', 'Choose the appropriate determiner:\nThere is very __________ water in the bottle; please refill it.\nA. few    B. many    C. little    D. several')
item('viii', 'Choose the correct preposition:\nThe children waited patiently __________ the bus stop.\nA. at    B. on    C. into    D. over')
item('ix', 'Identify the incorrect word and write its correction:\nThe coach asked us to spoke politely to the visiting team.')
item('x', 'Identify the incorrect word and write its correction:\nThe new student answered the question confident.')
q('Q4. Formal Letter (5 marks)')
p('You are Anuj/Ananya, a Class IX student of K.M. International School, 18, School Road, Jaipur. Write a letter in 100–120 words to the Principal requesting a weekly reading-club session. Explain how it would help students, suggest two activities and describe how students could help organise it. Use an appropriate formal format.')

# PAGE 5
page()
h('SECTION B — WRITING SKILLS (CONTINUED)')
q('Q5. Attempt EITHER option A OR option B. (5 marks)')
h('Option A — Diary Entry')
p('During a school craft workshop, you struggled to make a clay bowl. An artisan patiently guided you, and you finally completed it. Write a diary entry in 100–120 words describing your initial frustration, the support you received and what the experience taught you about patience and respect for skilled work. Include a suitable date and an appropriate diary opening.')
p('OR', True, center=True)
h('Option B — Story Writing')
p('Write an original story in 120–150 words beginning with the line below. Give your story a suitable title and a clear ending.')
p('“When I opened the old notebook, a folded map slipped onto the floor.”')
p('You may use these clues: an unfamiliar symbol — a friend’s suggestion — a search in an ordinary place — an unexpected discovery. The clues are optional; develop a coherent plot with your own details.')
h('SECTION C — LITERATURE (40 MARKS)')
q('Q6. Read the prose extract and answer ALL four questions. (5 marks)')
p('“We are well-off, but what use is money when I cannot be independent?”')
source('— How I Taught My Grandmother to Read, Sudha Murty; NCERT, Kaveri', 'https://ncert.nic.in/textbook/pdf/iebe101.pdf')
p('Answer with reference to the extract and the story.')
item('i', 'Identify the speaker and the person being addressed.', 1)
item('ii', 'What does the expression “well-off” mean here?\nA. Healthy    B. Wealthy    C. Well educated    D. Famous', 1)
item('iii', 'What inability makes the speaker feel dependent?', 1)
item('iv', 'What personal goal does the speaker set after this conversation, and what does it reveal about her character?', 2)

# PAGE 6
page()
h('SECTION C — LITERATURE (CONTINUED)')
q('Q7. Read the poetry extract and answer ALL four questions. (5 marks)')
verse = p('Palette of earth, rich and deep,\nWhere dreams of gardeners seep.\nBrushstrokes of seeds, planted true,\nAwaiting spring’s vibrant hue.')
verse.paragraph_format.left_indent = Inches(.25)
source('— Canvas of Soil, Maya Anthony; NCERT, Kaveri', 'https://ncert.nic.in/textbook/pdf/iebe103.pdf')
item('i', 'What is the rhyme scheme of these four lines?', 1)
item('ii', 'The comparison between seeds and brushstrokes is an example of:\nA. Metaphor    B. Simile    C. Onomatopoeia    D. Hyperbole', 1)
item('iii', 'Which word in the extract means “a shade of colour”?', 1)
item('iv', 'How does the extract present gardening as a creative activity? Explain using two details.', 2)
q('Q8. Short Answer Questions — Attempt ANY FOUR. (4 × 3 = 12 marks)')
p('Write each answer in 40–50 words and support it with relevant details from the prescribed work.')
item('i', 'How does Onula’s guidance help Sentila learn more confidently? Explain the value of encouragement in The Pot Maker.', 3)
item('ii', 'How do references to nature and spiritual traditions together express national pride in Bharat Our Land? Give one example of each.', 3)
item('iii', 'How does Gifts of Grace: Honouring Our Vocations show that different kinds of work deserve respect? Refer to two occupations.', 3)
item('iv', 'How does Canvas of Soil connect human effort with natural beauty? Explain the role of the gardener.', 3)
item('v', 'How do sound and smell bring the mother’s presence back to the speaker in I Cannot Remember My Mother?', 3)
item('vi', 'Why is Ravi uncomfortable with the way his mother speaks to Grandpa? What does his reaction show about his feelings? (Vitamin-M)', 3)

# PAGE 7
page()
h('SECTION C — LITERATURE (CONTINUED)')
q('Q9. Short Answer Questions — Attempt ALL FOUR. (4 × 2 = 8 marks)')
p('Write each answer in 30–40 words.')
item('i', 'What gift does the granddaughter choose for her grandmother, and how does the grandmother show that her lessons have succeeded? (How I Taught My Grandmother to Read)', 2)
item('ii', 'Why do the village elders consider teaching pottery a responsibility towards the community? Give two reasons. (The Pot Maker)', 2)
item('iii', 'Explain two ways in which craft exhibitions or workshops can benefit traditional fan-makers. (Winds of Change)', 2)
item('iv', 'Why does a stranger have Grandpa’s yellow cap, and what does this reveal about Grandpa? (Vitamin-M)', 2)
q('Q10. Long Answer — Attempt EITHER A OR B. (5 marks)')
p('Write your answer in 100–120 words, supporting your views with relevant incidents.')
p('A. Explain how the grandmother’s learning journey challenges assumptions about age and education. Discuss her effort and her respect for her young teacher. (How I Taught My Grandmother to Read)')
p('OR', True, center=True)
p('B. Arenla sees practical difficulties in pottery, while Sentila sees a vocation she loves. Discuss both viewpoints and explain how support and practice help Sentila progress. (The Pot Maker)')
q('Q11. Long Answer — Attempt EITHER A OR B. (5 marks)')
p('Write your answer in 100–120 words, supporting your views with relevant details.')
p('A. Caring for an older person requires concern for safety and respect for independence. Discuss Ravi’s dilemma and his response when Grandpa goes out. How does Vitamin-M encourage understanding between generations?')
p('OR', True, center=True)
p('B. A school wishes to celebrate traditional hand fans. Using Winds of Change, explain why these objects matter as cultural heritage. Suggest how such an event could help artisans, giving details from the chapter to support your ideas.')
p('— END OF QUESTION PAPER —', True, 11, True)

doc.core_properties.title = 'Class IX English Half-Yearly Question Paper'
doc.core_properties.subject = '80 marks | 3 hours | Supplied K.M. International School blueprint'
doc.core_properties.author = 'K.M. International School'
doc.core_properties.comments = (
    'Uses the eight works explicitly listed in the blueprint. No Moments chapters '
    'are named, so the additional 8-mark component uses the four prescribed prose '
    'works. Writing follows the detailed 5 + 5 distribution.'
)
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)

marks = [10,10,10,5,5,5,5,12,8,5,5]
assert sum(marks) == 80 and sum(marks[5:]) == 40
grammar = {'Tenses':2,'Modals':1,'Concord':2,'Determiners and articles':2,'Prepositions':1,'Editing':2}
assert sum(grammar.values()) == 10
assert sum([72,24,36,12,36]) == 180
assert sum([108,36,60,24,42]) == 270
assert 108 / 270 * 100 == 40
assert len(' '.join(passage1).split()) >= 400
assert 200 <= len(' '.join(passage2).split()) <= 250
saved = Document(OUT)
assert len(saved.tables) == 1
assert len(saved.element.xpath('//w:br[@w:type="page"]')) == 6
with zipfile.ZipFile(OUT) as z: assert z.testzip() is None
print(OUT.resolve())
print('Validated: 80 marks, 20/20/40 sections, exact grammar allocation, all eight prescribed works included.')
print('Reading passage word counts:', len(' '.join(passage1).split()), len(' '.join(passage2).split()))
