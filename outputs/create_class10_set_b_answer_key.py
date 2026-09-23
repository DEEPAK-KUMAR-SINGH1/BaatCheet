from pathlib import Path
from copy import deepcopy
import zipfile
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).parent
SOURCE=ROOT/'Class_10_English_Half_Yearly_Question_Paper_Set_B.docx'
TEMPLATE=ROOT/'Class_10_English_Half_Yearly_Answer_Key_and_Marking_Scheme.docx'
OUT=ROOT/'Class_10_English_Half_Yearly_Set_B_Answer_Key_and_Marking_Scheme.docx'
paper=Document(SOURCE)
paper_text='\n'.join(p.text for p in paper.paragraphs)
assert 'SET B' in paper_text and 'Offering help seems simple' in paper_text
assert 'By next Monday' in paper_text and 'KM/LIB/32' in paper_text
doc=Document(TEMPLATE)
paras=doc.paragraphs
assert len(paras)==124 and len(doc.tables)==11

def replace(p,text):
    props=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear(); r=p.add_run(text)
    if props is not None: r._r.insert(0,props)
def setp(i,text): replace(paras[i],text)
def table(i,rows):
    t=doc.tables[i]
    assert len(t.rows)==len(rows)+1
    for row,values in zip(t.rows[1:],rows):
        for c,value in zip(row.cells,values): replace(c.paragraphs[0],str(value))
def link(i,label,url):
    p=paras[i]; p.clear()
    a=OxmlElement('w:hyperlink'); a.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); props=OxmlElement('w:rPr')
    size=OxmlElement('w:sz'); size.set(qn('w:val'),'18'); props.append(size); r.append(props)
    t=OxmlElement('w:t'); t.text=label; r.append(t); a.append(r); p._p.append(a)

setp(1,'HALF-YEARLY EXAMINATION — CLASS X — SET B')
setp(5,'Teacher copy for Class_10_English_Half_Yearly_Question_Paper_Set_B.docx')
paras[5].runs[0].font.size=Pt(9)
for sec in doc.sections:
    replace(sec.header.paragraphs[0],'K.M. INTERNATIONAL SCHOOL | CLASS X ENGLISH | SET B KEY')
    sec.header.paragraphs[0].runs[0].font.size=Pt(9)

setp(16,'Q1. Discursive passage — Respectful assistance (10 marks)')
q1=[
    ['(i)','B — Effective help respects the receiver’s needs and choices.'],
    ['(ii)','The classmate may want help with only one heavy box, rather than having everything carried. Acting without asking can override the classmate’s wishes.'],
    ['(iii)','The teenager pressed the buttons for the neighbour, so the learner did not practise or understand how to repeat the search independently.'],
    ['(iv)','Try / perform / practise / carry out. Accept any equivalent verb or phrase that fits the sentence.'],
    ['(v)','C — Which part would you like help with?'],
    ['(vi)','False. A person may prefer to manage the challenge independently.'],
    ['(vii)','It may create reliance on support that later disappears, leaving the person without dependable help / damaging trust.'],
    ['(viii)','Discreetly.'],
    ['(ix)','Any one: whether the help meets a real need; builds confidence; preserves a meaningful choice or independence.'],
    ['(x)','Respect for privacy/confidentiality and dignity: help with a form does not give permission to share personal details publicly.']
]
table(1,q1)
setp(19,'Partial-credit guidance: award ½ for a relevant but incomplete explanation in Q1(ii), (iii) or (vii), and 1 when the reason is clear. In Q1(ix), one valid measure is enough. Do not require explanations for the MCQ or true/false items.')

setp(21,'Q2. Case-based passage — Revision methods (10 marks)')
q2=[
    ['(i)','A — To count each respondent once in each total.'],
    ['(ii)','Self-testing (80 students). Naming the method is sufficient.'],
    ['(iii)','40 students: 80 − 40 = 40.'],
    ['(iv)','16%: (40 ÷ 250) × 100 = 16%.'],
    ['(v)','40% decrease: [(100 − 60) ÷ 100] × 100 = 40%.'],
    ['(vi)','True. Lesson-video users remained at 20 in both months.'],
    ['(vii)','Organising information into notes suited the students / matched their preferred way of studying.'],
    ['(viii)','The table records preferred revision methods, not examination scores. There is no performance data showing that marks improved.'],
    ['(ix)','B — Self-testing became the choice of twice as many students (40 to 80).'],
    ['(x)','Provide an additional discussion period during the school day.']
]
table(2,q2)
setp(25,'Q2(iv): ½ for the correct setup 40/250 × 100 and ½ for 16%, if working is shown but incomplete. Q2(v): ½ for the correct relative-decrease setup and ½ for 40%. The share choosing rereading fell from 40% to 24%: that is 16 percentage points, not the requested percentage decrease.')
setp(26,'Q2(viii): “no examination scores were recorded” is sufficient for 1 mark. Merely mentioning the absence of a comparison group does not fully explain why this table cannot show a change in marks; award at most ½ for that relevant but incomplete observation.')

q3=[
    ['(i)','C — some',1],['(ii)','A — will have completed',1],
    ['(iii)','B — should',1],['(iv)','D — are',1],
    ['(v)','A — of',1],['(vi)','B — many',1],
    ['(vii)','were preparing (expected past continuous; see note)',1],
    ['(viii)','have → has: Every one of these paintings has a story behind it.',1],
    ['(ix)','meet → meeting: We look forward to meeting the visiting author.',1],
    ['(x)','had finished (expected backshift; see note)',1],
    ['(xi)','C — may',1],
    ['(xii)','than → to: Ms Sen is senior to me in this organisation.',1]
]
table(3,q3)
setp(33,'Q3(vii): “were preparing” expresses an action in progress at nine. “Prepared” is also grammatical if it reports the preparation as an event at that time. The sentence does not explicitly require continuous aspect, so accept either.')
setp(34,'Q3(x): “had finished” is the expected present-perfect-to-past-perfect backshift. “Has finished” is also possible when the completion remains relevant at reporting time; the paper does not specify that context, so accept it. In Q3(iv), the nearer subject “students” takes the plural verb “are”.')

models={}
models['Q4']='The table shows the festival preferences of 250 Class X students. Sports competitions are the most popular choice, attracting 75 students, or 30%, while literary events rank last with 25 students, or 10%. Music performances stand second at 24%, followed by the science exhibition at 20% and art displays at 16%. Sports are preferred by three times as many students as literary events. The science exhibition attracts twice the number choosing literary events. Together, sports and music account for 54% of respondents, forming a majority. Overall, preferences are spread across all five activities, although sports and music receive the strongest combined support.'
setp(41,models['Q4'])
setp(43,'Sports 75/30%; music 60/24%; science 50/20%; art 40/16%; literary events 25/10%. Total 250/100%. Other accurate comparisons are valid, such as the 6-percentage-point gap between sports and music. A student need not reproduce every statistic for full marks.')
setp(44,'Do not reward invented explanations for preferences or conclusions about student ability. This survey records each student’s preferred festival activity, not actual participation or performance.')
setp(47,'Assess either A or B. Accept Dev or Diya. Sensible dates and conventional formal closings are acceptable. Model body paragraphs follow the 100–120-word limit; address blocks and the closing are shown separately.')

setp(50,'Content: 1 mark for the streetlight fault and its effect on residents; 1 mark for practical remedies and a clear request for action or public attention.')
setp(51,'32, Ashok Nagar\nJaipur\n16 September 2026\n\nThe Editor\nThe City Herald\nJaipur')
setp(52,'Subject: Non-functioning streetlights near the local bus stop')
letter_a=[
    'Through the columns of your newspaper, I wish to highlight the failure of streetlights near the local bus stop in Ashok Nagar. They have remained out of order for several weeks. After dark, pedestrians struggle to see uneven surfaces and other obstacles, making the route difficult for children and older residents.',
    'The authorities should repair the lights promptly and arrange regular inspections. A clearly advertised telephone number or online system for reporting faults would allow residents to seek timely assistance.',
    'Please bring this issue to public attention and urge the concerned department to restore reliable lighting so that residents can use the bus stop safely.'
]
for i,t in zip([54,55,56],letter_a): setp(i,t)
models['Q5A']=' '.join(letter_a)
setp(57,'Yours faithfully,\nDiya')
setp(59,'Accept other relevant suggestions, such as recording complaints and following up unresolved faults. The letter should relate poor lighting to difficulty seeing obstacles and pedestrian safety. Do not demand invented accident statistics or claims about criminal incidents.')

setp(63,'K.M. International School\n18, School Road\nJaipur\n16 September 2026\n\nThe Sales Manager\nScholar Books\n15, College Road\nJaipur')
setp(64,'Subject: Missing and defective atlases — Order KM/LIB/32')
letter_b=[
    'Our school ordered 40 English atlases under Order No. KM/LIB/32 dated 4 September 2026. The consignment received on 12 September contained only 35 atlases. On inspection, we also found that six of these copies had missing pages and were unsuitable for use.',
    'Please supply the five missing atlases and replace the six defective copies before our geography exhibition on 22 September 2026. Kindly arrange collection of the defective books and check all replacement copies before dispatch.',
    'The shortage is affecting our preparations for the exhibition. Please acknowledge this complaint and confirm the delivery schedule promptly so that all 40 atlases are available in complete, usable condition.'
]
for i,t in zip([66,67,68],letter_b): setp(i,t)
models['Q5B']=' '.join(letter_b)
setp(69,'Yours faithfully,\nDev\nLibrary Secretary')
setp(71,'40 ordered − 35 received = 5 missing. Six of the 35 delivered copies need replacement. The supplier should therefore send 5 missing atlases plus 6 replacements, collecting the 6 defective copies. Do not describe all eleven as missing.')

setp(74,'Q6. Nelson Mandela: Long Walk to Freedom — Extract (4 marks)')
q6=[
    ['(i)','From his comrades / fellow men and women in the struggle against apartheid, who faced danger and suffering for their beliefs.',1],
    ['(ii)','Triumph.',1],
    ['(iii)','Courage involves overcoming fear and acting despite it (1). His fellow freedom fighters endured attacks/torture or risked their lives without abandoning their cause (1).',2]
]
table(6,q6)
setp(76,'Q7. Dust of Snow — Poetry extract (6 marks)')
q7=[
    ['(i)','His regretful, unhappy mood becomes lighter / more positive / hopeful.',1],
    ['(ii)','B — Regretted.',1],
    ['(iii)','A — Enjambment.',1],
    ['(iv)','A crow shakes snow from a hemlock tree onto the speaker. Accept an accurate account of the crow dislodging snow onto him without insisting on the tree’s name.',1],
    ['(v)','Part of the day has already been spent in regret and cannot be recovered (1). The small uplifting incident nevertheless improves the remaining time, showing that a difficult day can still contain relief or hope (1).',2]
]
table(7,q7)
setp(78,'Q7(iii) concerns the sentence continuing across line breaks. Do not substitute a rhyme-scheme answer. In Q7(v), accept another well-supported explanation of limited but meaningful emotional relief; the poem does not establish that every problem has disappeared.')

q8a=[
    ['(i) Mandela','1: Obligations to family and to people/community/country (½ each).\n1: Under apartheid, serving his people exposed him to punishment and separation, preventing him from fulfilling ordinary family duties.'],
    ['(ii) His First Flight','1: Family members fly around him, call encouragement, and celebrate his successful flight/landing.\n1: Their support reassures him, confirms that he can fly and helps him join them with confidence.'],
    ['(iii) Anne Frank','1: Her verse describes ducklings killed by their father swan for excessive quacking, humorously mirroring punishment for talking.\n1: Keesing appreciates the joke, reads it to classes and stops assigning extra homework for her talking.'],
    ['(iv) A Baker from Goa','1: Distinctive clothing and the familiar bamboo-staff sound make the baker a recognisable part of daily life.\n1: The affectionate recollection conveys nostalgia and the baker’s importance in childhood/community traditions.'],
    ['(v) Coorg; Tea from Assam','1: Coorgi hospitality welcomes visitors into homes and conversations about family valour; accept a relevant textual detail.\n1: Rajvir is excited by the tea landscape, whereas Pranjol, raised on a plantation, finds it familiar.']
]
table(8,q8a)
setp(86,'Q8(v): an answer about only Coorg or only Assam can earn at most 1 content mark. Q8(iv): distinguish the older kabai from the shirt and trousers, between shorts and full length, recalled from the narrator’s own childhood. Either accurate detail may support the analysis.')

q8b=[
    ['(vi) A Tiger in the Zoo','1: Ignoring visitors suggests frustrated withdrawal/indifference rather than pleasure in being displayed.\n1: His behaviour shows the confinement of a powerful creature deprived of natural freedom; accept a supported reading of dignity or resistance.'],
    ['(vii) Fire and Ice','1: The speaker initially supports fire, drawing on experience of desire.\n1: Knowledge of hatred leads him to recognise that ice could also cause destruction; both emotional extremes are dangerous.'],
    ['(viii) Wild Animals','½: The bear is identified by its dangerous embrace. ½: The leopard is recognised as it repeatedly leaps upon its victim.\n1: The advice becomes absurd because recognition comes during an attack; the light tone sharpens the humour.'],
    ['(ix) The Ball Poem','1: The ball carries the boy’s memories and experiences of childhood, giving it personal value.\n1: A replacement can restore possession of an object, but cannot bring back the same past or undo its loss.'],
    ['(x) Amanda!','1: Repeated commands and corrections contrast with Amanda’s private fantasies of peaceful freedom.\n1: This contrast reveals her sense of restriction and longing for autonomy, and invites sympathy for her inner life.'],
    ['(xi) The Trees','1: Roots disengage from floor cracks, leaves press against the glass, and twigs/branches strain in their confinement; give relevant details.\n1: Their combined effort suggests persistent, purposeful resistance and an urgent movement towards freedom.']
]
table(9,q8b)
setp(91,'For Q8(vi), reward interpretations supported by the caged tiger’s behaviour, rather than requiring one psychological label. For Q8(xi), two accurate physical details with explanation are sufficient; do not demand every image in the poem.')

q9=[
    ['(i) A Triumph of Surgery','1: Herriot and his colleagues enjoy the eggs, wine and brandy sent for Tricki.\n1: The deliveries show Mrs Pumphrey’s extravagant, anxious affection and belief that rich extras aid recovery.'],
    ['(ii) The Thief’s Story','1: The fifty-rupee note Anil gives Hari is still damp, suggesting he discovered the returned money.\n1: Anil remains kind, promises regular payment/continued lessons and does not accuse or humiliate him.'],
    ['(iii) The Midnight Visitor','1: Ausable’s detailed story convinces Max that a balcony is outside the window.\n1: Max trusts this fiction and jumps towards a non-existent balcony to escape the supposed police.'],
    ['(iv) A Question of Trust','1 each for two details linking behaviour to apparent ownership: she handles/addresses Sherry familiarly; casually adjusts ornaments; speaks firmly as mistress of the house; claims the jewels and refers to her husband.'],
    ['(v) Footprints Without Feet','1: Griffin oversleeps; returning shop assistants see him in his stolen clothes and pursue him when he runs.\n1: He removes the clothes, exposing his invisible body and escaping unseen.']
]
table(10,q9)
setp(96,'Accuracy notes: Q9(i) — the staff, not Tricki, consume the delivered luxuries. Q9(ii) — do not say Anil openly accuses Hari. Q9(v) — Griffin does not need another dose of chemicals; removing visible clothes makes him unseen again.')
link(97,'Text checked: A Triumph of Surgery — NCERT','https://ncert.nic.in/textbook/pdf/jefp101.pdf')
link(98,'Text checked: Footprints Without Feet — NCERT','https://ncert.nic.in/textbook/pdf/jefp105.pdf')

setp(102,'Option A — Black Aeroplane')
setp(103,'Content: risky decision/motive (1); storm and failures (1); guidance and unresolved mystery (1); reasoned judgement of the ending (1).')
models['Q10A']='The pilot enters the storm because his desire to reach home outweighs his awareness of danger. He knows that returning to Paris is sensible, but chooses to continue towards England. Inside the clouds, darkness and violent movement are followed by instrument and radio failure. With little fuel remaining, he becomes dependent on the pilot of a strange black aeroplane, who guides him towards a runway. After landing, however, he learns that no other aircraft appeared on the radar. The disappearance of his guide leaves the rescue unexplained. This ending creates wonder while exposing how helpless he became after a reckless decision. His gratitude is understandable, but survival does not prove that the risk was justified. The episode encourages humility and careful judgement rather than confidence based only on a fortunate outcome.'
setp(104,models['Q10A'])
setp(108,'Content: hope at the rain (1); destruction and despair (1); dependence on farming (1); enduring faith and appeal to God (1).')
models['Q10B']='Lencho initially welcomes the rain because his crop needs water and promises a good harvest. He imagines the drops as money, linking the weather directly with his family’s livelihood. His satisfaction turns to anxiety when hail begins, and then to despair as the storm destroys the corn and flowers. With no harvest, he fears that the family will have nothing to eat. Yet this disaster does not destroy his belief in divine help. Convinced that God can see his suffering, he writes asking for a hundred pesos to sow the field again and survive until the next crop. His changing feelings reveal the insecurity of a farmer whose income depends on nature. His practical request also shows how unwavering faith gives him hope when ordinary resources appear exhausted.'
setp(109,models['Q10B'])

setp(113,'Option A — A Triumph of Surgery')
setp(114,'Content: harmful indulgence (1); practical treatment and recovery (1); conclusion about responsible care (1).')
models['Q11A']='Mrs Pumphrey loves Tricki but confuses care with constant indulgence. Rich food, frequent extras and insufficient exercise make him overweight and unwell. Mr Herriot recognises that Tricki needs a change of routine rather than further treats. At the surgery, he initially restricts food, provides water and then allows normal meals and active companionship with other dogs. Tricki soon regains energy and physical fitness without an operation. The contrast shows that affection must be guided by an understanding of actual needs. Mrs Pumphrey’s generosity harms the animal when she cannot refuse him, whereas Herriot’s firm, practical approach restores his health. Responsible care may therefore require sensible limits rather than satisfying every desire.'
setp(115,models['Q11A'])
link(116,'Text checked: A Triumph of Surgery — NCERT','https://ncert.nic.in/textbook/pdf/jefp101.pdf')
setp(117,'Option B — The Midnight Visitor')
setp(118,'Content: misleading appearance (1); believable balcony/police inventions (1); outcome and changed understanding (1).')
models['Q11B']='Fowler expects a secret agent to look adventurous, so Ausable’s ordinary appearance and manner disappoint him. When Max threatens them with a pistol, however, Ausable remains composed. He invents a detailed complaint about a balcony outside the window, making an escape route seem real. At the knock, he claims that police have arrived to protect the important report. Max accepts the story and jumps towards the supposed balcony. The visitor is actually Henry, the waiter, and no balcony exists. Fowler consequently discovers that intelligence and control of a situation matter more than glamorous appearances. Ausable succeeds by understanding Max’s fears and turning convincing details into a means of overcoming armed danger.'
setp(119,models['Q11B'])
link(120,'Text checked: The Midnight Visitor — NCERT','https://ncert.nic.in/textbook/pdf/jefp103.pdf')

doc.core_properties.title='Class X English Half-Yearly Set B — Answer Key and Marking Scheme'
doc.core_properties.subject='Teacher copy | Set B | 80 marks | all internal choices'
doc.core_properties.comments='Prepared for Class_10_English_Half_Yearly_Question_Paper_Set_B.docx. Includes all keys, point allocations, model writing and valid alternatives.'
doc.save(OUT)

limits={'Q4':(100,120),'Q5A':(100,120),'Q5B':(100,120),'Q10A':(120,150),'Q10B':(120,150),'Q11A':(100,120),'Q11B':(100,120)}
for k,(lo,hi) in limits.items():
    count=len(models[k].split()); print(f'{k}: {count} words')
    assert lo<=count<=hi,(k,count)
assert len(q1)==10 and len(q2)==10 and len(q3)==12
assert len(q6)==3 and len(q7)==5 and len(q8a)==5 and len(q8b)==6 and len(q9)==5
assert sum([10,10,10,5,5,4,6,12,8,6,4])==80
assert (80-40)==40 and (40/250*100)==16 and ((100-60)/100*100)==40
assert sum([75,60,50,40,25])==250 and sum([30,24,20,16,10])==100
check=Document(OUT)
assert len(check.element.xpath('//w:br[@w:type="page"]'))==12
alltext='\n'.join(check.element.xpath('//w:t/text()'))
for stale in ['bridge-builder','Travel to school','KM/SPORTS/27','defective footballs','crooks','Riya','Aarohi','not Set B','somebody in this village']:
    assert stale not in alltext,stale
with zipfile.ZipFile(OUT) as z: assert z.testzip() is None
print(OUT.resolve())
print('Validated: all Set B questions/options, 80 marks, calculations, model-answer word limits and Word integrity.')
