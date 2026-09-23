from pathlib import Path
from copy import deepcopy
import re
import zipfile
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

folder = Path(__file__).parent
src = folder / 'Class_10_English_Half_Yearly_Question_Paper.docx'
out = folder / 'Class_10_English_Half_Yearly_Question_Paper_Set_B.docx'
doc = Document(src)
paras = doc.paragraphs
assert len(paras) == 134, 'Inspect the changed source before editing.'
assert paras[53].text.startswith('Q3.') and paras[116].text.startswith('Q9.')

def setp(index, text):
    p = paras[index]
    props = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear()
    r = p.add_run(text)
    if props is not None: r._r.insert(0, props)

def source(index, label, url):
    p = paras[index]; p.clear()
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '18'); props.append(size)
    run.append(props)
    text = OxmlElement('w:t'); text.text = label; run.append(text)
    link.append(run); p._p.append(link)

romans = ['i','ii','iii','iv','v','vi','vii','viii','ix','x','xi','xii']
def questions(start, items, offset=0):
    for j, value in enumerate(items):
        text, marks = value if isinstance(value, tuple) else (value,1)
        setp(start+j,f'({romans[j+offset]})  {text}  [{marks}]')

def fill_table(index, rows):
    for row, values in zip(doc.tables[index].rows,rows):
        for cell, text in zip(row.cells,values):
            p = cell.paragraphs[0]
            props = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
            p.clear(); r = p.add_run(str(text))
            if props is not None: r._r.insert(0,props)

setp(1,'HALF-YEARLY EXAMINATION — SET B')
for sec in doc.sections:
    p=sec.header.paragraphs[0]
    p.text='K.M. INTERNATIONAL SCHOOL  |  CLASS X — ENGLISH (184)  |  SET B'
    p.runs[0].font.size=Pt(9)

passage1 = [
    '1. Offering help seems simple: notice a difficulty and remove it. Yet useful assistance begins with understanding what another person actually needs. A student who carries every material for a classmate may believe that the task has become easier. The classmate, however, may have wanted only help with one heavy box. When we act without asking, kindness can unintentionally become control. Respectful help makes room for the other person’s wishes, abilities and choices.',
    '2. At a community workshop, a teenager helped an older neighbour learn to use a digital catalogue. Whenever the neighbour hesitated, the teenager quickly pressed the correct buttons. The book appeared on the screen, but the neighbour remained uncertain about how to find it again. On the next attempt, the teenager explained one step at a time and waited while the neighbour tried. The search took longer, yet the learner could repeat it independently. Finishing a task and helping someone learn are not always the same achievement.',
    '3. Good support also requires careful listening. Sometimes a person wants practical assistance; sometimes the person needs an explanation or simply time to think aloud. Asking a clear question can prevent a misunderstanding. Listening means allowing an answer to shape our response, even when it differs from our original plan. It also means accepting that an offer may be declined. A refusal need not be treated as ingratitude; people may prefer to manage a challenge in their own way.',
    '4. There are limits to what any helper can do. Promising more than we can provide may create dependence on support that later disappears. It is better to explain honestly what time, knowledge or resources are available. Privacy matters too. Someone who asks for help with a form has not automatically agreed to have personal details shared with others. Trust grows when assistance is reliable and information is handled discreetly. Public praise for the helper should never become more important than the dignity of the person receiving support.',
    '5. The best measure of assistance is therefore not how visible or impressive it appears. We should ask whether it meets a real need and leaves the other person with confidence and a meaningful choice. Sometimes this involves teaching a skill; at other times it involves sharing a burden that cannot easily be removed. Thoughtful help combines generosity with patience. It recognises that people are partners in finding solutions, even when they need different amounts or kinds of support.'
]
for i,t in enumerate(passage1,12): setp(i,t)
questions(21,[
    'Which statement best expresses the central idea?\nA. Help is useful only when it saves time.\nB. Effective help respects the receiver’s needs and choices.\nC. People should always accept assistance.\nD. Visible acts of kindness matter more than quiet support.',
    'Why might carrying every material for a classmate fail to meet the classmate’s wishes?',
    'Why did the neighbour remain unsure after the teenager’s first demonstration?',
    'Complete the statement: In the second attempt, the teenager allowed the neighbour to __________ the steps rather than merely watch.',
    'Which question best reflects the approach recommended in paragraph 3?\nA. Why can’t you do this yourself?\nB. Why have you rejected my excellent plan?\nC. Which part would you like help with?\nD. Will you tell everyone that I helped you?',
    'State whether the following is TRUE or FALSE: Declining help necessarily shows a lack of gratitude.',
    'Why should a helper avoid promising more support than can actually be provided?',
    'Find a word in paragraph 4 that means “carefully, without revealing private information”.',
    'According to paragraph 5, what is one useful way to judge whether assistance has been effective?',
    'A student shares a photograph of a neighbour’s completed personal form to display a helpful act. Identify the principle in paragraph 4 that the student has ignored.'
])

passage2 = [
    '1. A school surveyed 250 Class X students about their main method of revision in July and again in August. Each student selected the single method used most often during the previous week. The same students answered both surveys. The table presents hypothetical data prepared for this examination; it does not describe the results of an actual study.',
    '2. Between the surveys, teachers demonstrated how to create short self-tests and encouraged students to discuss difficult ideas with classmates. Students were free to choose their own revision methods. The school also provided sample questions and a room for small discussion groups during a supervised after-school session.',
    '3. Some students reported that self-testing helped them notice gaps in understanding. Others preferred preparing notes because organising information suited them. A few students who wanted group discussions could not stay after school because of transport arrangements. The school therefore considered offering an additional discussion period during the school day.',
    '4. The table shows reported preferences, not examination scores or time spent studying. An increase in the use of a method does not prove that the method improved marks. Since there was no comparison group, the school could not conclude that the demonstrations alone caused the changes. It planned to collect further feedback before deciding how to continue the support.'
]
for i,t in zip([32,33,34,40],passage2): setp(i,t)
fill_table(0,[
    ['Main revision method','July: students','August: students'],
    ['Rereading textbooks',100,60],['Preparing notes',60,50],['Self-testing',40,80],
    ['Group discussion',30,40],['Watching lesson videos',20,20],['Total',250,250]
])
questions(41,[
    'Why did each student choose only one main revision method?\nA. To count each respondent once in each total.\nB. To prevent the use of other methods.\nC. To identify the highest-scoring student.\nD. To measure the time spent on every subject.',
    'Which revision method was most popular in August?',
    'How many more students selected self-testing in August than in July?',
    'What percentage of the students selected group discussion in August?',
    'By what percentage did the number selecting rereading textbooks decrease from July to August?',
    'State whether the following is TRUE or FALSE: The number selecting lesson videos remained unchanged.',
    'Give one reason mentioned in the passage for preferring the preparation of notes.',
    'Why cannot the table establish that students’ examination marks improved?',
    'Which conclusion is supported by the data?\nA. Every student who stopped rereading joined a discussion group.\nB. Self-testing became the choice of twice as many students.\nC. Preparing notes was the least popular method in August.\nD. The demonstrations were the only cause of every change.',
    'What proposed action could help students unable to stay after school?'
])

questions(55,[
    'Choose the correct determiner:\nThe guide gave us __________ useful information about the museum.\nA. many    B. a few    C. some    D. each',
    'Choose the correct verb form:\nBy next Monday, the team __________ all the survey reports.\nA. will have completed    B. had completed    C. was completing    D. had been completing',
    'Choose the modal that gives advice:\nYou __________ check the instructions before starting the experiment.\nA. might not    B. should    C. need not    D. cannot',
    'Choose the correct verb:\nNeither the teacher nor the students __________ ready to leave.\nA. is    B. was    C. has been    D. are',
    'Choose the correct preposition:\nThe parents were proud __________ their daughter’s achievement.\nA. of    B. at    C. for    D. by',
    'Choose the appropriate determiner:\nThere are not __________ chairs in this room.\nA. much    B. many    C. a little    D. each',
    'Complete the sentence with the correct form of the verb in brackets:\nAt nine last night, we __________ (prepare) the invitation cards.',
    'Identify the error and write its correction:\nEvery one of these paintings have a story behind it.',
    'Identify the error and write its correction:\nWe look forward to meet the visiting author.',
    'Report the statement by completing the sentence:\nNitin said, “I have finished my homework.”\nNitin said that he __________ his homework.',
    'Choose the modal expressing formal permission:\nStudents __________ use the reference section after obtaining the librarian’s permission.\nA. must    B. need    C. may    D. ought',
    'Identify the error and write its correction:\nMs Sen is senior than me in this organisation.'
])

setp(70,'The table shows the preferred activity for an inter-house festival, based on a survey of 250 Class X students. Each student selected one activity. Write an analytical paragraph in 100–120 words. Identify the main trends, compare the preferences and present an overall observation. Base your analysis only on the given data; do not invent reasons for the choices.')
fill_table(1,[
    ['Preferred festival activity','Number of students','Percentage'],
    ['Sports competitions',75,'30%'],['Music performances',60,'24%'],
    ['Science exhibition',50,'20%'],['Art displays',40,'16%'],
    ['Literary events',25,'10%'],['Total',250,'100%']
])
setp(76,'You are Dev/Diya, a resident of 32, Ashok Nagar, Jaipur. Streetlights near the local bus stop have not been working for several weeks, making it difficult for pedestrians to see obstacles after dark. Write a letter to the Editor of The City Herald, Jaipur, highlighting the problem. Explain its effect on residents and suggest prompt repairs, regular inspections and a clear system for reporting faults.')
setp(79,'You are Dev/Diya, Library Secretary of K.M. International School, 18, School Road, Jaipur. Your school ordered 40 English atlases from Scholar Books, 15, College Road, Jaipur, under Order No. KM/LIB/32 dated 4 September 2026. The delivery received on 12 September contained only 35 atlases, six of which had missing pages. Write a complaint to the Sales Manager requesting replacement of the defective copies and delivery of the missing atlases before the geography exhibition on 22 September 2026.')

setp(83,'“I learned that courage was not the absence of fear, but the triumph over it.”')
source(84,'— Nelson Mandela: Long Walk to Freedom; NCERT, First Flight','https://cdn1.byjus.com/wp-content/uploads/2019/11/NCERT-Book-for-Class-10-English-Chapter-2.pdf')
questions(86,[
    ('From whom did Mandela learn this understanding of courage?',1),
    ('Which word in the extract means “victory”?',1),
    ('Explain why a person who feels afraid can still be courageous. Relate your answer to the actions of Mandela’s fellow freedom fighters.',2)
])
setp(90,'Has given my heart\nA change of mood\nAnd saved some part\nOf a day I had rued.')
source(91,'— Dust of Snow, Robert Frost; NCERT, First Flight','https://cdn1.byjus.com/wp-content/uploads/2019/11/NCERT-Book-for-Class-10-English-Chapter-1.pdf')
questions(92,[
    ('How has the speaker’s mood changed?',1),
    ('Choose the meaning of “rued” in this context.\nA. Celebrated    B. Regretted    C. Forgotten    D. Planned',1),
    ('The continuation of a sentence across these line breaks illustrates:\nA. Enjambment    B. Simile    C. Onomatopoeia    D. Hyperbole',1),
    ('What small incident, described earlier in the poem, brings about this change?',1),
    ('Why does the speaker say that only a portion of the day was saved? Explain what this suggests about finding relief during a difficult day.',2)
])

questions(102,[
    ('What two kinds of obligation does Mandela describe, and why was it difficult for him to fulfil both under apartheid? (Nelson Mandela: Long Walk to Freedom)',3),
    ('How does the young seagull’s family respond once he begins to fly? Explain how this response contributes to his confidence. (His First Flight)',3),
    ('How does Anne use humour in her final essay for Mr Keesing, and how does it change his attitude towards her talking? (From the Diary of Anne Frank)',3),
    ('What do the baker’s distinctive dress and familiar arrival suggest about his place in the narrator’s childhood memories? (A Baker from Goa)',3),
    ('How is hospitality presented in Coorg, and how do Rajvir and Pranjol respond differently to the tea landscape in Assam? Give a relevant detail from each account. (Coorg; Tea from Assam)',3)
])
questions(108,[
    ('Why does the caged tiger ignore visitors? Explain how his behaviour helps the poet convey the effect of captivity. (A Tiger in the Zoo)',3),
    ('The speaker first favours one cause of destruction but then accepts another as equally powerful. Explain this development of thought. (Fire and Ice)',3),
    ('How does the poet’s advice for identifying the bear and the leopard create humour? Refer to the danger hidden in the advice. (How to Tell Wild Animals)',3),
    ('Why does the lost ball have a value that a new ball cannot simply replace? Explain the connection between the object and the boy’s past. (The Ball Poem)',3),
    ('How does the contrast between the adult’s instructions and Amanda’s silent thoughts shape the reader’s understanding of her life? (Amanda!)',3),
    ('How do the roots, leaves and branches work towards the trees’ escape? Explain how these images make the movement seem determined. (The Trees)',3)
],offset=5)
questions(118,[
    ('Why is Mr Herriot tempted to keep Tricki at the surgery longer than necessary? What do Mrs Pumphrey’s deliveries reveal about her affection? (A Triumph of Surgery)',2),
    ('What clue suggests that Anil knows about the theft, and how does he treat Hari Singh the next morning? (The Thief’s Story)',2),
    ('Why does Max believe the window offers a way to escape, and what mistake does he make? (The Midnight Visitor)',2),
    ('How does the young woman’s confident behaviour make Horace accept her claim to the house? Give two details. (A Question of Trust)',2),
    ('Why must Griffin escape from the London store the next morning, and how does he make himself invisible again? (Footprints Without Feet)',2)
])
setp(125,'A. The pilot’s desire to reach home leads him into danger, while an unexplained encounter helps him survive. Examine his decision to enter the storm, the difficulties that follow and the mystery surrounding his guide. How does the ending influence your judgement of his experience? (Black Aeroplane)')
setp(127,'B. How does the storm transform Lencho’s hopes? Trace his feelings from the first raindrops to his decision to write to God, explaining what his responses reveal about his livelihood and faith. (A Letter to God)')
setp(130,'A. Mrs Pumphrey’s affection and Mr Herriot’s practical care have very different effects on Tricki. Compare their approaches and explain what the story suggests about responsible care. Support your answer with Tricki’s condition and recovery. (A Triumph of Surgery)')
setp(132,'B. Fowler initially doubts whether Ausable resembles a secret agent, but the encounter with Max changes his understanding. Explain how Ausable uses calm thinking and believable details to overcome a dangerous situation. (The Midnight Visitor)')
setp(133,'— END OF QUESTION PAPER: SET B —')

doc.core_properties.title='Class X English Half-Yearly Question Paper — Set B'
doc.core_properties.subject='Parallel Set B | 80 marks | 3 hours | English Language & Literature (184)'
doc.core_properties.comments='Follows supplied Class X paper: 20/20/40 sections; Q3 10 of 12; Q8 two prose and two poetry; Q9 four of five; Q10 6 marks; Q11 4 marks. Literature allocation: prose 16, poetry 12, supplementary reader 12.'
doc.save(out)

check=Document(out)
text='\n'.join(p.text for p in check.paragraphs)
def block(a,b): return text[text.index(a):text.index(b)]
for a,b,count in [('Q3.','Q4.',12),('Q6.','Q7.',3),('Q7.','Q8.',5),('Q8.','Q9.',11),('Q9.','Q10.',5)]:
    assert len(re.findall(r'^\([ivx]+\)',block(a,b),re.M))==count
assert sum(map(int,re.findall(r'\[(\d+)\]',block('Q6.','Q7.'))))==4
assert sum(map(int,re.findall(r'\[(\d+)\]',block('Q7.','Q8.'))))==6
assert sum([10,10,10,5,5,4,6,12,8,6,4])==80
assert (4+6+6,6+6,8+4)==(16,12,12)
assert sum([100,60,40,30,20])==sum([60,50,80,40,20])==250
assert sum([75,60,50,40,25])==250 and sum([30,24,20,16,10])==100
assert 350<=len(' '.join(passage1).split())<=450
assert 200<=len(' '.join(passage2).split())<=275
assert len(check.tables)==2
assert len(check.element.xpath('//w:br[@w:type="page"]'))==7
with zipfile.ZipFile(out) as z: assert z.testzip() is None
print(out.resolve())
print('Verified: 80 marks, all question and choice counts, 16/12/12 literature allocation, and both data tables.')
print('Passage word counts:',len(' '.join(passage1).split()),len(' '.join(passage2).split()))
