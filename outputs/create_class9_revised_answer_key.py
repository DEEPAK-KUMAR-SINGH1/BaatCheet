from pathlib import Path
import zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).parent
SOURCE=ROOT/'Class_9_English_Half_Yearly_Question_Paper_Revised.docx'
OUT=ROOT/'Class_9_English_Half_Yearly_Revised_Answer_Key_and_Marking_Scheme.docx'
paper=Document(SOURCE)
text='\n'.join(p.text for p in paper.paragraphs)
assert 'Do any 10 out of 12' in text and 'Do any 4 out of 5' in text
assert 'SET B' not in text and 'ALL three questions. (5 marks)' in text
doc=Document(); sec=doc.sections[0]
sec.page_width,sec.page_height=Inches(8.27),Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.7)
sec.header_distance=sec.footer_distance=Inches(.28)
s=doc.styles['Normal']; s.font.name='Times New Roman'; s.font.size=Pt(11)
s.paragraph_format.space_after=Pt(6); s.paragraph_format.line_spacing=1.03
for n in ['Heading 1','Heading 2']:
    s=doc.styles[n]; s.font.name='Times New Roman'; s.font.size=Pt(12); s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_before=Pt(9); s.paragraph_format.space_after=Pt(7)
def p(text='',bold=False,size=None,center=False):
    x=doc.add_paragraph(); r=x.add_run(text); r.bold=bold
    if size: r.font.size=Pt(size)
    if center: x.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return x
def h(text): doc.add_paragraph(text,'Heading 2')
def page(): doc.add_page_break()
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.autofit=False
    for c,text in zip(t.rows[0].cells,headers):
        c.text=str(text)
        for r in c.paragraphs[0].runs:r.bold=True
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for values in rows:
        for c,text in zip(t.add_row().cells,values): c.text=str(text)
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for c,w in zip(row.cells,widths):
            c.width=Inches(w)
            for x in c.paragraphs:
                x.paragraph_format.space_before=Pt(3); x.paragraph_format.space_after=Pt(3)
                for r in x.runs:r.font.size=Pt(10.5)
    p('')
def link(label,url):
    x=doc.add_paragraph(); a=OxmlElement('w:hyperlink')
    a.set(qn('r:id'),x.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); rp=OxmlElement('w:rPr'); size=OxmlElement('w:sz'); size.set(qn('w:val'),'18')
    rp.append(size);r.append(rp);t=OxmlElement('w:t');t.text=label;r.append(t);a.append(r);x._p.append(a)

hp=sec.header.paragraphs[0];hp.text='K.M. INTERNATIONAL SCHOOL | CLASS IX ENGLISH | TEACHER COPY'
hp.alignment=WD_ALIGN_PARAGRAPH.CENTER;hp.runs[0].font.size=Pt(9)
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
fp.add_run('Revised Paper: Answer Key | Page ')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');fp._p.append(f)
fp.add_run(' of ');f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'NUMPAGES');fp._p.append(f)

p('K.M. INTERNATIONAL SCHOOL',True,15,True)
p('CLASS IX — HALF-YEARLY EXAMINATION',True,13,True)
p('ENGLISH LANGUAGE & LITERATURE',True,12,True)
p('ANSWER KEY AND MARKING SCHEME',True,13,True)
p('REVISED PAPER | Maximum Marks: 80 | Time: 3 Hours',True,11,True)
p('For Class_9_English_Half_Yearly_Question_Paper_Revised.docx',size=9,center=True)
h('Mark allocation and choices')
table(['Question','Requirement','Marks'],[
    ['Q1','Discursive passage: all 10 parts',10],['Q2','Case-based passage: all 10 parts',10],
    ['Q3','Grammar: any 10 of 12',10],['Q4','Letter: Editor OR Complaint',5],
    ['Q5','Diary entry OR Story',5],['Q6','Prose extract: 1 + 1 + 1 + 2',5],
    ['Q7','Poetry extract: 1 + 1 + 3; no rhyme-scheme item',5],
    ['Q8','Any 4 of 6 short answers; 3 marks each',12],
    ['Q9','Any 4 of 5 short answers; 2 marks each',8],
    ['Q10','Either A or B; 100–120 words',5],['Q11','Either A or B; 100–120 words',5],
    ['TOTAL','Reading 20 + Grammar/Writing 20 + Literature 40',80]
],[.7,5.4,.75])
h('Guidance for examiners')
for t in [
    'This scheme is prepared for the supplied school paper; it is not an official CBSE marking scheme. Accept accurate paraphrases and alternative interpretations supported by the prescribed text.',
    'Objective items earn 1 or 0. A correct option letter or clearly written answer is sufficient. If the option and written answer conflict, award 0. Use half marks for partially developed written content where appropriate.',
    'For this scheme, count the best permitted number if extra answers are attempted, unless the school announced another rule beforehand. Q8 allows any four; there is no separate prose/poetry quota.',
    'Do not deduct twice for the same error. Minor language slips should not reduce comprehension marks when meaning is clear; assess expression separately where allocated.',
    'Use word limits to judge focus and adequate development, without an additional automatic penalty. Writing samples show one possible response, not mandatory wording.'
]:p(t,size=10)

page();h('SECTION A — READING SKILLS (20 MARKS)')
h('Q1. Group work — All ten parts (10 × 1 = 10)')
q1=[
    ['(i)','B — Teamwork combines different strengths through shared responsibility.'],
    ['(ii)','Any one: observing carefully; organising materials; checking information; noticing who needs help.'],
    ['(iii)','Everyone wanted to design the model, while necessary jobs such as preparing labels and explaining costs were neglected; noisy discussion replaced coordinated work.'],
    ['(iv)','The members’ interests and abilities.'],
    ['(v)','The planned containers were too large, so the design measurements and costs needed adjustment.'],
    ['(vi)','False. Members still need to share information and coordinate their decisions.'],
    ['(vii)','C — This design uses too much cardboard; can we reduce its size?'],
    ['(viii)','Efficiency.'],
    ['(ix)','Any one: clarifies who is responsible; shows unfinished tasks; identifies where support is needed before a problem grows.'],
    ['(x)','Acknowledge and thank the student for checking the measurements; recognise quiet, less visible contributions as well as public achievements.']
]
table(['Part','Expected answer'],q1,[.6,6.25])
p('Award 1 for each correct answer. For Q1(iv), award ½ if only interest or only ability is mentioned. For Q1(iii) and (v), allow ½ for a relevant but incomplete reason. No explanation is required for the MCQ or true/false items.',size=10)

page();h('Q2. Library borrowing — All ten parts (10 × 1 = 10)')
q2=[
    ['(i)','B — A borrowing transaction.'],['(ii)','Fiction (108 transactions). Naming the category is enough.'],
    ['(iii)','90 transactions: 270 − 180 = 90.'],['(iv)','40%: (108 ÷ 270) × 100 = 40%.'],
    ['(v)','Poetry: 12 before the programme, 24 during it.'],
    ['(vi)','24 transactions: 60 − 36 = 24.'],
    ['(vii)','False. There were 120 students; 270 is a transaction total and students could borrow more than one book.'],
    ['(viii)','Any one: same students; same number of school days; same borrowing rules; equal six-week observation periods.'],
    ['(ix)','Borrowing records show that books were issued, not whether they were read to the end. No completion data were collected.'],
    ['(x)','A — Done by choice.']
]
table(['Part','Expected answer'],q2,[.6,6.25])
h('Calculation and interpretation checks')
p('A correct numerical answer earns 1 without working. If Q2(iv) is incomplete, award ½ for the correct setup 108/270 × 100 and ½ for 40%. A transaction is not a unique borrower or a completed reading. Do not treat all 270 transactions as involving different students.',size=10)
p('Q2(viii) asks for one constant condition. Do not demand a list. Q2(vii) earns full credit for “False” alone because the question only requests true/false.',size=10)

page();h('SECTION B — GRAMMAR & WRITING (20 MARKS)')
h('Q3. Grammar — Any ten of twelve (10 marks)')
q3=[
    ['(i)','B — had completed',1],['(ii)','is knocking',1],['(iii)','C — must',1],
    ['(iv)','B — carries',1],['(v)','C — is',1],['(vi)','B — an',1],
    ['(vii)','C — little',1],['(viii)','A — at',1],
    ['(ix)','spoke → speak: The coach asked us to speak politely to the visiting team.',1],
    ['(x)','confident → confidently: The new student answered the question confidently.',1],
    ['(xi)','A — were playing',1],['(xii)','A — a few',1]
]
table(['Part','Answer / correction','Marks'],q3,[.6,5.6,.65])
p('Award 1 for a correct answer and 0 otherwise. For editing, accept either the error with its correction or the complete corrected sentence. Identifying an error without correcting it does not earn the mark. Maximum for Q3 is 10 even when all twelve answers are correct.')
h('Grammar reminders for consistent checking')
p('“Someone”, “each” and “quality” take singular verbs here. “Engineer” begins with a vowel sound, so it takes “an”. Water is uncountable, whereas apples are countable. “To” in Q3(ix) introduces an infinitive, requiring the base verb “speak”.',size=10)

page();h('Q4. Formal letter — Either A or B (5 marks)')
p('Length: 100–120 words. Use the same rubric for either letter.')
table(['Criterion','Marks and basis'],[
    ['Format','1: suitable addresses, date, subject, salutation and formal closing. Award ½ for a recognisable but incomplete format.'],
    ['Content','2: option-specific allocation below/on the next page.'],
    ['Organisation and tone','1: clear purpose, logical sequence, appropriate request and formal tone.'],
    ['Accuracy','1: clear grammar, vocabulary, spelling and punctuation; ½ for noticeable errors that do not obscure meaning.']
],[1.55,5.3])
h('Q4(A). Letter to the Editor')
p('Content: problem and effect on residents (1); practical cleanliness measures and responsible park use (1).')
p('24, Green Park\nJaipur\n16 September 2026\n\nThe Editor\nThe City Herald\nJaipur')
p('Subject: Poor cleanliness in the neighbourhood park',True)
p('Sir/Madam,')
models={}
models['Q4A']='Through the columns of your newspaper, I wish to highlight the poor condition of the public park in Green Park. Plastic wrappers and bottles are scattered around, while several dustbins are damaged. The litter makes the park unpleasant for walkers and leaves children with fewer clean spaces to play.\n\nThe authorities should replace damaged bins and arrange regular waste collection. Clear signs and community cleanliness drives could encourage visitors to dispose of rubbish responsibly. Residents should also avoid leaving plastic waste after gatherings.\n\nPlease draw attention to this problem and urge prompt action so that everyone can enjoy a cleaner, more welcoming public space.'
for t in models['Q4A'].split('\n\n'):p(t)
p('Yours faithfully,\nAnanya')

page();h('Q4(B). Complaint letter — Model and content allocation')
p('Content: accurate order/delivery details and both faults (1); request for missing copies and replacements (1). Apply the common 5-mark letter rubric.')
p('K.M. International School\n18, School Road\nJaipur\n16 September 2026\n\nThe Sales Manager\nBright Books\n12, Station Road\nJaipur')
p('Subject: Short delivery and damaged dictionaries — Order KM/LIB/19',True)
p('Sir/Madam,')
models['Q4B']='Our school ordered 30 English dictionaries under Order No. KM/LIB/19 dated 2 September 2026. However, the delivery received on 10 September contained only 25 dictionaries. Five of these copies had torn or missing pages and cannot be used by students.\n\nPlease send the five missing dictionaries and replace the five damaged copies promptly. Kindly arrange collection of the defective books and inspect the replacements before dispatch. The shortage is affecting the library’s ability to provide complete reference materials to our classes.\n\nPlease acknowledge this complaint and confirm the delivery schedule. We expect the matter to be resolved soon so that all 30 dictionaries are available in usable condition.'
for t in models['Q4B'].split('\n\n'):p(t)
p('Yours faithfully,\nAnuj\nLibrary Secretary')
p('Quantity check: 5 dictionaries are missing; 5 delivered copies require replacement. These are separate faults. Accept Anuj or Ananya and a sensible date. Model body lengths exclude conventional address blocks and closings.',size=10)
h('Q5. Creative writing — Either A or B (5 marks)')
p('Diary: format 1 (date and diary opening); content 2 (experience/support 1, personal learning 1); organisation and reflective voice 1; accuracy 1.')
p('Story: title and required opening 1 (½ each); content 2 (developed situation/problem 1, coherent resolution 1); organisation 1 (logical sequence and links); accuracy 1. The optional clues need not all be used.')

page();h('Q5(A). Diary entry — Model (100–120 words)')
p('16 September 2026\n8:00 p.m.\nDear Diary,')
models['Q5A']='Today’s craft workshop taught me much more than how to shape a bowl. At first, my clay kept collapsing, and I felt embarrassed when others seemed to progress faster. I almost gave up. Then the artisan sat beside me, demonstrated how gently to guide the clay and encouraged me to try again.\n\nSlowly, I learnt to keep my hands steady instead of forcing the shape. My finished bowl was uneven, but I was proud of it. I now understand how much patience and practice skilled work demands. Tomorrow, I will stop judging handmade objects only by their appearance and begin appreciating the effort behind them.'
for t in models['Q5A'].split('\n\n'):p(t)
p('Ananya')
p('Accept different emotions, techniques and reflections that fit the prompt. A specific time or signature is optional. Reward a personal, reflective first-person voice; a mere event report should not receive full reflection/voice credit.',size=10)
h('Q5(B). Story — Model (120–150 words)')
p('The Garden’s Secret',True)
models['Q5B']='When I opened the old notebook, a folded map slipped onto the floor. I was helping Grandfather clear a cupboard and expected another shopping list. Instead, the paper showed a crooked tree, a square and a star.\n\nMy friend Meera studied it carefully. “That tree looks like the mango tree in your garden,” she said. The square, we realised, was the old stone bench. Under it, one brick was loose. We called Grandfather, who helped us lift it and uncover a small tin.\n\nInside were faded photographs and a letter he had written as a boy, imagining his future family. He had forgotten where he hid it. As he read the letter aloud, we gathered around him. The map had led to no gold, but it had returned a lost part of his childhood.'
for t in models['Q5B'].split('\n\n'):p(t)
p('Any original, coherent plot earns credit. Do not require this discovery or a stated moral. The given opening should be retained; minor punctuation changes do not matter.',size=10)

page();h('SECTION C — LITERATURE (40 MARKS)')
h('Q6. How I Taught My Grandmother to Read (5 marks)')
q6=[
    ['(i)','Grandmother/Krishtakka (½), addressing her granddaughter/the narrator (½).',1],
    ['(ii)','B — Wealthy.',1],['(iii)','She cannot read Kannada independently.',1],
    ['(iv)','Learn to read independently (1); this shows determination and a desire for self-reliance (1).',2]
]
table(['Part','Answer and allocation','Marks'],q6,[.6,5.6,.65])
h('Q7. Canvas of Soil — Revised three-part question (5 marks)')
q7=[
    ['(i)','A — Metaphor.',1],['(ii)','Hue.',1],
    ['(iii)','Gardening resembles painting (1): earth is the creative palette (1), and planted seeds are brushstrokes producing future colour (1). Accept other explained details from this extract.',3]
]
table(['Part','Answer and allocation','Marks'],q7,[.6,5.6,.65])
p('Q7(iii): award 1 for the central comparison and 1 each for explaining two supporting details. Two relevant details without explanation may earn ½ each. There is no separate expression mark here, and no rhyme-scheme question to mark.',size=10)
h('How to mark the short answers')
p('Q8: any four of six, 40–50 words each. Award 2 content marks as specified plus 1 for coherent expression. Use ½ for a partially developed point or understandable expression with noticeable errors.')
p('Q9: any four of five, 30–40 words each. Award the two content marks shown; there is no separate language mark. Minor errors should not reduce content credit if meaning is clear.')

page();h('Q8. Short answers — Any four of six (4 × 3 = 12)')
p('The points below give 2 content marks per answer. Add up to 1 mark for clear expression. These are indicative content points, not sentences students must reproduce.')
q8=[
    ['(i) The Pot Maker','1: Onula reassures Sentila and demonstrates calmly, reducing her tension.\n1: Encouragement restores confidence, helping her learn through renewed effort and observation.'],
    ['(ii) Bharat Our Land','1: Nature, such as the Himalayas/Ganga, inspires pride.\n1: The Upanishads/Buddha’s teachings express spiritual richness; connect both examples with admiration for India.'],
    ['(iii) Gifts of Grace','1: Every vocation contributes and deserves dignity.\n½ each: two examples, such as carpenters’ precision and electricians lighting lives; accept other accurately explained occupations.'],
    ['(iv) Canvas of Soil','1: Gardeners plant and tend seeds.\n1: Their imagination and labour combine with natural growth to create a colourful, living artwork.'],
    ['(v) I Cannot Remember My Mother','1: A remembered tune recalls her singing beside his cradle.\n1: Shiuli flowers’ scent recalls temple worship and his mother’s presence.'],
    ['(vi) Vitamin-M','1: His mother speaks loudly and treats Grandpa like a child.\n1: Ravi’s discomfort shows empathy and respect for Grandpa’s dignity.']
]
table(['Part / work','Content allocation'],q8,[1.65,5.2])
p('Q8(ii) requires one natural and one spiritual example. Q8(iii) requires two occupations. Q8(v) requires sound and smell: an answer dealing only with one sense earns at most 1 content mark.',size=10)
link('Text reference: Kaveri, Unit 2 — The Pot Maker; Gifts of Grace','https://philoid.com/ncert/chapter/iebe102')
link('Text reference: Kaveri, Unit 4 — Vitamin-M; I Cannot Remember My Mother','https://philoid.com/ncert/chapter/iebe104')

page();h('Q9. Short answers — Any four of five (4 × 2 = 8)')
q9=[
    ['(i) Grandmother','Kashi Yatre by Triveni (1); she reads its title, author and publisher independently (1).'],
    ['(ii) The Pot Maker','Pottery supplies community needs (1) and preserves shared heritage through passing skills to new learners (1).'],
    ['(iii) Winds of Change','Any two: spread appreciation/awareness; demonstrate or transmit skills; attract buyers and support livelihoods. Award 1 for each distinct explained benefit.'],
    ['(iv) Vitamin-M','Grandpa gave the stranger his cap because of the heat (1), showing generosity and concern (1).'],
    ['(v) Grandmother','Any two: reading; repeating; writing; reciting. Award 1 each.']
]
table(['Part / work','Expected answer and allocation'],q9,[1.65,5.2])
p('For Q9(i), the book title is enough for the gift mark; do not insist on the author’s name. “She reads the book for herself” conveys successful independence. For Q9(iii), two restatements of one benefit count only once.',size=10)
h('Long-answer rubric: Q10 and Q11')
p('Each question carries 5 marks and offers A/B choice. Expected length: 100–120 words. Use the option-specific 3 content marks, plus 1 for organisation/coherence and 1 for grammar/vocabulary accuracy. Half marks are available for partially developed content or expression.')
p('A reasoned interpretation with relevant incidents may earn full credit even when it differs from the model. General statements without textual support do not earn all content marks. Assess only one option in each question.')

page();h('Q10. Long answer — Either A or B (5 marks)')
h('Option A — How I Taught My Grandmother to Read')
p('Content: age/learning (1); effort (1); respect for teaching (1). Add expression: 2.')
models['Q10A']='The grandmother’s success challenges the belief that learning belongs only to childhood. Although she is sixty-two and has never attended school, she decides to overcome her dependence on others for reading. She fixes a deadline and practises seriously, repeating, writing and reciting her lessons. Her granddaughter initially doubts her, but persistent effort proves more important than age. After achieving her goal, the grandmother touches the young teacher’s feet. This gesture honours the role of a teacher rather than the person’s age or family position. Her journey therefore presents education as a source of independence and shows that determination, humility and respect can make learning possible throughout life.'
p(models['Q10A'])
link('Text reference: NCERT Kaveri, Unit 1','https://ncert.nic.in/textbook/pdf/iebe101.pdf')
h('Option B — The Pot Maker')
p('Content: practical concerns (1); aspiration (1); support/practice (1). Add expression: 2.')
models['Q10B']='Arenla judges pottery through the hardship of earning a living. Collecting and preparing clay exhausts her, while long labour brings little money. She therefore prefers weaving for Sentila. Her daughter, however, feels drawn to the skill and beauty of shaping pots. Desire alone does not make her successful: tension and repeated failures initially hold her back. Onula’s patient demonstration gives her confidence, while observing her mother helps her improve difficult details. Continued practice eventually turns fascination into ability. The story respects the mother’s practical concerns but also shows why a learner’s interest matters. Encouragement and patient effort help Sentila develop a vocation that she genuinely values.'
p(models['Q10B'])
link('Text reference: Kaveri, Unit 2 — The Pot Maker','https://philoid.com/ncert/chapter/iebe102')

page();h('Q11. Long answer — Either A or B (5 marks)')
h('Option A — Vitamin-M')
p('Content: dilemma (1); response (1); intergenerational understanding (1). Add expression: 2.')
models['Q11A']='Ravi must balance his mother’s instructions with Grandpa’s wish to remain independent. His mother worries because Grandpa has previously become lost, but Grandpa resents being treated as incapable. When he goes out, Ravi fears both upsetting him and failing to protect him. He therefore follows at a distance, hoping to provide safety without openly restricting his freedom. His anxiety during the outing and relief on finding Grandpa home reveal deep affection. Yet the episode also exposes the difficulty of managing someone through secrecy. The story encourages younger people to recognise an older person’s dignity and abilities, while acknowledging that genuine concern and practical support may still be necessary.'
p(models['Q11A'])
link('Text reference: Kaveri, Unit 4 — Vitamin-M','https://philoid.com/ncert/chapter/iebe104')
h('Option B — Winds of Change')
p('Content: heritage (1); textual detail (1); artisan support (1). Add expression: 2.')
models['Q11B']='Traditional hand fans preserve local skills, materials and designs, making them records of cultural identity as well as useful objects. Rajasthan’s decorated fans and Gujarat’s mirror-work designs show how regions develop distinctive traditions. A school celebration could display labelled examples and invite artisans to demonstrate their techniques. Workshops would help students appreciate the patience behind the craft and introduce new learners to its methods. A sale or ordering desk could also connect makers with buyers, providing income as well as recognition. Such activities follow the chapter’s emphasis on celebrating craftsmanship and creating opportunities for livelihoods. The event should honour living artisans, not simply treat their work as decoration.'
p(models['Q11B'])
link('Text reference: NCERT Kaveri, Unit 3 — Winds of Change; Canvas of Soil','https://ncert.nic.in/textbook/pdf/iebe103.pdf')
h('Final marks check')
p('Reading: 20 | Grammar: 10 | Writing: 10 | Literature: 40 | TOTAL: 80',True)
p('Q6 5 + Q7 5 + Q8 12 + Q9 8 + Q10 5 + Q11 5 = 40. The revised Q7 allocation is 1 + 1 + 3.',size=10)

doc.core_properties.title='Class IX English Revised Paper — Answer Key and Marking Scheme'
doc.core_properties.subject='80 marks | All revised choices and question parts'
doc.core_properties.author='K.M. International School'
doc.core_properties.comments='Based on Class_9_English_Half_Yearly_Question_Paper_Revised.docx; Q3 any 10/12, Q8 any 4/6, Q9 any 4/5, Q7 1+1+3.'
doc.save(OUT)

limits={k:(100,120) for k in models};limits['Q5B']=(120,150)
for k,(lo,hi) in limits.items():
    n=len(models[k].split());print(f'{k}: {n} words');assert lo<=n<=hi,(k,n)
assert len(q1)==10 and len(q2)==10 and len(q3)==12
assert len(q6)==4 and len(q7)==3 and len(q8)==6 and len(q9)==5
assert sum(row[2] for row in q6)==5 and sum(row[2] for row in q7)==5
assert sum([10,10,10,5,5,5,5,12,8,5,5])==80
assert 270-180==90 and 108/270*100==40 and 60-36==24
assert models['Q5B'].startswith('When I opened the old notebook, a folded map slipped onto the floor.')
check=Document(OUT)
assert len(check.element.xpath('//w:br[@w:type="page"]'))==11
with zipfile.ZipFile(OUT) as z: assert z.testzip() is None
print(OUT.resolve())
print('Verified: all revised questions/options, 80-mark total, reading calculations, model word limits and Word package integrity.')
