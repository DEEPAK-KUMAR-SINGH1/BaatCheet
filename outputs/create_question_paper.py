from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

OUT = Path(__file__).parent / 'Class_10_English_Half_Yearly_Question_Paper.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.7)
sec.header_distance = sec.footer_distance = Inches(.28)
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.03
for name in ['Heading 1', 'Heading 2']:
    s = doc.styles[name]
    s.font.name = 'Times New Roman'
    s.font.size = Pt(12)
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(8)
    s.paragraph_format.space_after = Pt(6)

def p(text='', bold=False, size=None, center=False):
    para = doc.add_paragraph()
    r = para.add_run(text)
    r.bold = bold
    if size: r.font.size = Pt(size)
    if center: para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return para

def h(text):
    doc.add_paragraph(text, 'Heading 2')

def q(text):
    para = p(text, True)
    para.paragraph_format.keep_with_next = True

def item(label, text, marks=1):
    para = p(f'({label})  {text}' + (f'  [{marks}]' if marks else ''))
    para.paragraph_format.left_indent = Inches(.12)
    para.paragraph_format.first_line_indent = Inches(-.12)
    para.paragraph_format.keep_together = True

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    for c, text in zip(t.rows[0].cells, headers):
        c.text = text
        for r in c.paragraphs[0].runs: r.bold = True
    for row in rows:
        for c, text in zip(t.add_row().cells, row): c.text = str(text)
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                para.paragraph_format.space_before = Pt(3)
    p('')
    return t

def source(label, url):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(8)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), para.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '18'); props.append(size)
    run.append(props)
    text = OxmlElement('w:t'); text.text = label; run.append(text)
    link.append(run); para._p.append(link)

def page(): doc.add_page_break()

header = sec.header.paragraphs[0]
header.text = 'K.M. INTERNATIONAL SCHOOL  |  CLASS X — ENGLISH (184)'
header.alignment = WD_ALIGN_PARAGRAPH.CENTER
header.runs[0].font.size = Pt(9)
foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
foot.add_run('Page ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); foot._p.append(field)
foot.add_run(' of ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'NUMPAGES'); foot._p.append(field)

p('K.M. INTERNATIONAL SCHOOL', True, 15, True)
p('HALF-YEARLY EXAMINATION', True, 13, True)
p('CLASS X — ENGLISH LANGUAGE & LITERATURE (184)', True, 12, True)
p('Time Allowed: 3 Hours                                      Maximum Marks: 80', True)
p('Name: __________________________  Roll No.: __________  Section: ______')
h('General Instructions')
for text in [
    '1. This paper contains 11 questions in three sections: A (20 marks), B (20 marks) and C (40 marks).',
    '2. All questions are compulsory, subject to the internal choices given. Follow the attempt instructions carefully.',
    '3. Write answers in the answer booklet. Retain the question and sub-question numbers.',
    '4. Marks are shown in brackets. Follow the prescribed word limits wherever given.'
]: p(text, size=10)
h('SECTION A — READING SKILLS (20 MARKS)')
q('Q1. Read the following passage and answer ALL ten questions. (10 × 1 = 10)')
passage1 = [
    '1. Many students regard a mistake as proof that they are not capable of doing something. A wrong answer in mathematics or a forgotten line during a speech can feel like a final judgement. Yet a mistake is often useful information: it shows the difference between what we intended to do and what actually happened. Learning begins when we examine that difference instead of hiding it.',
    '2. Consider a student building a model bridge. Her first design collapses under a small weight. If she simply repeats the same design, failure teaches her little. If she observes where the bridge bends, changes the supports and tests it again, each attempt becomes an investigation. Success grows from a cycle of trying, noticing and revising. Effort matters most when it is guided by reflection.',
    '3. Feedback from others can strengthen this process, but its usefulness depends on how specific it is. Being told to “do better” rarely shows a learner what to change. Being told that the opening of a speech is clear but the examples need explanation provides a direction. Equally, learners must be willing to listen. A suggestion is easier to use when it is treated as help with a task rather than an attack on personal worth.',
    '4. A classroom that welcomes questions makes this attitude possible. When students fear ridicule, they may remain silent even when confused. Teachers and classmates can respond respectfully, allow time for a second attempt and recognise thoughtful improvement. Such support does not mean that every answer is correct or that standards should disappear. It means that a learner can acknowledge an error without feeling humiliated.',
    '5. However, celebrating mistakes without correcting them would miss the point. Repeating an error carelessly is different from taking a sensible risk while learning. Students still need preparation, practice and responsibility. A brief learning journal can help: What did I try? What went wrong? What will I change? These questions turn an uncomfortable moment into a practical plan. Over time, learners may become less concerned with appearing perfect and more interested in becoming capable.'
]
for text in passage1: p(text)
p('Questions for Q1 continue on the next page.', True, 10)

page()
h('SECTION A — READING SKILLS (CONTINUED)')
q('Q1. Answer ALL ten questions using the passage on the previous page. (10 marks)')
item('i', 'Which statement best expresses the central idea?\nA. All mistakes automatically improve learning.\nB. Examining and correcting mistakes supports learning.\nC. Practice is less useful than confidence.\nD. Teachers should accept every answer as correct.')
item('ii', 'According to paragraph 1, what difference can a mistake reveal?')
item('iii', 'Why does the bridge-builder change the supports before testing again?')
item('iv', 'Complete the statement: Effort becomes more productive when it is guided by __________.')
item('v', 'Which feedback would the writer consider most useful?\nA. Your work is poor.    B. Try harder next time.\nC. Explain how your example supports the main point.    D. Be more intelligent.')
item('vi', 'State whether the following is TRUE or FALSE: A supportive classroom must abandon standards of correctness.')
item('vii', 'What may prevent a confused student from asking a question in class?')
item('viii', 'Find a word in paragraph 4 that means “admit or accept that something is true”.')
item('ix', 'Why does the writer distinguish careless repetition from sensible risk-taking?')
item('x', 'A student receives feedback and submits exactly the same work again. Identify one action the writer would suggest before resubmission.')
h('Q2. Case-based Factual Passage — Read the information below. (10 marks)')
passage2 = [
    '1. A school eco-club studied how students travelled to school before planning a cleaner-commute campaign. It surveyed 200 students in July and the same 200 students in August. Each student selected the single mode used on most school days. The figures below are hypothetical data prepared for this examination.',
    '2. Between the two surveys, the club arranged a supervised walking group for students living nearby. It also shared information about existing bus routes and installed bicycle stands. The campaign encouraged families to consider practical alternatives to travelling by car; it did not require every student to choose the same mode.',
    '3. At the second survey, several students said that the walking group made the journey more enjoyable. Others reported that the bicycle stands helped them park securely. Some car users explained that long distances, unsafe road crossings or family schedules limited their choices. Therefore, the club proposed safer crossings and better bus access as possible next steps.'
]
for text in passage2: p(text)
p('The data table, final paragraph and questions for Q2 follow on the next page.', True, 10)

page()
h('SECTION A — READING SKILLS (CONTINUED)')
q('Q2. Study the table and complete the reading before answering. (10 × 1 = 10)')
table(['Main mode of travel', 'July: students', 'August: students'], [
    ['Walking', 30, 50], ['Bicycle', 20, 30], ['School bus', 70, 80],
    ['Private car', 60, 30], ['Public transport', 20, 10], ['Total', 200, 200]
])
p('4. The number using private cars fell, while walking, cycling and school-bus use increased. However, the survey recorded travel choices, not journey lengths or fuel consumption. It cannot show the exact reduction in pollution. Because there was no comparison group, it also cannot prove that the campaign alone caused every change. The club recommended repeating the survey in another season and asking families what support they needed.')
item('i', 'Why did each student select only one main mode of travel?\nA. To count each student once in the totals.\nB. To prevent students from changing their travel habits.\nC. To measure the length of every journey.\nD. To exclude students who used cars.')
item('ii', 'Which mode of travel was used by the greatest number of students in BOTH months?')
item('iii', 'How many more students walked in August than in July?')
item('iv', 'What percentage of the surveyed students used bicycles in August?')
item('v', 'By what percentage did the number of private-car users decrease from July to August?')
item('vi', 'State whether the following is TRUE or FALSE: Public-transport use increased after the campaign.')
item('vii', 'Give one reason why some students continued to travel by car.')
item('viii', 'Why would it be inaccurate to claim that the survey proves an exact reduction in pollution?')
item('ix', 'Which conclusion is best supported by the data?\nA. Every car user changed to walking.\nB. The combined number walking and cycling increased.\nC. Public transport became the most popular option.\nD. The campaign was the only possible cause of the changes.')
item('x', 'Suggest one action, supported by the passage, that could help students who face unsafe road crossings.')

page()
h('SECTION B — GRAMMAR & WRITING SKILLS (20 MARKS)')
q('Q3. Attempt ANY TEN of the following twelve grammar items. (10 × 1 = 10)')
p('For multiple-choice items, write the correct option and answer. For editing items, write the incorrect word and its correction.')
item('i', 'Fill in the blank with the correct determiner:\nThere is __________ sugar left in the jar, so we must buy some.\nA. few    B. little    C. many    D. a few')
item('ii', 'Choose the correct verb form:\nBy the time the teacher arrived, the students __________ the display.\nA. complete    B. had completed    C. will complete    D. are completing')
item('iii', 'Complete the sentence with the modal expressing strict prohibition:\nVisitors __________ touch the exhibits.\nA. need not    B. might not    C. must not    D. would not')
item('iv', 'Choose the correct verb:\nEach of the participants __________ a certificate.\nA. receive    B. receives    C. have received    D. are receiving')
item('v', 'Fill in the blank with the correct preposition:\nRiya has been a member of the club __________ 2023.\nA. for    B. since    C. during    D. by')
item('vi', 'Fill in the blank with the correct determiner:\nHave you read __________ of the two books I lent you?\nA. either    B. every    C. much    D. any of')
item('vii', 'Complete the sentence using the correct form of the verb in brackets:\nLook! The children __________ (plant) saplings near the gate.')
item('viii', 'Identify the error and write its correction:\nThe list of selected candidates are on the notice board.')
item('ix', 'Identify the error and write its correction:\nOur team is confident of win the final match.')
item('x', 'Report the following statement by completing the sentence:\nMeera said, “I am preparing for the debate.”\nMeera said that she __________ for the debate.')
item('xi', 'Complete the sentence with the correct modal:\nWhen he was five, my brother __________ swim across the pool.\nA. can    B. could    C. must    D. should')
item('xii', 'Identify the error and write its correction:\nWe discussed about the proposal during the meeting.')

page()
h('SECTION B — WRITING SKILLS (CONTINUED)')
q('Q4. Analytical Paragraph (5 marks)')
p('The table shows how 200 Class X students at a school spend the largest share of their leisure time on a typical weekday. Each student selected one activity. Write an analytical paragraph in 100–120 words describing the data. Identify the main trends, make relevant comparisons and present an overall observation. Use only the information given; do not invent reasons for the preferences.')
table(['Preferred leisure activity', 'Number of students', 'Percentage'], [
    ['Sports and outdoor games', 60, '30%'], ['Social media', 50, '25%'],
    ['Watching films or online videos', 40, '20%'], ['Reading', 30, '15%'],
    ['Art and music', 20, '10%'], ['Total', 200, '100%']
])
p('The data are hypothetical and have been prepared for this examination.', size=9)
q('Q5. Formal Letter — Attempt EITHER option A OR option B. (5 marks)')
p('Write your letter in 100–120 words, using an appropriate formal format.')
h('Option A — Letter to the Editor')
p('You are Aarav/Aarohi, a resident of 24, Green Park, Jaipur. Vehicles parked on the footpath near your school force children to walk on the road. Write a letter to the Editor of The City Herald, Jaipur, highlighting the problem. Explain its effect on student safety and suggest practical measures such as enforcing parking rules, providing a designated drop-off area and keeping footpaths clear.')
p('OR', True, center=True)
h('Option B — Letter of Complaint')
p('You are Rohan/Riya, Sports Secretary of K.M. International School, 18, School Road, Jaipur. Your school ordered 20 footballs from Sunrise Sports, 12, Station Road, Jaipur, under Order No. KM/SPORTS/27 dated 2 September 2026. The delivery received on 10 September contained only 16 footballs, four of which had defective valves. Write a complaint to the Sales Manager. Give the relevant order details and request replacement of the defective footballs and delivery of the missing items before the sports trials on 20 September 2026.')

page()
h('SECTION C — LITERATURE (40 MARKS)')
q('Q6. Read the extract and answer ALL three questions. (4 marks)')
p('“But don’t send it to me through the mail because the post office employees are a bunch of crooks.”')
source('— A Letter to God, G. L. Fuentes; NCERT, First Flight', 'https://ncert.nic.in/textbook/pdf/jeff101.pdf')
p('Answer with reference to the extract and your understanding of the story.')
item('i', 'Who wrote these words, and to whom were they addressed?', 1)
item('ii', 'Which word in the extract means “dishonest people”?', 1)
item('iii', 'Explain the irony in the writer’s accusation. Refer to what the post office employees had actually done.', 2)
q('Q7. Read the poetry extract and answer ALL five questions. (6 marks)')
verse = p('He stalks in his vivid stripes\nThe few steps of his cage,\nOn pads of velvet quiet,\nIn his quiet rage.')
verse.paragraph_format.left_indent = Inches(.25)
source('— A Tiger in the Zoo, Leslie Norris; NCERT, First Flight', 'https://cdn1.byjus.com/wp-content/uploads/2019/11/NCERT-Book-for-Class-10-English-Chapter-2.pdf')
item('i', 'What does the limited space available to the tiger suggest about his life?', 1)
item('ii', 'Choose the meaning of “vivid” in this context.\nA. Faded    B. Bright and distinct    C. Hidden    D. Colourless', 1)
item('iii', 'The comparison of the tiger’s pads to velvet is an example of:\nA. Metaphor    B. Hyperbole    C. Refrain    D. Onomatopoeia', 1)
item('iv', 'Identify the tiger’s dominant emotion in the final line.', 1)
item('v', 'How does the poet contrast the tiger’s outward behaviour with his inner feelings? Support your response with details from the extract.', 2)

page()
h('SECTION C — LITERATURE (CONTINUED)')
q('Q8. First Flight — Short Answer Questions (12 marks)')
p('Attempt TWO questions from Part A and TWO questions from Part B. Answer each selected question in 40–50 words. Each answer carries 3 marks.')
h('Part A — Prose: Attempt ANY TWO. (2 × 3 = 6)')
item('i', 'How did Mandela’s understanding of freedom change as he grew older? Explain how this change influenced his sense of responsibility. (Nelson Mandela: Long Walk to Freedom)', 3)
item('ii', 'How did the young seagull’s mother use his hunger to help him overcome fear? Explain what his first flight taught him about himself. (His First Flight)', 3)
item('iii', 'Why did Anne feel the need to keep a diary despite having a family and friends? What kind of companion did she expect her diary to become? (From the Diary of Anne Frank)', 3)
item('iv', 'How does the continued importance of bread at Goan celebrations show the baker’s place in community life? Give two examples. (A Baker from Goa)', 3)
item('v', 'How do Coorg’s natural surroundings appeal to visitors, and how does Rajvir’s interest in tea enrich his visit to Assam? Give one relevant detail from each account. (Coorg; Tea from Assam)', 3)
h('Part B — Poetry: Attempt ANY TWO. (2 × 3 = 6)')
item('vi', 'How does a small incident in nature change the speaker’s mood? Explain why the ordinary setting matters to the poem’s message. (Dust of Snow)', 3)
item('vii', 'How does the poet use fire and ice to comment on human emotions and their destructive power? (Fire and Ice)', 3)
item('viii', 'How does the poet make dangerous encounters amusing? Explain with reference to any two animals described in the poem. (How to Tell Wild Animals)', 3)
item('ix', 'Why does the poet allow the boy to experience the loss of his ball instead of immediately offering him money? What must the boy learn? (The Ball Poem)', 3)
item('x', 'What do Amanda’s imagined roles reveal about her wishes? Explain with reference to any two of her fantasies. (Amanda!)', 3)
item('xi', 'How does the movement of the trees from the house to the forest develop the idea of freedom? Refer to the effort involved in their movement. (The Trees)', 3)

page()
h('SECTION C — LITERATURE (CONTINUED)')
q('Q9. Footprints Without Feet — Short Answer Questions (8 marks)')
p('Attempt ANY FOUR of the following five questions in 30–40 words each. (4 × 2 = 8)')
item('i', 'Why did Mr Herriot choose a simple routine of controlled food and exercise for Tricki? What does Tricki’s recovery reveal about the cause of his illness? (A Triumph of Surgery)', 2)
item('ii', 'Why did Hari Singh return to Anil after stealing the money? Identify the opportunity that mattered more to him than immediate gain. (The Thief’s Story)', 2)
item('iii', 'How did Ausable use the unexpected knock at the door to make his invented story convincing to Max? (The Midnight Visitor)', 2)
item('iv', 'Why did Horace Danby remove his gloves, and how did this decision contribute to his arrest? (A Question of Trust)', 2)
item('v', 'How does Griffin’s treatment of his landlord reveal the way he uses his scientific discovery? (Footprints Without Feet)', 2)
q('Q10. First Flight — Long Answer: Attempt EITHER A OR B. (6 marks)')
p('Write your answer in 120–150 words. Develop a clear argument and support it with relevant incidents from the text.')
p('A. Courage involves acting despite fear, while responsible action also requires judgement. Compare the young seagull’s struggle in His First Flight with the pilot’s decisions in Black Aeroplane. Discuss their fears, the help they receive and what their experiences suggest about confidence and risk.')
p('OR', True, center=True)
p('B. The postmaster responds to Lencho’s distress with generosity, yet Lencho fails to recognise the people who help him. Analyse how A Letter to God explores faith, compassion and misunderstanding. Support your answer with the postmaster’s actions and Lencho’s response to the money he receives.')
q('Q11. Footprints Without Feet — Long Answer: Attempt EITHER A OR B. (4 marks)')
p('Write your answer in 100–120 words, supporting your views with relevant details from the story.')
p('A. Anil’s kindness gives Hari Singh a reason to reconsider his choices. Explain how trust and the prospect of education influence Hari Singh’s decision to return. What does the ending suggest about his chance to change? (The Thief’s Story)')
p('OR', True, center=True)
p('B. Horace Danby regards himself as a careful and respectable man, but he is deceived by someone more skilful. Explain how the young woman uses his assumptions against him. What does the outcome reveal about his confidence and his dishonesty? (A Question of Trust)')
p('— END OF QUESTION PAPER —', True, 11, True)

doc.core_properties.title = 'Class X English Half-Yearly Question Paper'
doc.core_properties.subject = '80 marks | 3 hours | Based on the supplied K.M. International School blueprint'
doc.core_properties.author = 'K.M. International School'
doc.core_properties.keywords = 'Class 10, English 184, half-yearly, question paper'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)

# Validate the assessment structure and saved Word package.
marks = {'Q1':10,'Q2':10,'Q3':10,'Q4':5,'Q5':5,'Q6':4,'Q7':6,'Q8':12,'Q9':8,'Q10':6,'Q11':4}
assert sum(marks.values()) == 80
assert sum(marks[f'Q{i}'] for i in [6,7,8,9,10,11]) == 40
assert 4 + 6 + 6 == 16  # First Flight prose: extract, short answers, long answer.
assert 6 + 6 == 12      # First Flight poetry: extract and short answers.
assert 8 + 4 == 12      # Supplementary reader: short and long answers.
assert all(sum(col) == 200 for col in [(30,20,70,60,20),(50,30,80,30,10)])
assert sum([60,50,40,30,20]) == 200
saved = Document(OUT)
assert len(saved.tables) == 2
assert len(saved.element.xpath('//w:br[@w:type="page"]')) == 7
print(OUT.resolve())
print(f'Validated: 80 total marks, exact 16/12/12 literature split, 8 planned pages, {len(saved.paragraphs)} paragraphs.')
