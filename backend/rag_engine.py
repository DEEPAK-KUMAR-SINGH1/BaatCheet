import base64
import csv
import io
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.error
import urllib.request
from datetime import datetime
from typing import Generator

from config import GEMINI_API_KEY, load_env

load_env()

from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_mistralai import ChatMistralAI, MistralAIEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter

BASE_DIR = os.path.dirname(__file__)
CHROMA_BASE = os.path.join(BASE_DIR, "chroma_store")
UPLOADS_DIR = os.path.join(BASE_DIR, "uploads")

os.makedirs(CHROMA_BASE, exist_ok=True)
os.makedirs(UPLOADS_DIR, exist_ok=True)

logger = logging.getLogger(__name__)
GEMINI_OCR_MODEL = os.getenv("GEMINI_OCR_MODEL", "gemini-flash-latest")
_GEMINI_MODEL_CACHE: list[str] | None = None

llm = ChatMistralAI(model_name="mistral-large-2512", streaming=True)
embeddings = MistralAIEmbeddings(model="mistral-embed")

splitter = RecursiveCharacterTextSplitter(
    chunk_size=1000,
    chunk_overlap=150,
    separators=["\n\n", "\n", ".", "!", "?", ",", " "],
)


def _collection_name(workspace_id: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9_-]", "_", workspace_id)[:48]
    safe = safe.strip("_-") or "default"
    return f"thread_{safe}"


def get_vectorstore(workspace_id: str) -> Chroma:
    persist_dir = os.path.join(CHROMA_BASE, workspace_id)
    return Chroma(
        collection_name=_collection_name(workspace_id),
        embedding_function=embeddings,
        persist_directory=persist_dir,
    )


def extract_text_from_pdf_pages(file_path: str) -> list[dict]:
    import fitz

    pages = []
    pages_needing_ocr = []
    with fitz.open(file_path) as doc:
        for index, page in enumerate(doc, start=1):
            text = page.get_text().strip()
            if text:
                pages.append({"page": index, "text": text})
            else:
                pages_needing_ocr.append((index, page))

        if pages_needing_ocr and GEMINI_API_KEY:
            with tempfile.TemporaryDirectory() as tmpdir:
                for index, page in pages_needing_ocr:
                    image_path = os.path.join(tmpdir, f"page-{index}.png")
                    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    pix.save(image_path)
                    text = extract_text_from_image(image_path).strip()
                    if text:
                        pages.append({"page": index, "text": text})

    return sorted(pages, key=lambda page: page["page"])


def _normalize_gemini_model(model: str) -> str:
    return model.removeprefix("models/").strip()


def _gemini_model_candidates() -> list[str]:
    candidates = [
        GEMINI_OCR_MODEL,
        "gemini-flash-latest",
        "gemini-2.5-flash",
        "gemini-2.0-flash",
        "gemini-1.5-flash",
    ]
    seen = set()
    result = []
    for candidate in candidates:
        normalized = _normalize_gemini_model(candidate)
        if normalized and normalized not in seen:
            seen.add(normalized)
            result.append(normalized)
    return result


def _discover_gemini_models() -> list[str]:
    global _GEMINI_MODEL_CACHE
    if _GEMINI_MODEL_CACHE is not None:
        return _GEMINI_MODEL_CACHE

    try:
        url = f"https://generativelanguage.googleapis.com/v1beta/models?key={GEMINI_API_KEY}"
        with urllib.request.urlopen(url, timeout=30) as resp:
            payload = json.loads(resp.read().decode("utf-8"))
        models = []
        for model in payload.get("models", []):
            methods = model.get("supportedGenerationMethods") or []
            if "generateContent" in methods:
                models.append(_normalize_gemini_model(model.get("name", "")))
        preferred = [model for model in models if "flash" in model and "tts" not in model]
        _GEMINI_MODEL_CACHE = preferred or models
    except Exception as exc:
        logger.error("Gemini model discovery failed: %s", exc)
        _GEMINI_MODEL_CACHE = []
    return _GEMINI_MODEL_CACHE


def extract_text_from_image(file_path: str) -> str:
    if not GEMINI_API_KEY:
        raise ValueError("Image OCR requires GEMINI_API_KEY in your .env file.")

    try:
        with open(file_path, "rb") as img_file:
            image_data = base64.b64encode(img_file.read()).decode("utf-8")

        ext = os.path.splitext(file_path)[1].lower()
        mime_type = {
            ".jpg": "image/jpeg",
            ".jpeg": "image/jpeg",
            ".png": "image/png",
            ".webp": "image/webp",
        }.get(ext, "image/jpeg")

        payload = json.dumps({
            "contents": [
                {
                    "parts": [
                        {
                            "text": "Extract all text from this image. Return only the extracted text."
                        },
                        {
                            "inline_data": {
                                "mime_type": mime_type,
                                "data": image_data,
                            }
                        },
                    ]
                }
            ]
        }).encode("utf-8")

        last_error = None
        tried = set()

        def try_models(models: list[str]):
            nonlocal last_error
            for model in models:
                if model in tried:
                    continue
                tried.add(model)
                for api_version in ("v1beta", "v1"):
                    url = (
                        f"https://generativelanguage.googleapis.com/{api_version}/models/"
                        f"{model}:generateContent?key={GEMINI_API_KEY}"
                    )
                    req = urllib.request.Request(
                        url=url,
                        data=payload,
                        headers={"Content-Type": "application/json"},
                        method="POST",
                    )
                    try:
                        with urllib.request.urlopen(req, timeout=60) as resp:
                            result = json.loads(resp.read().decode("utf-8"))
                        parts = result.get("candidates", [{}])[0].get("content", {}).get("parts", [])
                        return "\n".join(part.get("text", "") for part in parts).strip()
                    except urllib.error.HTTPError as http_err:
                        last_error = (
                            f"Gemini HTTP {http_err.code} for {model}: "
                            f"{http_err.read().decode('utf-8', errors='ignore')}"
                        )
                        logger.error(last_error)
                    except Exception as exc:
                        last_error = str(exc)
                        logger.error("Gemini image extraction failed with %s: %s", model, exc)
            return None

        text = try_models(_gemini_model_candidates())
        if text is not None:
            return text
        text = try_models(_discover_gemini_models())
        if text is not None:
            return text

        raise ValueError(last_error or "Gemini image extraction failed.")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Image OCR failed: {exc}") from exc


def _read_text_file(file_path: str) -> str:
    with open(file_path, "rb") as handle:
        data = handle.read()

    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace").strip()


def _extract_csv(file_path: str) -> str:
    text = _read_text_file(file_path)
    if not text:
        return ""

    try:
        dialect = csv.Sniffer().sniff(text[:4096])
    except csv.Error:
        dialect = csv.excel

    rows = list(csv.reader(io.StringIO(text), dialect))
    if not rows:
        return ""

    headers = [cell.strip() for cell in rows[0]]
    has_headers = len(rows) > 1 and any(headers)
    lines = []
    if has_headers:
        lines.append("Columns: " + ", ".join(header or f"column_{idx + 1}" for idx, header in enumerate(headers)))
        for row_number, row in enumerate(rows[1:], start=1):
            pairs = []
            for idx, value in enumerate(row):
                header = headers[idx] if idx < len(headers) and headers[idx] else f"column_{idx + 1}"
                if value.strip():
                    pairs.append(f"{header}={value.strip()}")
            if pairs:
                lines.append(f"Row {row_number}: " + "; ".join(pairs))
    else:
        lines = [", ".join(cell.strip() for cell in row if cell.strip()) for row in rows]
    return "\n".join(line for line in lines if line.strip()).strip()


def _extract_docx(file_path: str) -> str:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise ValueError("DOCX support requires python-docx") from exc

    doc = DocxDocument(file_path)
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _run_text_command(command: list[str], output_path: str | None = None) -> str:
    result = subprocess.run(
        command,
        capture_output=True,
        text=True,
        timeout=60,
        check=False,
    )
    if result.returncode != 0:
        raise ValueError((result.stderr or result.stdout or "command failed").strip())
    if output_path:
        return _read_text_file(output_path)
    return (result.stdout or "").strip()


def _extract_doc_with_soffice(file_path: str) -> str:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        raise ValueError("LibreOffice/soffice not found")

    with tempfile.TemporaryDirectory() as tmpdir:
        _run_text_command([
            executable,
            "--headless",
            "--convert-to",
            "txt:Text",
            "--outdir",
            tmpdir,
            file_path,
        ])
        output = os.path.join(
            tmpdir,
            os.path.splitext(os.path.basename(file_path))[0] + ".txt",
        )
        if not os.path.exists(output):
            raise ValueError("LibreOffice did not create a text file")
        return _read_text_file(output)


def _extract_doc_with_antiword(file_path: str) -> str:
    executable = shutil.which("antiword") or shutil.which("catdoc")
    if not executable:
        raise ValueError("antiword/catdoc not found")
    return _run_text_command([executable, file_path])


def _extract_doc_with_word(file_path: str) -> str:
    if sys.platform != "win32":
        raise ValueError("Microsoft Word automation is only available on Windows")
    try:
        import pythoncom
        import win32com.client
    except ImportError as exc:
        raise ValueError("Microsoft Word automation requires pywin32") from exc

    with tempfile.TemporaryDirectory() as tmpdir:
        output = os.path.join(tmpdir, "document.txt")
        pythoncom.CoInitialize()
        word = None
        doc = None
        try:
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            doc = word.Documents.Open(os.path.abspath(file_path), ReadOnly=True)
            doc.SaveAs2(os.path.abspath(output), FileFormat=7)
        finally:
            if doc is not None:
                doc.Close(False)
            if word is not None:
                word.Quit()
            pythoncom.CoUninitialize()
        return _read_text_file(output)


def _clean_binary_text(parts: list[str]) -> str:
    lines = []
    seen = set()
    for part in parts:
        for line in part.splitlines():
            line = re.sub(r"[\x00-\x1f\x7f-\x9f]", " ", line)
            line = re.sub(r"\s+", " ", line).strip()
            if len(line) < 4:
                continue
            if sum(ch.isalnum() for ch in line) < max(3, len(line) // 4):
                continue
            if line not in seen:
                seen.add(line)
                lines.append(line)
    return "\n".join(lines).strip()


def _extract_doc_from_ole(file_path: str) -> str:
    try:
        import olefile
    except ImportError as exc:
        raise ValueError("OLE fallback requires olefile") from exc

    if not olefile.isOleFile(file_path):
        raise ValueError("Not an OLE .doc file")

    parts = []
    with olefile.OleFileIO(file_path) as ole:
        for stream_name in ("WordDocument", "1Table", "0Table", "Data"):
            if not ole.exists(stream_name):
                continue
            data = ole.openstream(stream_name).read()
            parts.append(data.decode("utf-16le", errors="ignore"))
            parts.append(data.decode("latin-1", errors="ignore"))
    return _clean_binary_text(parts)


def _extract_doc(file_path: str) -> str:
    errors = []
    for extractor in (
        _extract_doc_with_soffice,
        _extract_doc_with_antiword,
        _extract_doc_with_word,
        _extract_doc_from_ole,
    ):
        try:
            text = extractor(file_path).strip()
            if text:
                return text
        except Exception as exc:
            errors.append(str(exc))

    detail = "; ".join(error for error in errors if error)
    raise ValueError(
        "Could not extract text from this legacy .doc file. Install LibreOffice, "
        "antiword/catdoc, or Microsoft Word with pywin32, or convert the file to DOCX. "
        f"Details: {detail}"
    )



def _extract_xlsx(file_path: str) -> str:
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError("XLSX support requires openpyxl") from exc

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets:
        rows.append(f"Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            values = [str(value) for value in row if value is not None]
            if values:
                rows.append(", ".join(values))
    return "\n".join(rows).strip()


def _extract_pptx(file_path: str) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ValueError("PPTX support requires python-pptx") from exc

    prs = Presentation(file_path)
    slides = []
    for idx, slide in enumerate(prs.slides, start=1):
        text = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text.append(shape.text.strip())
        if text:
            slides.append(f"Slide {idx}\n" + "\n".join(text))
    return "\n\n".join(slides).strip()


def extract_document_pages(file_path: str, filename: str) -> list[dict]:
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf_pages(file_path)
    if ext in {".jpg", ".jpeg", ".png", ".webp"}:
        text = extract_text_from_image(file_path)
        return [{"page": 1, "text": text}] if text.strip() else []
    if ext in {".txt", ".md"}:
        text = _read_text_file(file_path)
        return [{"page": 1, "text": text}] if text else []
    if ext == ".csv":
        text = _extract_csv(file_path)
        return [{"page": 1, "text": text}] if text else []
    if ext == ".doc":
        text = _extract_doc(file_path)
        return [{"page": 1, "text": text}] if text else []
    if ext == ".docx":
        text = _extract_docx(file_path)
        return [{"page": 1, "text": text}] if text else []
    if ext == ".xlsx":
        text = _extract_xlsx(file_path)
        return [{"page": 1, "text": text}] if text else []
    if ext == ".pptx":
        text = _extract_pptx(file_path)
        return [{"page": 1, "text": text}] if text else []
    raise ValueError(f"Unsupported file type: {ext}")


def extract_text(file_path: str, filename: str) -> str:
    return "\n\n".join(page["text"] for page in extract_document_pages(file_path, filename)).strip()


def _chunk_pages(pages_or_text) -> list[dict]:
    if isinstance(pages_or_text, str):
        pages = [{"page": 1, "text": pages_or_text}]
    else:
        pages = pages_or_text

    chunks = []
    chunk_index = 0
    for page in pages:
        for chunk in splitter.split_text(page.get("text") or ""):
            clean = chunk.strip()
            if not clean:
                continue
            chunks.append({
                "chunk_index": chunk_index,
                "page": page.get("page"),
                "content": clean,
            })
            chunk_index += 1
    return chunks


def add_document_to_store(
    workspace_id: str,
    pages_or_text,
    doc_name: str,
    doc_id: str,
    return_chunks: bool = False,
):
    chunks = _chunk_pages(pages_or_text)
    docs = [
        Document(
            page_content=chunk["content"],
            metadata={
                "doc_id": doc_id,
                "doc_name": doc_name,
                "chunk_index": chunk["chunk_index"],
                "chunk": chunk["chunk_index"],
                "page": chunk.get("page"),
                "workspace_id": workspace_id,
                "thread_id": workspace_id,
            },
        )
        for chunk in chunks
    ]
    if docs:
        vs = get_vectorstore(workspace_id)
        vs.add_documents(docs, ids=[f"{doc_id}:{chunk['chunk_index']}" for chunk in chunks])
    if return_chunks:
        return len(chunks), chunks
    return len(chunks)


def _score_value(score):
    if score is None:
        return None
    try:
        return float(score)
    except Exception:
        return None


def _source_from_chunk(row: dict, label: str, score: float | None = None) -> dict:
    return {
        "doc_id": row.get("doc_id"),
        "filename": row.get("filename", "document"),
        "page": row.get("page"),
        "chunk_index": row.get("chunk_index"),
        "score": score,
        "snippet": (row.get("content") or "")[:500].strip(),
        "label": label,
    }


def _context_from_sources(sources: list[dict], content_by_key: dict[tuple[str, int], str]):
    context_parts = []
    for source in sources:
        page_label = f", page {source['page']}" if source.get("page") else ""
        key = (source.get("doc_id"), source.get("chunk_index"))
        context_parts.append(
            f"[{source['label']}: {source['filename']}{page_label}, chunk {source.get('chunk_index')}]\n"
            f"{content_by_key.get(key, source.get('snippet', ''))}"
        )
    return "\n\n---\n\n".join(context_parts), sources


def _fallback_chunk_search(workspace_id: str, query: str, k: int = 5):
    from db import execute

    rows = execute(
        """
        SELECT dc.doc_id, dc.chunk_index, dc.page, dc.content, d.filename
        FROM document_chunks dc
        JOIN documents d ON d.doc_id = dc.doc_id
        WHERE d.workspace_id=%s AND d.status='indexed'
        ORDER BY d.uploaded_at ASC, dc.chunk_index ASC
        """,
        (workspace_id,),
        fetch="all",
    )
    if not rows:
        return "", []

    terms = [
        term for term in re.findall(r"[A-Za-z0-9_]{3,}", query.lower())
        if term not in {"the", "and", "for", "with", "from", "this", "that"}
    ][:12]

    scored = []
    for original_index, row in enumerate(rows):
        content = (row.get("content") or "").lower()
        score = 0
        if query.strip() and query.lower() in content:
            score += 20
        score += sum(content.count(term) for term in terms)
        scored.append((score, -original_index, row))

    scored.sort(key=lambda item: (item[0], item[1]), reverse=True)
    best_rows = [
        row for score, _, row in scored[:k]
        if score > 0 or not terms
    ]
    if not best_rows:
        best_rows = [row for _, _, row in scored[:k]]

    sources = []
    content_by_key = {}
    for row in best_rows:
        label = f"S{len(sources) + 1}"
        normalized_score = None
        if terms:
            normalized_score = min(1.0, sum((row.get("content") or "").lower().count(term) for term in terms) / len(terms))
        source = _source_from_chunk(row, label, normalized_score)
        sources.append(source)
        content_by_key[(row.get("doc_id"), row.get("chunk_index"))] = row.get("content") or ""
    return _context_from_sources(sources, content_by_key)


def retrieve_context_with_sources(workspace_id: str, query: str, k: int = 5):
    try:
        vs = get_vectorstore(workspace_id)
    except Exception:
        return _fallback_chunk_search(workspace_id, query, k=k)

    try:
        results = vs.similarity_search_with_relevance_scores(query, k=k)
    except Exception:
        try:
            results = vs.similarity_search_with_score(query, k=k)
        except Exception:
            return _fallback_chunk_search(workspace_id, query, k=k)

    context_parts = []
    sources = []
    content_by_key = {}
    seen = set()
    for idx, item in enumerate(results, start=1):
        if isinstance(item, tuple):
            doc, score = item
        else:
            doc, score = item, None
        meta = doc.metadata or {}
        doc_id = meta.get("doc_id")
        chunk_index = meta.get("chunk_index", meta.get("chunk"))
        key = (doc_id, chunk_index)
        if key in seen:
            continue
        seen.add(key)
        filename = meta.get("doc_name", "document")
        page = meta.get("page")
        snippet = doc.page_content[:500].strip()
        source = {
            "doc_id": doc_id,
            "filename": filename,
            "page": page,
            "chunk_index": chunk_index,
            "score": _score_value(score),
            "snippet": snippet,
            "label": f"S{len(sources) + 1}",
        }
        sources.append(source)
        content_by_key[(doc_id, chunk_index)] = doc.page_content
        page_label = f", page {page}" if page else ""
        context_parts.append(
            f"[{source['label']}: {filename}{page_label}, chunk {chunk_index}]\n{doc.page_content}"
        )
    if sources:
        return "\n\n---\n\n".join(context_parts), sources
    return _fallback_chunk_search(workspace_id, query, k=k)


def retrieve_context(workspace_id: str, query: str, k: int = 5) -> str:
    context, _ = retrieve_context_with_sources(workspace_id, query, k=k)
    return context


def delete_thread_store(workspace_id: str):
    persist_dir = os.path.join(CHROMA_BASE, workspace_id)
    if os.path.exists(persist_dir):
        shutil.rmtree(persist_dir)


def delete_doc_from_store(workspace_id: str, doc_id: str):
    vs = get_vectorstore(workspace_id)
    try:
        vs.delete(where={"doc_id": doc_id})
    except Exception:
        pass


RAG_SYSTEM = """You are a highly intelligent AI assistant. Today's date is {date}.

You have access to uploaded workspace documents. Use them when relevant, cite document names, and say clearly when a document answer is not present in the provided context.
"""


def stream_rag_response(
    thread_id: str,
    user_message: str,
    chat_history: list,
    workspace_id: str | None = None,
) -> Generator[str, None, None]:
    from engine import stream_response

    workspace_id = workspace_id or thread_id
    context = retrieve_context(workspace_id, user_message, k=5)
    if context:
        context = RAG_SYSTEM.format(date=datetime.now().strftime("%d %B %Y, %A")) + "\n\n" + context
    yield from stream_response(thread_id, user_message, document_context=context)
