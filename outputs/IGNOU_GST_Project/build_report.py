from pathlib import Path
import re, json, math
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor, white
import fitz

ROOT=Path(__file__).parent
ASSETS=ROOT/'assets'; ASSETS.mkdir(exist_ok=True)
BLUE='#214F6B'; TEAL='#348A88'; GOLD='#C89E52'

def graphic(name,kind,labels,values=None):
    pdf=ASSETS/(name+'.pdf'); c=canvas.Canvas(str(pdf),pagesize=(620,260))
    c.setFillColor(white);c.rect(0,0,620,260,fill=1,stroke=0)
    if kind=='bar':
        maximum=max(values)*1.15
        n=len(values); bw=480/n
        c.setStrokeColor(HexColor('#ccd5db'));c.line(65,40,590,40)
        for i,(lab,val) in enumerate(zip(labels,values)):
            x=80+i*bw; h=170*val/maximum
            c.setFillColor(HexColor(TEAL if i%2 else BLUE));c.rect(x,40,bw*.58,h,fill=1,stroke=0)
            c.setFont('Helvetica-Bold',11);c.drawCentredString(x+bw*.29,48+h,f'{val:g}')
            c.setFillColor(HexColor('#263746'));c.setFont('Helvetica',10);c.drawCentredString(x+bw*.29,22,lab)
    elif kind=='split':
        boxes=[(215,175,190,60,labels[0]),(65,55,210,75,labels[1]),(345,55,210,75,labels[2])]
        c.setStrokeColor(HexColor(GOLD));c.setLineWidth(2)
        c.line(310,175,170,130);c.line(310,175,450,130)
        for x,y,w,h,lab in boxes:
            c.setFillColor(HexColor(BLUE if y>150 else TEAL));c.roundRect(x,y,w,h,8,fill=1,stroke=0)
            c.setFillColor(white);c.setFont('Helvetica-Bold',12)
            for j,l in enumerate(lab.split('\n')):c.drawCentredString(x+w/2,y+h-23-j*17,l)
        c.setFillColor(HexColor('#455765'));c.setFont('Helvetica',10);c.drawCentredString(310,25,'Illustrates ordinary domestic supplies; specific legal conditions and exceptions apply.')
    elif kind=='flow':
        n=len(labels); w=min(145,540/n-12); y=100
        for i,lab in enumerate(labels):
            x=20+i*(580/n)
            c.setFillColor(HexColor(BLUE if i%2==0 else TEAL));c.roundRect(x,y,w,95,8,fill=1,stroke=0)
            c.setFillColor(white);c.setFont('Helvetica-Bold',11)
            lines=lab.split('\n')
            for j,l in enumerate(lines):c.drawCentredString(x+w/2,y+60-j*17,l)
            if i<n-1:
                c.setStrokeColor(HexColor(GOLD));c.setLineWidth(3);c.line(x+w+3,y+47,x+580/n-5,y+47)
                c.line(x+580/n-12,y+54,x+580/n-5,y+47);c.line(x+580/n-12,y+40,x+580/n-5,y+47)
        c.setFont('Helvetica',11);c.setFillColor(HexColor('#455765'));c.drawCentredString(310,55,'Analytical framework; arrows indicate a possible mechanism, not a measured effect.')
    elif kind=='workshop':
        for x,label in [(20,'PRODUCTION'),(225,'RECORDS'),(430,'MARKET')]:
            c.setFillColor(HexColor('#eef3f5'));c.roundRect(x,30,175,210,8,fill=1,stroke=0)
            c.setFillColor(HexColor(BLUE));c.setFont('Helvetica-Bold',12);c.drawCentredString(x+87,49,label)
        c.setFillColor(HexColor(TEAL));c.rect(40,95,132,80,fill=1,stroke=0)
        path=c.beginPath();path.moveTo(40,175);path.lineTo(75,205);path.lineTo(75,175);path.lineTo(110,205);path.lineTo(110,175);path.lineTo(145,205);path.lineTo(172,175);path.close();c.drawPath(path,fill=1,stroke=0)
        c.setFillColor(white)
        for x in [55,95,135]:c.rect(x,130,20,23,fill=1,stroke=0)
        c.setFillColor(HexColor(BLUE));c.roundRect(247,115,132,84,4,fill=1,stroke=0);c.rect(302,92,22,23,fill=1,stroke=0);c.rect(270,86,85,8,fill=1,stroke=0)
        c.setFillColor(white);c.rect(258,127,110,60,fill=1,stroke=0)
        c.setFillColor(HexColor(TEAL));c.setFont('Helvetica-Bold',13);c.drawCentredString(313,166,'GST');c.setFont('Helvetica',10);c.drawCentredString(313,146,'Invoice / Return')
        c.setFillColor(HexColor(TEAL));c.rect(455,100,125,70,fill=1,stroke=0)
        c.setFillColor(HexColor(GOLD));c.rect(447,170,141,24,fill=1,stroke=0)
        c.setFillColor(white);c.rect(468,112,35,40,fill=1,stroke=0);c.rect(522,100,32,52,fill=1,stroke=0)
    c.save();d=fitz.open(pdf);d[0].get_pixmap(matrix=fitz.Matrix(2,2),alpha=False).save(ASSETS/(name+'.png'));d.close()

graphic('enterprise','workshop',[])
graphic('gst_structure','split',['Taxable supply','Within state\nCGST + SGST','Interstate\nIGST'])
graphic('credit_chain','flow',['Input supplier\nTax invoice','SME buyer\nEligible credit','Output supply\nNet payment'])
graphic('framework','flow',['GST design\nRules + portal','Firm capability\nPeople + records','Business effects\nCost + liquidity'])
graphic('research_flow','flow',['Select\nsources','Check period\nand population','Compare\nevidence','Interpret\nwith limits'])
graphic('collections','bar',['2023-24','2024-25'],[20.18249,22.08861])
graphic('gva','bar',['2017-18','2018-19','2019-20','2020-21','2021-22','2022-23'],[29.7,30.5,30.5,27.3,29.6,30.1])
graphic('exports','bar',['2019-20','2020-21','2021-22','2022-23','2023-24'],[49.75,49.35,45.03,43.59,45.73])
graphic('fixedcost','bar',['Rs 20 lakh','Rs 50 lakh','Rs 1 crore','Rs 2 crore'],[2.4,.96,.48,.24])
graphic('cashflow','flow',['Purchase\nInput tax paid','Sale on credit\nReceivable opens','Tax due\nCash gap possible','Buyer pays\nGap closes'])
graphic('sensitivity','bar',['15 days','30 days','60 days','90 days'],[986,1973,3945,5918])
graphic('sentiment','bar',['2022','2023','2024','2025'],[59,72,84,85])
graphic('control_cycle','flow',['Capture\ninvoice','Reconcile\nexceptions','Review\neligibility','File + retain\nevidence'])

def shade(cell,color):
    tcPr=cell._tc.get_or_add_tcPr(); s=OxmlElement('w:shd');s.set(qn('w:fill'),color);tcPr.append(s)

def textp(doc,text,style=None):
    p=doc.add_paragraph(text,style)
    if style is None:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

def hyperlink(doc,url):
    p=doc.add_paragraph(style='Source');p.paragraph_format.space_after=Pt(3)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'245E80');pr.append(co)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'20');pr.append(sz);r.append(pr)
    t=OxmlElement('w:t');t.text=url;r.append(t);h.append(r);p._p.append(h)

def table(doc,rows,compact=False):
    tab=doc.add_table(rows=0,cols=len(rows[0]));tab.style='Table Grid'
    widths=None
    if compact:
        widths=[5.5,.77] if len(rows[0])==2 else [.9,4.75,.62]
        tab.autofit=False
        for col,w in zip(tab.columns,widths):col.width=Inches(w)
    for i,row in enumerate(rows):
        cells=tab.add_row().cells
        for cell,value in zip(cells,row):
            p=cell.paragraphs[0];p.paragraph_format.line_spacing=1.0 if compact else 1.12;p.paragraph_format.space_after=Pt(2 if compact else 5);p.paragraph_format.space_before=Pt(2 if compact else 5)
            run=p.add_run(value);run.font.size=Pt(10)
            if i==0:shade(cell,'214F6B');run.font.bold=True;run.font.color.rgb=RGBColor(255,255,255)
        if widths:
            for cell,w in zip(cells,widths):cell.width=Inches(w)
        trPr=tab.rows[-1]._tr.get_or_add_trPr();el=OxmlElement('w:cantSplit');trPr.append(el)
        if i==0:rep=OxmlElement('w:tblHeader');trPr.append(rep)
    spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(0);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=Pt(1)
    spacer.add_run().font.size=Pt(1)

def build():
    chunks=[]
    for f in sorted(ROOT.glob('content_*.md')):chunks.extend(f.read_text(encoding='utf-8').strip().split('\n===\n'))
    doc=Document();sec=doc.sections[0];sec.page_height=Inches(11.69);sec.page_width=Inches(8.27)
    sec.top_margin=sec.bottom_margin=Inches(.85);sec.left_margin=Inches(1.1);sec.right_margin=Inches(.9)
    sec.header_distance=sec.footer_distance=Inches(.4)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
    normal.paragraph_format.line_spacing=2;normal.paragraph_format.space_after=Pt(6)
    for nm,size in [('Title',22),('Heading 1',17),('Heading 2',14),('Heading 3',12)]:
        s=doc.styles[nm];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string('214F6B');s.paragraph_format.line_spacing=1.2;s.paragraph_format.space_after=Pt(12)
    for nm in ['Caption','Source']:
        if nm not in doc.styles:doc.styles.add_style(nm,1)
        s=doc.styles[nm];s.font.name='Times New Roman';s.font.size=Pt(10);s.paragraph_format.line_spacing=1.1;s.paragraph_format.space_after=Pt(6)
    hdr=sec.header.paragraphs[0];hdr.text='MCOP-001  |  GST AND SMALL & MEDIUM ENTERPRISES';hdr.style='Source';hdr.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)
    manifest=[]
    for ix,chunk in enumerate(chunks):
        start_index=len(doc._element.body)
        lines=chunk.strip().splitlines();title=lines[0].lstrip('# ')
        manifest.append({'page':ix+1,'title':title,'words':len(chunk.split())})
        i=0
        while i<len(lines):
            line=lines[i].strip()
            if not line:i+=1;continue
            if line.startswith('|'):
                rows=[]
                while i<len(lines) and lines[i].strip().startswith('|'):
                    rows.append([s.strip() for s in lines[i].strip().strip('|').split('|')]);i+=1
                table(doc,rows,compact=ix in [4,5,6,7]);continue
            if line.startswith('!FIG '):
                name,caption,source=line[5:].split('|',2)
                p=doc.add_paragraph();p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(3);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(ASSETS/(name+'.png')),width=Inches(5.95))
                textp(doc,caption,'Caption');textp(doc,source,'Source')
            elif line.startswith('!SPACE'):doc.add_paragraph().paragraph_format.space_after=Pt(float(line.split()[1]))
            elif line.startswith('!CENTER '):
                p=textp(doc,line[8:]);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1.2;p.paragraph_format.space_after=Pt(8)
            elif line.startswith('!TITLE '):
                p=textp(doc,line[7:],'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            elif line.startswith('### '):textp(doc,line[4:],'Heading 3')
            elif line.startswith('## '):textp(doc,line[3:],'Heading 2')
            elif line=='# Project report':pass
            elif line.startswith('# '):textp(doc,line[2:],'Heading 1')
            elif line.startswith('~ '):textp(doc,line[2:],'Source')
            elif line.startswith('URL: '):hyperlink(doc,line[5:])
            else:textp(doc,line)
            i+=1
        new_paras=[e for e in list(doc._element.body)[max(0,start_index-1):] if e.tag==qn('w:p')]
        if new_paras:
            if ix:
                ppr=new_paras[0].find(qn('w:pPr'))
                if ppr is None:ppr=OxmlElement('w:pPr');new_paras[0].insert(0,ppr)
                pb=OxmlElement('w:pageBreakBefore');ppr.append(pb)
            mark=OxmlElement('w:bookmarkStart');mark.set(qn('w:id'),str(ix+1));mark.set(qn('w:name'),f'page_{ix+1:03d}')
            end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(ix+1))
            new_paras[0].insert(1 if new_paras[0].find(qn('w:pPr')) is not None else 0,mark);new_paras[0].append(end)
    doc.core_properties.title='Impact of Goods and Services Tax (GST) on Small and Medium Enterprises in India'
    doc.core_properties.subject='M.Com MCOP-001 secondary-data project report'
    doc.core_properties.author='Kajal'
    doc.core_properties.keywords='GST, SMEs, IGNOU, MCOP-001, secondary research'
    dest=ROOT/'Kajal_IGNOU_MCom_GST_Project.docx';doc.save(dest)
    (ROOT/'page_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
    print(json.dumps({'file':str(dest),'planned_pages':len(chunks),'words':sum(x['words'] for x in manifest)},indent=2))

if __name__=='__main__':build()
