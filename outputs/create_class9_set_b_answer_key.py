from pathlib import Path
import sys
import zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).parent
SOURCE = ROOT / 'Class_9_English_Half_Yearly_Question_Paper_Set_B.docx'
OUT = ROOT / 'Class_9_English_Half_Yearly_Set_B_Answer_Key_and_Marking_Scheme.docx'
paper = Document(SOURCE)
source_text = '\n'.join(p.text for p in paper.paragraphs)
assert 'SET B' in source_text and 'Do any 10 out of 12' in source_text
assert 'Do any 4 out of 5' in source_text and 'birthday surprise' in source_text

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.7)
sec.header_distance = sec.footer_distance = Inches(.28)
s = doc.styles['Normal']
s.font.name = 'Times New Roman'
s.font.size = Pt(11)
s.paragraph_format.space_after = Pt(6)
s.paragraph_format.line_spacing = 1.03
for name in ['Heading 1', 'Heading 2']:
    s = doc.styles[name]
    s.font.name = 'Times New Roman'
    s.font.size = Pt(12)
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(9)
    s.paragraph_format.space_after = Pt(7)

def p(text='', bold=False, size=None, center=False):
    x = doc.add_paragraph()
    r = x.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    if center:
        x.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return x

def h(text):
    doc.add_paragraph(text, 'Heading 2')

def page():
    doc.add_page_break()

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.autofit = False
    for c, text in zip(t.rows[0].cells, headers):
        c.text = str(text)
        for r in c.paragraphs[0].runs:
            r.bold = True
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for values in rows:
        for c, text in zip(t.add_row().cells, values):
            c.text = str(text)
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for c, width in zip(row.cells, widths):
            c.width = Inches(width)
            for x in c.paragraphs:
                x.paragraph_format.space_before = Pt(3)
                x.paragraph_format.space_after = Pt(3)
                for r in x.runs:
                    r.font.size = Pt(10.5)
    p()

def link(label, url):
    x = doc.add_paragraph()
    a = OxmlElement('w:hyperlink')
    a.set(qn('r:id'), x.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r')
    rp = OxmlElement('w:rPr')
    size = OxmlElement('w:sz')
    size.set(qn('w:val'), '18')
    rp.append(size)
    r.append(rp)
    t = OxmlElement('w:t')
    t.text = label
    r.append(t)
    a.append(r)
    x._p.append(a)

hp = sec.header.paragraphs[0]
hp.text = 'K.M. INTERNATIONAL SCHOOL | CLASS IX ENGLISH | SET B | TEACHER COPY'
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
hp.runs[0].font.size = Pt(9)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.add_run('Set B: Answer Key | Page ')
for field, suffix in [('PAGE', ' of '), ('NUMPAGES', '')]:
    f = OxmlElement('w:fldSimple')
    f.set(qn('w:instr'), field)
    fp._p.append(f)
    fp.add_run(suffix)

p('K.M. INTERNATIONAL SCHOOL', True, 15, True)
p('CLASS IX — HALF-YEARLY EXAMINATION — SET B', True, 13, True)
p('ENGLISH LANGUAGE & LITERATURE', True, 12, True)
p('ANSWER KEY AND MARKING SCHEME', True, 13, True)
p('Maximum Marks: 80 | Time Allowed: 3 Hours', True, 11, True)
p('For Class_9_English_Half_Yearly_Question_Paper_Set_B.docx', size=9, center=True)
h('Mark allocation and choices')
allocation = [10, 10, 10, 5, 5, 5, 5, 12, 8, 5, 5]
table(['Question', 'Requirement', 'Marks'], [
    ['Q1', 'Observation passage: all 10 parts', 10],
    ['Q2', 'Waste audit: all 10 parts', 10],
    ['Q3', 'Grammar: any 10 of 12', 10],
    ['Q4', 'Letter to the Editor OR Complaint Letter', 5],
    ['Q5', 'Diary Entry OR Story', 5],
    ['Q6', 'Prose extract: 1 + 1 + 1 + 2', 5],
    ['Q7', 'Poetry extract: 1 + 1 + 3', 5],
    ['Q8', 'Any 4 of 6; 3 marks each', 12],
    ['Q9', 'Any 4 of 5; 2 marks each', 8],
    ['Q10', 'Either A or B; 100–120 words', 5],
    ['Q11', 'Either A or B; 100–120 words', 5],
    ['TOTAL', 'Reading 20 + Grammar/Writing 20 + Literature 40', 80]
], [.7, 5.4, .75])
h('Guidance for examiners')
for text in [
    'This scheme is prepared for the supplied school paper; it is not an official CBSE marking scheme. Accept accurate paraphrases and reasonable interpretations supported by the text.',
    'Objective items earn 1 or 0. A correct option letter or clearly written answer is sufficient. If the option and written answer conflict, award 0. Award half marks for partially developed written content where specified or appropriate.',
    'Suggested over-attempt rule: assess the best permitted number, unless the school announced another rule beforehand. Apply the same rule consistently. Q8 has no separate prose/poetry quota.',
    'Do not deduct twice for the same error. Minor language slips should not reduce comprehension marks when meaning is clear. Assess expression separately where allocated.',
    'Word limits indicate expected focus and development; do not add an automatic penalty. Model responses illustrate acceptable answers, not compulsory wording.'
]:
    p(text, size=10)

page()
h('SECTION A — READING SKILLS (20 MARKS)')
h('Q1. Observation — All ten parts (10 × 1 = 10)')
q1 = [
    ['(i)', 'C — Careful attention and questioning deepen understanding.'],
    ['(ii)', 'The changing colour of the tree’s leaves, compared with the previous week.'],
    ['(iii)', 'Those students examined the actual shoe closely and noticed its uneven laces, worn heel or tear, instead of drawing a familiar shape from memory.'],
    ['(iv)', 'Expectations can make us see what we assume is there, overlooking the actual object’s particular details.'],
    ['(v)', 'Any one: comparing a character’s actions with their words while reading; noticing a colour change in a science activity; watching a ball’s bounce to choose a better receiving position.'],
    ['(vi)', 'False. A noticed detail still needs interpretation.'],
    ['(vii)', 'A — The corridor floor is wet.'],
    ['(viii)', 'Revise.'],
    ['(ix)', 'Selecting details relevant to a question keeps attention focused on useful evidence and helps answer that question. Recording everything may include irrelevant information.'],
    ['(x)', 'Respect for other people’s privacy; curiosity does not justify accessing their private messages.']
]
table(['Part', 'Expected answer'], q1, [.6, 6.25])
p('Award 1 per correct response. For Q1(iii), close observation rather than reliance on memory is the key distinction; do not require every feature of the shoe. For Q1(v), one relevant example is sufficient. Q1(vi) earns full credit for “False” alone.', size=10)
p('For explanatory items, award ½ for a relevant but incomplete idea. Answers need not copy the passage word for word.', size=10)

page()
h('Q2. School waste audit — All ten parts (10 × 1 = 10)')
q2 = [
    ['(i)', 'B — Waste sent from the lunch area for disposal.'],
    ['(ii)', 'Food scraps: 60 kg in July and 42 kg in August. Naming the category is sufficient.'],
    ['(iii)', '60 kg: 180 − 120 = 60.'],
    ['(iv)', '35%: (42 ÷ 120) × 100 = 35%.'],
    ['(v)', '50%: [(36 − 18) ÷ 36] × 100 = 50%.'],
    ['(vi)', 'Other waste (½), by 6 kg (½): 30 − 24 = 6.'],
    ['(vii)', 'False. Material sent for composting was excluded.'],
    ['(viii)', 'Any one: same number of students; same weighing scale; same measurement time each afternoon; equal five-day observation periods.'],
    ['(ix)', 'Two weeks cannot show whether the improvement is permanent; repeated audits would test whether reduced disposal continues over time.'],
    ['(x)', 'B — Separated into groups.']
]
table(['Part', 'Expected answer'], q2, [.6, 6.25])
h('Calculation and interpretation notes')
p('Correct numerical answers earn full credit without working. In Q2(iv) and (v), award ½ for a correct setup and ½ for the correct result if working is shown. For Q2(v), use July’s plastic waste (36 kg) as the base; the question asks for a percentage decrease, not a decrease in kilograms.', size=10)
p('The table records waste sent for disposal. It does not measure all material used or prove that all waste generation fell. Q2(viii) requires only one controlled condition.', size=10)

page()
h('SECTION B — GRAMMAR & WRITING SKILLS (20 MARKS)')
h('Q3. Grammar — Any ten of twelve (10 marks)')
q3 = [
    ['(i)', 'A — has lived', 1],
    ['(ii)', 'was reading', 1],
    ['(iii)', 'C — need not', 1],
    ['(iv)', 'B — is', 1],
    ['(v)', 'D — are', 1],
    ['(vi)', 'B — a', 1],
    ['(vii)', 'C — many', 1],
    ['(viii)', 'D — in', 1],
    ['(ix)', 'went → go: We did not go to the market yesterday.', 1],
    ['(x)', 'sweet → sweetly: The child sang sweetly at the annual function.', 1],
    ['(xi)', 'C — will be travelling', 1],
    ['(xii)', 'D — any', 1]
]
table(['Part', 'Answer / correction', 'Marks'], q3, [.6, 5.6, .65])
p('Award 1 per correct answer and 0 otherwise. For editing, accept the incorrect word with its correction or the complete corrected sentence. Merely identifying an error without correcting it earns 0. Maximum: 10 marks, even if all twelve responses are correct.')
h('Grammar notes')
p('The residence began in the past and continues now, requiring the present perfect. “Neither” takes a singular verb here; “shoes” takes a plural verb. “University” begins with a consonant sound, so “a” is correct. After “did not”, use the base verb. “Sweetly” is the adverb modifying “sang”.', size=10)

page()
h('Q4. Formal letter — Either A or B (5 marks)')
p('Expected length: 100–120 words. Apply the following common rubric to either option.')
table(['Criterion', 'Marks and basis'], [
    ['Format', '1: appropriate addresses, date, subject, salutation and formal closing. Award ½ for a recognisable but incomplete format.'],
    ['Content', '2: use the option-specific allocation below or on the next page.'],
    ['Organisation and tone', '1: clear purpose, logical development and suitable formal tone.'],
    ['Accuracy', '1: appropriate grammar, vocabulary, spelling and punctuation; ½ for noticeable errors that do not obscure meaning.']
], [1.55, 5.3])
h('Q4(A). Letter to the Editor — Model')
p('Content: late-night noise and its effects on students/older residents (1); practical noise-reduction measures and consideration for others (1).')
p('16, Shanti Nagar\nJaipur\n16 September 2026\n\nThe Editor\nThe City Herald\nJaipur')
p('Subject: Disturbance caused by late-night loudspeakers', True)
p('Sir/Madam,')
models = {}
models['Q4A'] = 'Through your newspaper, I wish to draw attention to the frequent use of loudspeakers late at night in Shanti Nagar. Students find it difficult to concentrate or prepare for examinations, while older residents lose much-needed sleep. The disturbance continues even after repeated requests from neighbours.\n\nResidents should agree on suitable timings for celebrations and keep the volume low. Organisers could use indoor venues and inform neighbours beforehand. The residents’ association should encourage cooperation and provide a contact for reporting repeated disturbances.\n\nPlease highlight this issue and urge everyone to celebrate responsibly. A little consideration can protect peaceful study and rest without taking away the joy of community events.'
for text in models['Q4A'].split('\n\n'):
    p(text)
p('Yours faithfully,\nKavya')

page()
h('Q4(B). Complaint Letter — Model')
p('Content: accurate order/delivery details and both defects (1); request for replacements and missing items before the trials (1). Apply the common letter rubric.')
p('K.M. International School\n18, School Road\nJaipur\n16 September 2026\n\nThe Sales Manager\nChampion Sports\n8, Market Road\nJaipur')
p('Subject: Short delivery and damaged rackets — Order KM/SP/24', True)
p('Sir/Madam,')
models['Q4B'] = 'Our school ordered 20 badminton rackets under Order No. KM/SP/24 dated 3 September 2026. However, the delivery received on 11 September contained only 16 rackets. Three of these have cracked handles and are unsuitable for use.\n\nPlease deliver the four missing rackets and replace the three damaged ones before our school trials. Kindly arrange collection of the defective items and check the replacements carefully before dispatch. The shortage is disrupting practice and may prevent students from participating fully in the selection process.\n\nPlease acknowledge this complaint and confirm a prompt delivery schedule. We expect all 20 rackets to be available in good condition before the trials begin.'
for text in models['Q4B'].split('\n\n'):
    p(text)
p('Yours faithfully,\nKabir\nSports Secretary')
p('Quantity check: four missing rackets and three damaged rackets requiring replacement are separate issues. Accept Kabir or Kavya and a sensible date. Word counts for model letters refer to the body, excluding address blocks and conventional closings.', size=10)
h('Q5. Creative writing — Either A or B (5 marks)')
p('Diary: format 1 (date and diary opening); content 2 (quiz experience and memorable moment 1, personal learning 1); organisation and reflective voice 1; accuracy 1.')
p('Story: title and required opening 1 (½ each); content 2 (developed situation/problem 1, coherent resolution 1); organisation 1; accuracy 1. The suggested clues are optional.')

page()
h('Q5(A). Diary Entry — Model (100–120 words)')
p('16 September 2026\n8:30 p.m.\nDear Diary,')
models['Q5A'] = 'Today I represented our school in an inter-school quiz for the first time. My hands trembled as we took our seats, and I worried that I would forget everything. Then a difficult question about space came up. I remembered something from a library book, pressed the buzzer and answered correctly! Hearing the applause was wonderful.\n\nOur team did not win, but my teacher praised my courage and reminded us how much we had learnt. I returned home disappointed about the result yet proud that I had contributed. I now understand that success includes facing fear and supporting teammates. Next time, I will prepare more widely and trust myself.'
for text in models['Q5A'].split('\n\n'):
    p(text)
p('Kabir')
p('Accept any plausible difficult quiz question and relevant reflection. A time and signature are optional. The team must not be portrayed as the winner; credit the other fulfilled content points if this detail is missed.', size=10)
h('Q5(B). Story — Model (120–150 words)')
p('A Return Worth Waiting For', True)
models['Q5B'] = 'The bus had just left when I noticed a small bag lying beside the empty bench. A label read “Mrs Mehta”, but there was no address. My own bus was approaching, and I had an important practice session at school. Should I leave the bag and hurry away?\n\nA tea seller saw my hesitation and pointed towards the transport office. Together, we handed the bag to the clerk, who contacted the driver of the bus that had just departed. Meanwhile, I called my teacher and explained the delay.\n\nTwenty minutes later, a breathless woman arrived in an auto. After describing the bag and its contents, she received it from the clerk. Her important documents were safe. She thanked us warmly. I missed the start of practice, but my teacher smiled when I arrived. “Some decisions matter more than being first,” she said.'
for text in models['Q5B'].split('\n\n'):
    p(text)
p('Accept any original, coherent plot with a clear ending. Do not require these characters or this resolution. Retain the supplied opening; minor punctuation changes are acceptable.', size=10)

page()
h('SECTION C — LITERATURE (40 MARKS)')
h('Q6. The Pot Maker — All four parts (5 marks)')
q6 = [
    ['(i)', 'Onula/Aunty (½), addressing Sentila (½).', 1],
    ['(ii)', 'B — Reassuring.', 1],
    ['(iii)', 'Repeated failures leave Sentila tense and discouraged.', 1],
    ['(iv)', 'Watch her mother shape the mouth (1); she observes the slower strokes and added rim (1).', 2]
]
table(['Part', 'Expected answer and allocation', 'Marks'], q6, [.6, 5.6, .65])
p('In Q6(iv), a clearly explained relevant action earns the response mark; do not require all technical details. Award ½ for an incomplete but relevant response.', size=10)
h('Q7. Bharat Our Land — All three parts (5 marks)')
q7 = [
    ['(i)', 'B — Personification.', 1],
    ['(ii)', 'Generous.', 1],
    ['(iii)', 'The rhetorical question implies no equal (1), praises Ganga’s generosity/grace (1), and invites shared pride or admiration (1).', 3]
]
table(['Part', 'Expected answer and allocation', 'Marks'], q7, [.6, 5.6, .65])
p('For Q7(iii), students need not use the technical term “rhetorical question” if they explain its function. Merely naming the device without explanation is insufficient for the first mark. No separate expression mark applies.', size=10)
h('Short-answer marking')
p('Q8: any four of six, 40–50 words each. Award the specified 2 content marks plus 1 for clear, coherent expression. Use half marks for partial development or understandable expression with noticeable errors.')
p('Q9: any four of five, 30–40 words each. Award the two content marks indicated; no separate language mark. The following pages give marking points rather than full-length model responses.')

page()
h('Q8. Short answers — Any four of six (4 × 3 = 12)')
p('Award 1 for each content point below, plus up to 1 for expression. Accept equivalent explanations supported by the work.')
q8 = [
    ['(i) The Pot Maker', 'Coordinated hand/spatula movement shapes clay (1); excessive or insufficient firing ruins pots (1).'],
    ['(ii) Bharat Our Land', 'Repetition creates a celebratory tone (1) and unites readers in patriotic appreciation (1).'],
    ['(iii) Gifts of Grace', 'Singing suggests joyful work (1); distinct occupational voices combine in national harmony (1).'],
    ['(iv) Canvas of Soil', 'Planting expresses the gardener’s hopes (1); awaiting spring shows patient faith that effort will become colourful growth (1).'],
    ['(v) I Cannot Remember My Mother', 'He senses his mother’s gaze across the sky (1), suggesting enduring, comforting affection (1).'],
    ['(vi) Vitamin-M', 'Grandpa recalls complex chess games despite forgetting names (1); Ravi wonders at this contrast and recognises his skill (1).']
]
table(['Part / work', 'Content allocation'], q8, [1.65, 5.2])
p('Q8(i): the answer should cover both shaping and firing. Two details confined to only one process earn at most 1 content mark. Q8(iii): explain the relationship between individual voices and the whole, rather than merely listing jobs.', size=10)
p('Q8(vi): a clear explanation of the learner’s surprise at the contrast is sufficient for the response mark. Do not require a medical explanation of memory.', size=10)
link('Text reference: Kaveri, Unit 1 — Bharat Our Land', 'https://ncert.nic.in/textbook/pdf/iebe101.pdf')
link('Text reference: Kaveri, Unit 2 — The Pot Maker; Gifts of Grace', 'https://philoid.com/ncert/chapter/iebe102')
link('Text reference: Kaveri, Unit 3 — Canvas of Soil', 'https://ncert.nic.in/textbook/pdf/iebe103.pdf')
link('Text reference: Kaveri, Unit 4 — Vitamin-M; I Cannot Remember My Mother', 'https://philoid.com/ncert/chapter/iebe104')

page()
h('Q9. Short answers — Any four of five (4 × 2 = 8)')
q9 = [
    ['(i) Grandmother', 'Similar age/life identification (1); shared unfulfilled wish to visit Kashi (1). Accept reasoned admiration for the heroine’s generosity.'],
    ['(ii) The Pot Maker', 'Carrying clay uphill causes backache (1); pounding it soft is exhausting (1).'],
    ['(iii) Winds of Change', 'Bamboo — Bihar (1); palm leaves — Bengal or Odisha (1). Alternatives: cotton — Gujarat; brass — Rajasthan; sola — Bengal.'],
    ['(iv) Vitamin-M', 'He falls in his garden (1) and remains unaided overnight, making solitary living unsafe (1).'],
    ['(v) Grandmother', 'She cannot read the new episode (1) and feels helpless, embarrassed to ask others (1).']
]
table(['Part / work', 'Expected answer and allocation'], q9, [1.65, 5.2])
p('Q9(i): accept two distinct text-supported reasons; repeated versions of the same reason count once. Q9(ii): other specific difficulties in obtaining or preparing the material may replace the examples.', size=10)
p('Q9(iii): each complete material–region pair earns 1; a correct material with a missing or incorrect region earns ½. Require two different materials. Q9(iv): credit the incident that caused the move, rather than a different later incident.', size=10)
h('Long-answer rubric: Q10 and Q11')
p('Each question carries 5 marks and has an A/B choice. Expected length: 100–120 words. Award 3 content marks as indicated for the selected option, 1 for organisation/coherence and 1 for grammar/vocabulary accuracy. Use half marks for partially developed content or expression.')
p('The model answers illustrate the expected depth. Accept alternative wording and supported interpretations. General statements without relevant incidents cannot earn all content marks. Assess only one option per question.')

page()
h('Q10. Long answer — Either A or B (5 marks)')
h('Option A — How I Taught My Grandmother to Read')
p('Content: initial doubt (1); change with evidence (1); lesson learnt (1). Organisation and accuracy: 2.')
models['Q10A'] = 'Initially, the granddaughter laughs at the ambition of a woman aged sixty-two to learn the alphabet. She assumes that age and household duties will prevent success. Teaching changes this judgement: her grandmother completes extensive homework and practises steadily. The granddaughter rewards her achievement with a novel and feels proud when she reads its title and other details independently. The experience teaches her to assess willingness and effort instead of making assumptions about age. It also reveals the dignity of teaching when her grandmother honours her as a teacher. Learning becomes a shared achievement, requiring commitment from the student and affection and guidance from the teacher.'
p(models['Q10A'])
link('Text reference: NCERT Kaveri, Unit 1', 'https://ncert.nic.in/textbook/pdf/iebe101.pdf')
h('Option B — The Pot Maker')
p('Content: elders (1); Arenla (1); Onula (1), each explained. Organisation and accuracy: 2.')
models['Q10B'] = 'The village elders treat pottery as common heritage and remind the family that experts must teach willing learners. Their intervention challenges the idea that inherited knowledge is private property. Arenla provides practical training in gathering and preparing clay; later, her skilled movements become a model for observation. Onula contributes the encouragement that Sentila needs after repeated disappointment. She demonstrates patiently and directs the girl back to her mother for further learning. These contributions complement one another: community responsibility preserves the opportunity, practical instruction supplies technique, and sympathetic support builds confidence. Tradition survives through active teaching and attentive practice, rather than through possession of knowledge alone.'
p(models['Q10B'])
link('Text reference: Kaveri, Unit 2 — The Pot Maker', 'https://philoid.com/ncert/chapter/iebe102')

page()
h('Q11. Long answer — Either A or B (5 marks)')
h('Option A — Vitamin-M')
p('Content: birthday reversal (1); detective joke (1); changed assumptions (1). Organisation and accuracy: 2.')
models['Q11A'] = 'The birthday surprise reverses the family’s assumptions. Ravi’s mother thinks Grandpa has forgotten when Ravi was born, but he is celebrating his own birthday by giving everyone presents. She has forgotten the date, making his suggestion that she needs help with memory amusingly appropriate. The detective collection adds another joke: Grandpa recommends lessons in following someone without being fooled. This hints that he noticed Ravi’s secret pursuit and understood more than Ravi realised. The humour comes from the supposedly helpless elder outwitting his anxious protectors. His memory, generosity and playful intelligence show why the family should avoid judging his entire ability through occasional lapses.'
p(models['Q11A'])
p('Accept cautious wording about the ending’s implication; it need not be treated as an explicit confession by the character.', size=10)
link('Text reference: Kaveri, Unit 4 — Vitamin-M', 'https://philoid.com/ncert/chapter/iebe104')
h('Option B — Winds of Change')
p('Content: change over time (1); first regional example (1); second regional example (1). Link examples to identity. Organisation and accuracy: 2.')
models['Q11B'] = 'Hand fans began as practical objects for personal comfort and acquired importance through trade, ceremony and artistic display. Modern technology has reduced their everyday cooling role, while decorative and commercial uses help the craft continue. Regional designs still reveal local traditions. Rajasthan’s appliqué fans combine shaped fabric pieces with ornamental stitching, displaying its textile skills. Gujarat’s cotton fans use mirror work, giving them a distinctive decorative character. These examples show that an object can change its function without losing its cultural meaning. Preserving such craftsmanship keeps regional identities visible and creates opportunities for artisans, provided appreciation is accompanied by support for the people who make the fans.'
p(models['Q11B'])
link('Text reference: NCERT Kaveri, Unit 3 — Winds of Change', 'https://ncert.nic.in/textbook/pdf/iebe103.pdf')
h('Final marks check')
p('Reading: 20 | Grammar: 10 | Writing: 10 | Literature: 40 | TOTAL: 80', True)
p('Literature: Q6 5 + Q7 5 + Q8 12 + Q9 8 + Q10 5 + Q11 5 = 40.', size=10)

doc.core_properties.title = 'Class IX English Set B — Answer Key and Marking Scheme'
doc.core_properties.subject = '80 marks | All questions and internal choices'
doc.core_properties.author = 'K.M. International School'
doc.core_properties.comments = 'Prepared for the supplied Set B paper. Q3 any 10/12; Q8 any 4/6; Q9 any 4/5; Q7 1+1+3.'

for key, text in models.items():
    low, high = (120, 150) if key == 'Q5B' else (100, 120)
    count = len(text.split())
    print(f'{key}: {count} words')
    assert low <= count <= high, (key, count)
assert list(map(len, [q1, q2, q3, q6, q7, q8, q9])) == [10, 10, 12, 4, 3, 6, 5]
assert sum(row[2] for row in q6) == 5 and sum(row[2] for row in q7) == 5
assert sum(allocation) == 80
data = [[c.text for c in row.cells] for row in paper.tables[0].rows]
assert data[-1][1:] == ['180', '120']
assert 180 - 120 == 60 and 42 / 120 * 100 == 35 and (36 - 18) / 36 * 100 == 50
assert models['Q5B'].startswith('The bus had just left when I noticed a small bag lying beside the empty bench.')
assert len(doc.element.xpath('//w:br[@w:type="page"]')) == 11
doc.save(OUT)
check = Document(OUT)
assert len(check.tables) == 9
with zipfile.ZipFile(OUT) as package:
    assert package.testzip() is None
print(OUT.resolve())
print('Verified: Set B question coverage, all choices, 80 marks, calculations, model word limits and Word package integrity.')
