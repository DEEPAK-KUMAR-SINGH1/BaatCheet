from pathlib import Path
from copy import deepcopy
import re
import zipfile
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

folder = Path(__file__).parent
src = folder / 'Class_9_English_Half_Yearly_Question_Paper_Revised.docx'
out = folder / 'Class_9_English_Half_Yearly_Question_Paper_Set_B.docx'
doc = Document(src)
paras = doc.paragraphs
assert len(paras) == 125, 'Source structure changed; inspect before applying replacements.'
assert paras[49].text.startswith('Q3.') and paras[107].text.startswith('Q9.')

def setp(index, text):
    p = paras[index]
    props = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear()
    r = p.add_run(text)
    if props is not None: r._r.insert(0, props)

def setsource(index, label, url):
    p = paras[index]
    p.clear()
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '18'); props.append(size)
    run.append(props)
    text = OxmlElement('w:t'); text.text = label; run.append(text)
    link.append(run); p._p.append(link)

romans = ['i','ii','iii','iv','v','vi','vii','viii','ix','x','xi','xii']
def questions(start, items):
    for j, value in enumerate(items):
        body, marks = value if isinstance(value, tuple) else (value, 1)
        setp(start+j, f'({romans[j]})  {body}  [{marks}]')

setp(1, 'HALF-YEARLY EXAMINATION — SET B')
for section in doc.sections:
    section.header.paragraphs[0].text = 'K.M. INTERNATIONAL SCHOOL  |  CLASS IX — ENGLISH  |  SET B'
setp(7, '2. All questions are compulsory, subject to the internal choices provided. Attempt only the required number of parts.')

passage1 = [
    '1. Many people believe that being observant means having unusually sharp eyesight. Yet observation involves more than seeing clearly. It means paying attention, asking questions and noticing relationships between details. A student may walk past a tree every morning without noticing its changing leaves. Another may pause to compare its colour with the previous week. Both have seen the same tree, but only one has begun to investigate it. Ordinary surroundings can become a source of learning when we look at them with curiosity.',
    '2. During an art lesson, a teacher placed an old shoe on a table and asked the class to draw it. Several students quickly sketched a familiar shoe shape from memory. Others looked closely at the uneven laces, worn heel and small tear near the toe. Their drawings were not necessarily neater, but they captured the particular object in front of them. The activity showed that what we expect to see can sometimes prevent us from noticing what is actually there.',
    '3. Careful observation is useful beyond the art room. While reading, a student can notice how a character’s actions differ from the character’s words. During a science activity, a small change in colour may raise an important question. On the playground, watching how a ball bounces can suggest a better position for receiving it. In each case, attention provides information that a hurried glance might miss. However, noticing a detail is only the beginning; the observer must consider what it means.',
    '4. This is why observation should be separated from assumption. Seeing wet ground outside a classroom is an observation. Deciding immediately that it has rained is an assumption: someone may have washed the corridor or spilt water. A thoughtful observer considers more than one explanation and looks for further evidence. There is nothing wrong with making an initial guess, provided one is willing to revise it. Confidence becomes more useful when it is accompanied by the ability to admit uncertainty.',
    '5. Developing this habit does not require expensive equipment. Keeping a brief notebook, drawing a familiar object or listening carefully during a conversation can offer daily practice. The purpose is not to record every possible detail, but to select those relevant to a question. Equally, observation should respect other people’s privacy; curiosity does not justify reading someone else’s messages. Used responsibly, attention helps us understand our surroundings and other people more accurately. It replaces quick judgement with a willingness to discover.'
]
for i, text in enumerate(passage1,12): setp(i,text)
questions(20, [
    'Which statement best expresses the main idea?\nA. Good observation depends only on eyesight.\nB. Familiar places offer little opportunity for learning.\nC. Careful attention and questioning deepen understanding.\nD. A confident first impression is always correct.',
    'What does the second student notice about the tree that the first student overlooks?',
    'Why do some drawings show more of the shoe’s particular features than others?',
    'According to paragraph 2, how can expectations interfere with observation?',
    'Give one example from paragraph 3 of how observation helps a student outside an art lesson.',
    'State whether this statement is TRUE or FALSE: Noticing a detail automatically explains its meaning.',
    'Which statement is an observation rather than an assumption?\nA. The corridor floor is wet.\nB. It must have rained this morning.\nC. Someone was careless with a bucket.\nD. The cleaner has just washed the corridor.',
    'Find a word in paragraph 4 that means “change or reconsider an opinion”.',
    'Why does the writer suggest selecting relevant details rather than recording everything?',
    'A student opens a classmate’s private messages out of curiosity. Which principle in paragraph 5 does this action violate?'
])

passage2 = [
    '1. A school eco-club studied the waste sent from its lunch area for disposal. Student volunteers weighed five categories of waste during one five-day school week in July. After this first audit, the school introduced reusable containers, a paper-reuse tray and a composting facility for suitable food scraps. The volunteers repeated the measurements during a five-day week in August. The figures in the table are hypothetical data prepared for this examination.',
    '2. Both weeks had the same number of students using the lunch area. Volunteers used the same weighing scale and measured the waste at the same time each afternoon. The table shows the total mass sent for disposal over each week, not the total amount of material used in the school. Materials retained for reuse or placed in the composting facility were excluded from these totals.',
    '3. Waste was segregated into five categories before weighing. The club welcomed the fall in the overall total but noticed that one category had increased. Its members proposed inspecting the contents of that category before choosing further action. They also cautioned that two weeks of measurements could not establish a permanent change in habits. Repeating the audit over several months would provide stronger evidence about whether the improvement continued.'
]
for i, text in enumerate(passage2,33): setp(i,text)
rows = [
    ['Waste category', 'July: mass sent for disposal (kg)', 'August: mass sent for disposal (kg)'],
    ['Paper and cardboard', '48', '24'], ['Plastic packaging', '36', '18'],
    ['Food scraps', '60', '42'], ['Metal cans', '12', '6'],
    ['Other waste', '24', '30'], ['Total', '180', '120']
]
for row, values in zip(doc.tables[0].rows, rows):
    for cell, text in zip(row.cells, values):
        p = cell.paragraphs[0]
        props = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
        p.clear(); r=p.add_run(text)
        if props is not None: r._r.insert(0,props)
questions(37, [
    'What did the eco-club measure?\nA. All materials purchased by the school\nB. Waste sent from the lunch area for disposal\nC. The number of students using reusable containers\nD. The amount of compost produced',
    'Which waste category had the greatest mass in BOTH weeks?',
    'Calculate the reduction in the total mass sent for disposal from July to August.',
    'What percentage of the August total consisted of food scraps?',
    'By what percentage did plastic-packaging waste decrease from July to August?',
    'Which category increased, and by how many kilograms?',
    'State whether this statement is TRUE or FALSE: The August total includes all material placed in the composting facility.',
    'Mention one condition kept the same in both weeks to make the comparison fair.',
    'Why did the club recommend repeating the audit over several months?',
    'In paragraph 3, “segregated” means:\nA. burnt completely    B. separated into groups\nC. mixed together    D. transported elsewhere'
])

questions(51, [
    'Choose the correct tense:\nNisha __________ in this town since 2021 and still lives here.\nA. has lived    B. is living    C. will live    D. had lived',
    'Fill in the blank with the correct form of the verb in brackets:\nAt seven yesterday evening, I __________ (read) when the lights went out.',
    'Choose the modal that expresses lack of necessity:\nYou __________ bring a dictionary; one will be provided to everyone.\nA. must not    B. cannot    C. need not    D. should not',
    'Choose the correct verb:\nNeither of the two answers __________ correct.\nA. are    B. is    C. have been    D. were',
    'Choose the correct verb:\nThe shoes under the bench __________ mine.\nA. is    B. was    C. has been    D. are',
    'Choose the correct article:\nMy cousin studies at __________ university in Delhi.\nA. an    B. a    C. no article    D. both a and an',
    'Choose the appropriate determiner:\nHow __________ notebooks do we need for the workshop?\nA. much    B. little    C. many    D. any',
    'Choose the correct preposition:\nKabir is interested __________ learning a new language.\nA. on    B. at    C. for    D. in',
    'Identify the incorrect word and write its correction:\nWe did not went to the market yesterday.',
    'Identify the incorrect word and write its correction:\nThe child sang sweet at the annual function.',
    'Choose the correct verb form:\nAt this time tomorrow, our team __________ to the tournament.\nA. travelled    B. has travelled    C. will be travelling    D. was travelling',
    'Choose the appropriate determiner:\nThere isn’t __________ milk left for tea.\nA. many    B. several    C. few    D. any'
])

setp(66, 'You are Kabir/Kavya, a resident of 16, Shanti Nagar, Jaipur. Loudspeakers are used late into the night in your neighbourhood, disturbing students and older residents. Write a letter to the Editor of The City Herald, Jaipur, explaining the problem and suggesting practical steps to reduce noise and encourage consideration for others.')
setp(69, 'You are Kabir/Kavya, Sports Secretary of K.M. International School, 18, School Road, Jaipur. Your school ordered 20 badminton rackets from Champion Sports, 8, Market Road, Jaipur, under Order No. KM/SP/24 dated 3 September 2026. The delivery received on 11 September contained only 16 rackets, three of which had cracked handles. Write a complaint to the Sales Manager requesting replacement of the damaged rackets and delivery of the missing items before the school trials.')
setp(74, 'You represented your school in an inter-school quiz for the first time. Although your team did not win, you answered a difficult question and received encouragement from your teacher. Write a diary entry in 100–120 words describing your nervousness, a memorable moment and what you learnt from the experience. Include a suitable date and an appropriate diary opening.')
setp(78, '“The bus had just left when I noticed a small bag lying beside the empty bench.”')
setp(79, 'You may use these clues: a name on a label — a difficult decision — help from a stranger — a grateful owner. The clues are optional; develop a coherent plot with your own details.')

setp(82, '“Don’t worry, little one, I shall teach you how to make a perfect pot.”')
setsource(83, '— The Pot Maker, Temsula Ao; NCERT, Kaveri', 'https://philoid.com/ncert/chapter/iebe102')
questions(85, [
    ('Identify the speaker and the person being addressed.',1),
    ('Which word best describes the speaker’s tone?\nA. Threatening    B. Reassuring    C. Mocking    D. Impatient',1),
    ('Why does the listener need reassurance at this moment?',1),
    ('What advice does the speaker later give about completing the pot’s mouth, and how does the listener act on it?',2)
])
setp(92, 'The generous Ganga is ours—\nwhich other river can match her grace?')
setsource(93, '— Bharat Our Land, Subramania Bharati; NCERT, Kaveri', 'https://ncert.nic.in/textbook/pdf/iebe101.pdf')
questions(94, [
    ('Giving a river human qualities is an example of:\nA. Onomatopoeia    B. Personification    C. Simile    D. Alliteration',1),
    ('Which adjective in the extract suggests a willingness to give freely?',1),
    ('Explain how the poet’s question expresses admiration rather than seeking information. Refer to the river’s qualities and the effect on the reader.',3)
])
questions(99, [
    ('Why do shaping and firing clay require care? Use two details from The Pot Maker to explain the skill involved.',3),
    ('How does the repeated call to praise the country affect the tone and message of Bharat Our Land?',3),
    ('How do the images of singing and different voices create a sense of unity in Gifts of Grace: Honouring Our Vocations?',3),
    ('How does waiting for spring connect the gardener’s work with hope in Canvas of Soil?',3),
    ('What does the speaker feel while looking into the distant sky? Explain how this image conveys affection in I Cannot Remember My Mother.',3),
    ('How does Grandpa’s knowledge of chess challenge a simple judgement about his memory? Explain Ravi’s response. (Vitamin-M)',3)
])
questions(109, [
    ('Why is the grandmother drawn to the main character of Kashi Yatre? Give two reasons. (How I Taught My Grandmother to Read)',2),
    ('Mention two difficulties Arenla faces in obtaining and preparing clay. (The Pot Maker)',2),
    ('Name two materials used in traditional hand fans and identify a region associated with each. (Winds of Change)',2),
    ('What incident leads Grandpa’s daughter to bring him to live with her, and why does it worry her? (Vitamin-M)',2),
    ('Why does the grandmother feel distressed when her granddaughter is away at a wedding? (How I Taught My Grandmother to Read)',2)
])
setp(116, 'A. At first, the granddaughter doubts her grandmother’s ambition; later, she takes pride in her achievement. Trace this change and explain what the experience teaches the granddaughter about learning and teaching. (How I Taught My Grandmother to Read)')
setp(118, 'B. Traditional knowledge survives when skilled people share it. Discuss how the village elders, Arenla and Onula contribute, in different ways, to passing pottery skills to Sentila. (The Pot Maker)')
setp(121, 'A. Explain how the birthday surprise and the gift of detective stories create humour at the end of Vitamin-M. How do these events challenge the family’s assumptions about Grandpa?')
setp(123, 'B. The role of traditional hand fans has changed with time, yet their designs still express regional identity. Discuss this change with two examples of regional craftsmanship from Winds of Change.')
setp(124, '— END OF QUESTION PAPER: SET B —')

doc.core_properties.title = 'Class IX English Half-Yearly Question Paper — Set B'
doc.core_properties.subject = 'Parallel Set B | 80 marks | 3 hours | Revised question pattern'
doc.core_properties.comments = 'Set B follows the eight prescribed works and revised choices: Q3 10 of 12; Q9 4 of 5; Q4 Editor/Complaint; Q7 three parts, 1+1+3 marks, without a rhyme-scheme question.'
doc.save(out)

check = Document(out)
text = '\n'.join(p.text for p in check.paragraphs)
def block(a,b): return text[text.index(a):text.index(b)]
for a,b,count in [('Q3.','Q4.',12),('Q6.','Q7.',4),('Q7.','Q8.',3),('Q8.','Q9.',6),('Q9.','Q10.',5)]:
    assert len(re.findall(r'^\([ivx]+\)',block(a,b),re.M)) == count
assert sum(map(int,re.findall(r'\[(\d+)\]',block('Q7.','Q8.')))) == 5
assert 'rhyme scheme' not in text.lower()
assert sum([10,10,10,5,5,5,5,12,8,5,5]) == 80
assert sum([48,36,60,12,24]) == 180
assert sum([24,18,42,6,30]) == 120
assert round(42/120*100) == 35
assert 400 <= len(' '.join(passage1).split()) <= 450
assert 200 <= len(' '.join(passage2).split()) <= 250
assert len(check.tables) == 1
assert len(check.element.xpath('//w:br[@w:type="page"]')) == len(doc.element.xpath('//w:br[@w:type="page"]'))
with zipfile.ZipFile(out) as z: assert z.testzip() is None
print(out.resolve())
print('Verified: 80 marks; all revised internal choices; 1+1+3 marks in Q7; new passages and questions.')
print('Passage word counts:',len(' '.join(passage1).split()),len(' '.join(passage2).split()))
